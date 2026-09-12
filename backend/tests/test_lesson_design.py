from __future__ import annotations

import tempfile
import json
import unittest
from unittest.mock import Mock, patch
from pathlib import Path
from zipfile import ZipFile

from backend.app.config import AppConfig
from backend.app.runtime import WebGISRuntime
from backend.app.store import RuntimeStore


class LessonDesignServiceTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        root_dir = Path(__file__).resolve().parents[2]
        config = AppConfig(root_dir=root_dir)
        config.data_dir = Path(self.temp_dir.name) / "backend" / "data"
        config.state_dir = config.data_dir / "state"
        config.uploads_dir = config.data_dir / "uploads"
        config.outputs_dir = config.data_dir / "outputs"
        config.state_file = config.state_dir / "runtime.json"
        config.minimax_api_key = ""
        config.minimax_token_plan_key = ""
        config.ensure_dirs()
        self.store = RuntimeStore(config.state_file)
        self.runtime = WebGISRuntime(config=config, store=self.store)
        self.project = self.runtime.create_project()["project_id"]
        self.addCleanup(self.temp_dir.cleanup)

    def test_shanghai_seed_keeps_old_draft_and_original_scenes_questions_homework(self) -> None:
        service = self.runtime.classroom.lesson_design
        old = service.create_or_resume(self.project, "local_admin")
        old_before = json.dumps(old.to_dict(), ensure_ascii=False, sort_keys=True)
        base = self.store.get_lesson("lesson_builtin_population_shanghai_world")
        base_before = json.dumps(base.to_dict(), ensure_ascii=False, sort_keys=True)
        design = service.create_or_resume(self.project, "local_admin", base.lesson_id)
        self.assertNotEqual(design.design_id, old.design_id)
        self.assertEqual(len(design.draft["stages"]), 8)
        self.assertEqual(sum(len(s["questions"]) for s in design.draft["stages"]), 12)
        self.assertEqual(design.draft["stages"], base.stages)
        self.assertEqual(design.draft["homework"], base.plan["homework"])
        self.assertEqual(design.draft["stages"][0]["scene"]["view"]["center"], [121.47, 31.23])
        self.assertEqual(service.create_or_resume(self.project, "local_admin", base.lesson_id).design_id, design.design_id)
        design.draft["stages"][0]["title"] = "教师草稿修改"
        self.assertEqual(json.dumps(base.to_dict(), ensure_ascii=False, sort_keys=True), base_before)
        self.assertEqual(json.dumps(self.store.get_lesson_design(old.design_id).to_dict(), ensure_ascii=False, sort_keys=True), old_before)

    def test_legacy_structured_objectives_resume_as_text_without_read_mutation(self) -> None:
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        stored = self.store.get_lesson_design(design["design_id"])
        stored.draft["objectives"] = [{"id": 1, "statement": "描述人口分布", "type": "基础"}, "比较人口密度"]
        stored.draft["core_questions"] = {"core": "为什么分布不均？", "sub_questions": [{"statement": "哪里人口密集？", "objective_refs": [1]}]}
        self.store.upsert_lesson_design(stored)
        resumed = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        self.assertEqual(resumed["draft"]["objectives"], ["描述人口分布", "比较人口密度"])
        self.assertEqual(resumed["draft"]["core_questions"]["sub_questions"], ["哪里人口密集？"])
        self.assertEqual(resumed["draft"]["structured_text_originals"]["objectives"][0]["type"], "基础")
        self.assertIsInstance(self.store.get_lesson_design(design["design_id"]).draft["objectives"][0], dict)
        self.assertEqual(resumed["revision"], design["revision"])

    def test_direct_objective_and_core_question_edits_accept_section_payloads(self) -> None:
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        result = self.runtime.classroom.resolve_lesson_design(design["design_id"], "objectives", "edit", "", 0, ["描述分布", "比较密度"])["design"]
        self.assertEqual(result["draft"]["objectives"], ["描述分布", "比较密度"])
        result = self.runtime.classroom.resolve_lesson_design(design["design_id"], "core_questions", "edit", "", result["revision"], {"core": "为何不均？", "sub_questions": ["稠密区在哪？"]})["design"]
        self.assertEqual(result["draft"]["core_questions"]["core"], "为何不均？")
        with self.assertRaisesRegex(ValueError, "缺少可读文字"):
            self.runtime.classroom.resolve_lesson_design(design["design_id"], "objectives", "edit", "", result["revision"], [{"id": 1}])
        self.assertEqual(self.store.get_lesson_design(design["design_id"]).draft["objectives"], ["描述分布", "比较密度"])

    def test_model_text_objects_are_normalized_at_merge_and_unknown_shapes_rejected(self) -> None:
        service = self.runtime.classroom.lesson_design
        payload = service._validate_model_payload({"section_patch": {"objectives": [{"statement": "描述分布"}], "core_questions": {"sub_questions": [{"text": "人口在哪里？"}]} }})
        draft = {}
        service._merge_patch(draft, payload["section_patch"])
        self.assertEqual(draft["objectives"], ["描述分布"])
        self.assertEqual(draft["core_questions"]["sub_questions"], ["人口在哪里？"])
        self.assertIsNone(service._validate_model_payload({"section_patch": {"objectives": [{"id": 1}]}}))

    def test_explicit_schedule_is_not_silently_replaced_on_model_failure(self) -> None:
        service = self.runtime.classroom.lesson_design
        design = service.create_or_resume(self.project, "local_admin")
        stored = self.store.get_lesson_design(design.design_id)
        stored.draft["stages"] = [{"title": "已有课堂活动", "minutes": 40}]
        self.store.upsert_lesson_design(stored)
        message = "请生成40分钟教学过程。环节安排为世界人口分布读图5分钟、总量与密度辨析8分钟、胡焕庸线及成因探究15分钟、上海迁移应用8分钟、总结评价4分钟。"
        schedule = service._requested_schedule(message)
        self.assertEqual([s["minutes"] for s in schedule], [5, 8, 15, 8, 4])
        bad_results = [None, {"section_patch": {"stages": service._make_stages("人口分布", 40)}},
                       {"section_patch": {"stages": [{**item, "minutes": 8} for item in schedule]}}]
        for result in bad_results:
            with self.subTest(result=result), patch.object(service, "_ask_minimax", return_value=result):
                with self.assertRaisesRegex(ValueError, "原草稿已保留"):
                    service.turn(design.design_id, message, 0, "process")
                unchanged = service.get(design.design_id)
                self.assertEqual(unchanged.revision, 0)
                self.assertEqual(unchanged.turns, [])
                self.assertEqual(unchanged.draft["stages"], [{"title": "已有课堂活动", "minutes": 40}])
        with patch.object(service, "_ask_minimax", return_value={"section_patch": {"stages": schedule}, "reply": "已按五个环节起草"}):
            result = service.turn(design.design_id, message, 0, "process")
        self.assertEqual(result["draft"]["stages"], schedule)
        self.assertEqual(result["generation_mode"], "model")

    def test_process_generation_has_headroom_and_diagnostics_do_not_log_response_secrets(self) -> None:
        service = self.runtime.classroom.lesson_design
        design = service.create_or_resume(self.project, "local_admin")
        client = Mock()
        client.chat_completion.return_value = '{"section_patch": {}}'
        service.minimax_client = client
        service._ask_minimax(design, "process", "生成课堂过程")
        self.assertEqual(client.chat_completion.call_args.kwargs["extra_payload"]["max_completion_tokens"], 12288)
        self.assertEqual(client.chat_completion.call_args.kwargs["timeout"], 90.0)
        client.chat_completion.side_effect = TimeoutError("sensitive-provider-response")
        with self.assertLogs("backend.app.services.lesson_design", level="WARNING") as captured:
            self.assertIsNone(service._ask_minimax(design, "process", "生成课堂过程"))
        self.assertIn("TimeoutError", " ".join(captured.output))
        self.assertNotIn("sensitive-provider-response", " ".join(captured.output))

    def test_invalid_model_json_is_corrected_once_and_failures_are_bounded(self) -> None:
        service = self.runtime.classroom.lesson_design
        design = service.create_or_resume(self.project, "local_admin")
        client = Mock()
        service.minimax_client = client
        client.chat_completion.side_effect = ['{"section_patch":', '{"reply":"修正后的草稿","section_patch":{"board_design":"读图与归因"}}']
        result = service._ask_minimax(design, "process", "设计课堂过程")
        self.assertEqual(result["section_patch"]["board_design"], "读图与归因")
        self.assertEqual(client.chat_completion.call_count, 2)
        self.assertIn("未通过系统校验", client.chat_completion.call_args.args[0][-1]["content"])
        client.reset_mock()
        client.chat_completion.side_effect = ['{"section_patch":[]}', '{"section_patch":[]}']
        self.assertIsNone(service._ask_minimax(design, "process", "设计课堂过程"))
        self.assertEqual(client.chat_completion.call_count, 2)
        client.reset_mock()
        client.chat_completion.side_effect = TimeoutError("timeout")
        self.assertIsNone(service._ask_minimax(design, "process", "设计课堂过程"))
        self.assertEqual(client.chat_completion.call_count, 1)

    def test_reading_design_validates_same_snapshot_without_model_or_store_changes(self) -> None:
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        before = self.store.state_file.read_bytes()
        with patch.object(self.runtime.classroom.lesson_design, "_ask_minimax") as ask:
            first = self.runtime.classroom.get_lesson_design(design["design_id"])
            second = self.runtime.classroom.get_lesson_design(design["design_id"])
        ask.assert_not_called()
        self.assertEqual(first["revision"], design["revision"])
        self.assertEqual(first["turns"], design["turns"])
        self.assertFalse(first["rehearsal_report"]["ready"])
        self.assertEqual(first["rehearsal_report"], second["rehearsal_report"])
        self.assertEqual(self.store.state_file.read_bytes(), before)

    def test_model_and_rules_cannot_rewrite_or_drop_protected_question_snapshots(self) -> None:
        service = self.runtime.classroom.lesson_design
        design = service.create_or_resume(self.project, "local_admin")
        original = {"question_id": "q1", "text": "哪些因素共同影响人口分布？", "source": "teacher_manual", "answer": "自然与人文因素", "explanation": "需基于地图解释"}
        record = self.store.get_lesson_design(design.design_id)
        record.draft["stages"] = [{"stage_id": "s1", "questions": [original]}]
        self.store.upsert_lesson_design(record)
        for candidate in ([], [{**original, "text": "改写题目"}], [original, original], [{**original, "answer": "错误答案"}]):
            with self.subTest(candidate=candidate), patch.object(service, "_ask_minimax", return_value={"section_patch": {"stages": [{"questions": candidate}]}}):
                with self.assertRaisesRegex(ValueError, "需保留的题目"):
                    service.turn(design.design_id, "修改教学活动", 0, "process")
                self.assertEqual(service.get(design.design_id).revision, 0)
                self.assertEqual(service.get(design.design_id).draft["stages"][0]["questions"], [original])
        with patch.object(service, "_ask_minimax", return_value=None):
            with self.assertRaisesRegex(ValueError, "需保留的题目"):
                service.turn(design.design_id, "修改教学活动", 0, "process")
        relocated = [{"stage_id": "s2", "questions": [original]}]
        with patch.object(service, "_ask_minimax", return_value={"section_patch": {"stages": relocated}}):
            result = service.turn(design.design_id, "调整活动顺序", 0, "process")
        self.assertEqual(result["draft"]["stages"], relocated)
        open_question = {"question_id": "q2", "text": "原开放题"}
        self.assertEqual(service._questions_to_preserve({"stages": [{"questions": [open_question]}]}, "保留现有三道开放题及原文"), [open_question])

    def test_dataset_context_uses_local_values_prioritizes_named_regions_and_omits_geometry(self) -> None:
        service = self.runtime.classroom.lesson_design
        data_path = Path(self.temp_dir.name) / "population.geojson"
        features = [{"properties": {"name": f"区域{i}", "population": i * 100, "source_year": "2020"}, "geometry": {"type": "Point", "coordinates": [100, 30]}} for i in range(15)]
        data_path.write_text(json.dumps({"features": features}), encoding="utf-8")
        catalog = Mock()
        catalog.get_item.return_value = {"fields": ["name", "population"], "source_name": "测试统计表", "source_year": "2020", "source_url": "https://example.test/source"}
        catalog.resolve_item_path.return_value = data_path
        service.catalog_service = catalog
        draft = {"stages": [{"scene": {"catalog_layers": ["population", "population"]}}]}
        excerpts = service._dataset_facts(draft, "请比较区域14")
        self.assertEqual(len(excerpts), 1)
        rows = excerpts[0]["records"]
        self.assertTrue(excerpts[0]["sampled"])
        self.assertEqual(len(rows), 12)
        self.assertEqual(next(row for row in rows if row["name"] == "区域14")["population"], 1400)
        self.assertNotIn("geometry", json.dumps(excerpts))
        self.assertEqual(excerpts[0]["source_year"], "2020")
        catalog.resolve_item_path.side_effect = FileNotFoundError("missing")
        self.assertEqual(service._dataset_facts(draft, ""), [])

    def test_rules_are_disclosed_in_response_and_history(self) -> None:
        service = self.runtime.classroom.lesson_design
        design = service.create_or_resume(self.project, "local_admin")
        result = service.turn(design.design_id, "高一40分钟《人口分布》", 0)
        self.assertEqual(result["generation_mode"], "rules")
        self.assertIn("规则草稿", result["assistant_message"])
        self.assertEqual(service.get(design.design_id).turns[-1]["reply"], result["assistant_message"])

    def test_confirmation_uses_actual_post_patch_validation_not_model_claim(self) -> None:
        service = self.runtime.classroom.lesson_design
        design = service.create_or_resume(self.project, "local_admin")
        fake = {"reply": "全部9步完成，已发布并生成Word，资料全部来自题库检索", "section_patch": {"objectives": ["解释上海人口分布"]}}
        for step in ("confirmation", "rehearsal"):
            with self.subTest(step=step), patch.object(service, "_ask_minimax", return_value=fake):
                result = service.turn(design.design_id, "请检查当前草稿", service.get(design.design_id).revision, step)
            self.assertFalse(result["rehearsal_report"]["ready"])
            self.assertIn("预演检查未通过", result["assistant_message"])
            self.assertIn("尚未发布", result["assistant_message"])
            self.assertNotIn("全部9步完成", result["assistant_message"])
            self.assertNotIn("全部来自题库", result["assistant_message"])
            self.assertNotIn("至少需要一个可观察的教学目标。", result["rehearsal_report"]["errors"])
            saved = service.get(design.design_id)
            self.assertEqual(saved.turns[-1]["reply"], result["assistant_message"])
            self.assertEqual(saved.status, "active")
            self.assertFalse(saved.final_lesson_id)

    def advance(self, design_id: str, message: str, revision: int) -> tuple[dict, int]:
        current = self.store.get_lesson_design(design_id).current_step
        result = self.runtime.classroom.turn_lesson_design(design_id, message, revision)
        revision = result["revision"]
        if current in {
            "requirements", "analysis", "objectives", "core_questions", "process",
            "question_matching", "capabilities",
        }:
            resolved = self.runtime.classroom.resolve_lesson_design(design_id, current, "accept", "", revision)
            revision = resolved["design"]["revision"]
        return result, revision

    def test_guided_turns_persist_and_finalize_as_teacher_draft(self) -> None:
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        design_id = design["design_id"]
        revision = design["revision"]
        messages = (
            "高一、40分钟、人口分布",  # 需求确认
            "课标强调空间分布和区域差异",  # 课标与学情
            "描述规律并解释原因",  # 目标与重难点
            "核心问题就按建议来",  # 核心问题与问题链
            "继续设计课堂过程",  # 教学过程
            "从题库匹配题目并补作业",  # 题目匹配
            "使用二维地图",  # GIS/AI能力
        )
        for message in messages:
            result, revision = self.advance(design_id, message, revision)
        result = self.runtime.classroom.turn_lesson_design(design_id, "运行预演", revision)
        revision = result["revision"]
        result = self.runtime.classroom.turn_lesson_design(design_id, "补充设计思路与反思", revision)
        revision = result["revision"]
        persisted = self.store.get_lesson_design(design_id)
        self.assertIsNotNone(persisted)
        self.assertEqual(persisted.revision, revision)
        self.assertEqual(len(persisted.turns), 9)
        result = self.runtime.classroom.finalize_lesson_design(design_id, revision)
        self.assertEqual(result["lesson"]["source"], "assistant_draft")
        self.assertEqual(result["lesson"]["metadata"]["project_id"], self.project)
        self.assertEqual(self.store.get_lesson_design(design_id).status, "finalized")

    def test_store_round_trip_keeps_design_and_lesson_plan(self) -> None:
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        result = self.runtime.classroom.turn_lesson_design(design["design_id"], "高一、40分钟、胡焕庸线", 0)
        self.assertEqual(result["revision"], 1)
        reloaded = RuntimeStore(self.store.state_file)
        restored = reloaded.get_lesson_design(design["design_id"])
        self.assertIsNotNone(restored)
        self.assertEqual(restored.draft["topic"], "胡焕庸线")

    def test_natural_requirement_sentence_keeps_short_book_title(self) -> None:
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        message = (
            "高一必修二《人口分布》，单课时40分钟。学生已经学过人口密度，但容易把人口总量、"
            "人口密度和实时人口混为一谈。请围绕胡焕庸线安排Top20数据探究和课后复盘。"
        )
        result = self.runtime.classroom.turn_lesson_design(design["design_id"], message, 0)
        self.assertEqual(result["draft"]["title"], "人口分布")
        self.assertEqual(result["draft"]["topic"], "人口分布")
        self.assertEqual(result["draft"]["grade"], "高一")
        self.assertEqual(result["draft"]["duration_minutes"], 40)
        self.assertEqual(result["draft"]["requirements"]["raw"], message)
        self.assertNotIn("学生已经学过", result["draft"]["title"])

    def test_rehearsal_blocks_a_plan_that_does_not_fill_the_declared_class_time(self) -> None:
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        record = self.store.get_lesson_design(design["design_id"])
        record.draft.update({
            "title": "人口分布",
            "grade": "高一",
            "duration_minutes": 40,
            "objectives": ["描述人口分布规律"],
            "stages": [
                {
                    "stage_id": "s1", "title": "导入", "minutes": 8,
                    "content": "观察人口分布图", "design_intent": "发现空间差异",
                    "questions": [{"text": "人口主要分布在哪里？"}],
                },
                {
                    "stage_id": "s2", "title": "探究", "minutes": 18,
                    "content": "分析胡焕庸线", "design_intent": "解释空间格局",
                    "questions": [{"text": "界线两侧为何差异明显？"}],
                },
                {
                    "stage_id": "s3", "title": "复盘", "minutes": 13,
                    "content": "完成证据链", "design_intent": "当堂检测目标",
                    "questions": [{"text": "如何用证据说明人口分布规律？"}],
                },
            ],
        })
        self.store.upsert_lesson_design(record)

        report = self.runtime.classroom.lesson_design.rehearse(design["design_id"])

        self.assertFalse(report["ready"])
        self.assertEqual(report["total_minutes"], 39)
        self.assertTrue(any("与课堂时长 40 分钟不一致" in item for item in report["errors"]))

    def test_direct_edit_revision_conflict_and_confirmed_section_stability(self) -> None:
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        design_id = design["design_id"]
        edited = self.runtime.classroom.resolve_lesson_design(
            design_id,
            "requirements",
            "edit",
            "",
            design["revision"],
            {"requirements": {"raw": "高一、45分钟、人口迁移"}},
        )["design"]
        accepted = self.runtime.classroom.resolve_lesson_design(
            design_id, "requirements", "accept", "", edited["revision"]
        )["design"]
        self.assertEqual(accepted["section_status"]["requirements"], "confirmed")
        with self.assertRaisesRegex(ValueError, "已更新"):
            self.runtime.classroom.resolve_lesson_design(
                design_id, "title", "edit", "", edited["revision"], "过期修改"
            )
        result = self.runtime.classroom.turn_lesson_design(
            design_id, "继续分析学生已有基础", accepted["revision"], "analysis"
        )
        self.assertEqual(result["draft"]["requirements"]["raw"], "高一、45分钟、人口迁移")
        self.assertEqual(result["section_status"]["requirements"], "confirmed")

    def test_turn_back_instruction_short_circuits_flow_control(self) -> None:
        """回归（用户验收反馈）：「返回上一步」是控制指令，不当需求文本解析。

        曾在第一步点「返回上一步」被规则回退当成课题（“我理解你想做一节
        ‘返回上一步’”）并生成 section_patch。现在短路：不进 LLM/规则解析。
        """
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        design_id = design["design_id"]

        # 第一步点「返回上一步」：不回退、不产生需求 patch
        result = self.runtime.classroom.turn_lesson_design(design_id, "返回上一步，重新讨论上一部分")
        self.assertIn("已在第一步", result["assistant_message"])
        self.assertNotIn("返回上一步", str(result["draft"].get("title") or ""))
        self.assertNotIn("返回上一步", str(result["draft"].get("topic") or ""))
        self.assertEqual(result["review_sections"], [])
        self.assertEqual(result["next_step"], "requirements")

        # 发一条真实需求并确认，推进到第二步，再点「返回上一步」
        advanced = self.runtime.classroom.turn_lesson_design(
            design_id, "高一、40分钟、人口分布", result["revision"]
        )
        accepted = self.runtime.classroom.resolve_lesson_design(
            design_id, "requirements", "accept", "", advanced["revision"]
        )["design"]
        self.assertEqual(accepted["current_step"], "analysis")
        back = self.runtime.classroom.turn_lesson_design(
            design_id, "返回上一步", accepted["revision"]
        )
        self.assertIn("已回到「需求确认」", back["assistant_message"])
        self.assertEqual(back["next_step"], "requirements")
        stored = self.store.get_lesson_design(design_id)
        self.assertEqual(stored.current_step, "requirements")
        # 回退不污染草稿：课题保持真实需求解析结果，而不是控制指令文本
        self.assertNotIn("返回上一步", str(stored.draft.get("title") or ""))

    def test_direct_edit_stages_stores_array_and_plan_items_stay_renderable(self) -> None:
        """回归（E2E 验收发现）：直接编辑教学过程必须把数组写入 draft.stages。

        前端曾把 ``{"stages": [...]}`` 包装对象作为 value 提交，SECTION 级
        resolve edit 会把它原样写入 ``draft.stages``，双端渲染环节列表时
        ``.map``/``stage.get`` 崩溃（前端白屏、后端 plan_items AttributeError）。
        """
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        design_id = design["design_id"]
        revision = design["revision"]
        for message in ("高一、40分钟、人口分布", "课标强调空间分布", "描述规律", "核心问题按建议", "继续设计课堂过程"):
            _, revision = self.advance(design_id, message, revision)
        stages_value = [
            {
                "stage_id": "s1", "title": "情境导入", "minutes": 10, "material": "人口密度图",
                "question_chain": ["分布有何差异？"], "teacher_activities": ["引导读图"],
                "student_activities": ["观察并描述"], "knowledge_conclusion": "分布不均",
                "design_intent": "从证据进入",
            },
            {
                "stage_id": "s2", "title": "成因探究", "minutes": 20, "material": "地形气候图",
                "question_chain": ["哪些因素？"], "teacher_activities": ["组织对比"],
                "student_activities": ["小组归因"], "knowledge_conclusion": "自然加人文",
                "design_intent": "由描述到解释",
            },
            {
                "stage_id": "s3", "title": "归纳迁移", "minutes": 10, "material": "板书要点",
                "question_chain": ["如何迁移？"], "teacher_activities": ["收束规律"],
                "student_activities": ["结论卡片"], "knowledge_conclusion": "方法迁移",
                "design_intent": "检查目标",
            },
        ]
        result = self.runtime.classroom.resolve_lesson_design(
            design_id, "stages", "edit", "", revision, stages_value
        )
        stored = self.store.get_lesson_design(design_id)
        self.assertIsInstance(stored.draft["stages"], list)
        self.assertEqual(
            [stage["title"] for stage in stored.draft["stages"]],
            ["情境导入", "成因探究", "归纳迁移"],
        )
        # 会话视图（plan_items）必须可渲染，不得再抛 'str' object has no attribute 'get'
        stored.current_step = "process"  # 教学过程步骤的 plan_items 会逐个映射环节
        plan_items = self.runtime.classroom.lesson_design.session_view(stored)["plan_items"]
        stage_items = [item for item in plan_items if item["section_key"] == "stages"]
        self.assertEqual(len(stage_items[0]["value"]), 3)
        self.assertEqual(stored.section_status.get("stages"), "proposed")

    def test_rehearsal_rejects_unknown_registered_capabilities(self) -> None:
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        with self.assertRaisesRegex(ValueError, "不可用系统能力"):
            self.runtime.classroom.resolve_lesson_design(
                design["design_id"],
                "capabilities",
                "edit",
                "",
                design["revision"],
                {"capabilities": [{"id": "invented-capability"}]},
            )
        record = self.store.get_lesson_design(design["design_id"])
        record.draft.update({
            "title": "人口分布",
            "objectives": ["描述人口分布规律"],
            "stages": [{
                "stage_id": "s1", "title": "读图", "minutes": 40,
                "content": "观察专题图", "design_intent": "形成空间认识",
                "questions": [{"text": "人口主要分布在哪里？"}],
                "scene": {
                    "templates": ["invented-template"],
                    "catalog_layers": ["invented-dataset"],
                    "globe": {"enabled": True, "themes": ["invented-theme"]},
                },
            }],
        })
        self.store.upsert_lesson_design(record)
        report = self.runtime.classroom.lesson_design.rehearse(design["design_id"])
        self.assertFalse(report["ready"])
        self.assertTrue(any("invented-template" in item for item in report["errors"]))
        self.assertTrue(any("invented-dataset" in item for item in report["errors"]))
        self.assertTrue(any("invented-theme" in item for item in report["errors"]))

    def test_cross_project_base_and_export_are_rejected(self) -> None:
        other_project = self.runtime.create_project()["project_id"]
        lesson = self.runtime.classroom.lesson_service.create_lesson({
            "title": "项目一课时",
            "subject": "地理",
            "grade": "高一",
            "objectives": [],
            "stages": [],
            "metadata": {"project_id": self.project},
            "plan": {},
        }, source="assistant_draft", owner_user_id="local_admin")
        with self.assertRaisesRegex(ValueError, "不属于当前项目"):
            self.runtime.classroom.create_lesson_design(other_project, "local_admin", lesson.lesson_id)
        with self.assertRaisesRegex(ValueError, "不属于当前项目"):
            self.runtime.classroom.export_lesson_docx(lesson.lesson_id, other_project)

    def test_docx_export_has_five_column_process_and_page_field(self) -> None:
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        revision = 0
        messages = (
            "高一、40分钟、人口分布",
            "课标",
            "描述规律",
            "核心问题就按建议来",
            "设计过程",
            "从题库匹配题目并补作业",
            "二维地图",
        )
        for message in messages:
            result, revision = self.advance(design["design_id"], message, revision)
        result = self.runtime.classroom.turn_lesson_design(design["design_id"], "运行预演", revision)
        revision = result["revision"]
        result = self.runtime.classroom.turn_lesson_design(design["design_id"], "补充设计思路与反思", revision)
        revision = result["revision"]
        lesson = self.runtime.classroom.finalize_lesson_design(design["design_id"], revision)["lesson"]
        exported = self.runtime.classroom.export_lesson_docx(lesson["lesson_id"], self.project, design["design_id"])
        path = Path(exported["artifact"]["path"])
        self.assertTrue(path.is_file())
        from docx import Document
        document = Document(path)
        self.assertGreaterEqual(len(document.tables), 2)
        self.assertEqual([cell.text for cell in document.tables[1].rows[0].cells], ["环节", "知识单元", "知识点", "具体内容/活动", "设计意图"])
        with ZipFile(path) as archive:
            footer = archive.read("word/footer1.xml").decode("utf-8")
            document_xml = archive.read("word/document.xml").decode("utf-8")
            core_xml = archive.read("docProps/core.xml").decode("utf-8")
        self.assertIn("PAGE", footer)
        self.assertIn("w:tblHeader", document_xml)
        self.assertIn('w:tblLayout w:type="fixed"', document_xml)
        self.assertIn("WebGIS-AI", core_xml)
        self.assertNotIn("张珂", path.read_bytes().decode("latin1", errors="ignore"))

    def test_full_requirement_paragraph_prefills_multiple_sections(self) -> None:
        """验收：一段完整需求一次生成多环节初稿，且全部保持待确认。"""
        service = self.runtime.classroom.lesson_design
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        message = (
            "设计一节胡焕庸线与中国人口分布课，45分钟，重点分析东南密集西北稀疏，"
            "并安排地图观察和课堂提问。"
        )
        result = service.turn(design["design_id"], message, design["revision"])
        draft = result["draft"]
        self.assertEqual(result["generation_mode"], "draft_rules")
        self.assertEqual(draft["title"], "胡焕庸线与中国人口分布")
        self.assertEqual(draft["duration_minutes"], 45)
        self.assertEqual(draft["requirements"]["raw"], message)
        self.assertGreaterEqual(len(draft["objectives"]), 3)
        self.assertTrue(draft["core_questions"]["core"])
        self.assertGreaterEqual(len(draft["core_questions"]["sub_questions"]), 2)
        self.assertTrue(any("胡焕庸" in stage["title"] for stage in draft["stages"]))
        self.assertEqual(sum(int(stage["minutes"]) for stage in draft["stages"]), 45)
        for stage in draft["stages"]:
            self.assertTrue(stage["material"])
            self.assertTrue(stage["question_chain"])
            self.assertTrue(stage["student_activities"])
            self.assertTrue(stage["knowledge_conclusion"])
        self.assertTrue(draft["homework"]["basic"])
        self.assertTrue(draft["homework"]["inquiry"])
        self.assertTrue(result["capability_bindings"])
        self.assertNotEqual(result["section_status"].get("stages"), "confirmed")
        for key in ("requirements", "objectives", "core_questions", "stages", "homework", "capabilities"):
            self.assertEqual(result["section_status"].get(key), "proposed")
        self.assertTrue(any("待确认" in item for item in [result["assistant_message"]]))
        self.assertIn("教学建议", result["assistant_message"])
        focus = result["focus_summary"]
        self.assertIn("教学过程", focus["changed_labels"])
        self.assertTrue(focus["next_confirm_sections"])
        self.assertTrue(focus["next_confirm_question"])
        # 未审核内容不能直接定稿
        with self.assertRaisesRegex(ValueError, "请先逐项确认"):
            service.finalize(design["design_id"], result["revision"])
        report = service.rehearse(design["design_id"])
        self.assertTrue(report["ready"], report["errors"])

    def test_full_draft_does_not_overwrite_existing_or_confirmed_sections(self) -> None:
        """完整初稿只填空缺章节；已确认章节和已有内容一律不覆盖。"""
        service = self.runtime.classroom.lesson_design
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        first = service.turn(design["design_id"], "高一、40分钟、人口迁移", design["revision"])
        stored = self.store.get_lesson_design(design["design_id"])
        stored.draft["objectives"] = ["教师目标一", "教师目标二", "教师目标三"]
        stored.draft["stages"] = [{
            "stage_id": "s1", "title": "教师环节", "minutes": 40, "material": "教材图",
            "question_chain": ["教师问题"], "teacher_activities": ["教师活动"],
            "student_activities": ["学生活动"], "knowledge_conclusion": "教师结论",
            "design_intent": "教师意图",
            "questions": [{"question_id": "s1q1", "type": "open", "text": "教师问题",
                           "options": [], "answer_index": None, "expected_points": [], "misconceptions": []}],
        }]
        self.store.upsert_lesson_design(stored)
        result = service.turn(
            design["design_id"],
            "生成完整初稿：设计一节人口迁移课，重点分析推拉因素",
            self.store.get_lesson_design(design["design_id"]).revision,
        )
        draft = result["draft"]
        self.assertEqual(draft["objectives"], ["教师目标一", "教师目标二", "教师目标三"])
        self.assertEqual(draft["stages"][0]["title"], "教师环节")
        # 空缺章节仍被预填
        self.assertTrue(draft["core_questions"]["core"])
        self.assertTrue(draft["curriculum_interpretation"])
        self.assertTrue(any("未覆盖" in item for item in [result["assistant_message"]]))
        # 已确认章节保持确认状态（requirements 在首轮已确认）
        accepted = self.runtime.classroom.resolve_lesson_design(
            design["design_id"], "requirements", "accept", "",
            self.store.get_lesson_design(design["design_id"]).revision,
        )
        self.assertEqual(accepted["design"]["section_status"]["requirements"], "confirmed")

    def test_scoped_turn_only_modifies_current_step_sections(self) -> None:
        """「只修改当前环节」：补丁只落当前步骤章节，改完留在本步等待复核。"""
        service = self.runtime.classroom.lesson_design
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        first = service.turn(design["design_id"], "高一40分钟《人口分布》", design["revision"])
        stored = self.store.get_lesson_design(design["design_id"])
        stored.draft["objectives"] = ["描述分布", "解释成因", "迁移方法"]
        stored.draft["stages"] = [{
            "stage_id": "s1", "title": "导入", "minutes": 10, "material": "人口分布图",
            "question_chain": ["差异？"], "teacher_activities": ["引导"],
            "student_activities": ["读图"], "knowledge_conclusion": "分布不均",
            "design_intent": "观察",
            "questions": [{"question_id": "s1q1", "type": "open", "text": "人口分布有何差异？",
                           "options": [], "answer_index": None, "expected_points": [], "misconceptions": []}],
        }]
        stored.draft["board_design"] = "板书"
        stored.current_step = "process"
        self.store.upsert_lesson_design(stored)
        revised_stage = dict(stored.draft["stages"][0], minutes=12)
        cross_section_patch = {
            "reply": "已调整", "section_patch": {"stages": [revised_stage], "objectives": ["越权目标"]},
            "next_step": "question_matching", "source_refs": [], "capability_bindings": [], "suggestions": [],
        }
        with patch.object(service, "_ask_minimax", return_value=cross_section_patch):
            result = service.turn(
                design["design_id"], "只修改当前环节：把导入调整为12分钟",
                self.store.get_lesson_design(design["design_id"]).revision, "process",
            )
        self.assertEqual(result["draft"]["objectives"], ["描述分布", "解释成因", "迁移方法"])
        self.assertEqual(result["draft"]["stages"][0]["minutes"], 12)
        self.assertEqual(result["section_status"].get("objectives"), "pending")
        self.assertEqual(result["next_step"], "process")
        self.assertEqual(self.store.get_lesson_design(design["design_id"]).pending_next_step, "")
        self.assertEqual(result["focus_summary"]["changed_labels"], ["教学过程"])
        # 空指令必须提示教师先写修改内容
        with self.assertRaisesRegex(ValueError, "只修改当前环节"):
            service.turn(
                design["design_id"], "只修改当前环节",
                self.store.get_lesson_design(design["design_id"]).revision, "process",
            )

    def test_focus_summary_supports_continuous_adoption(self) -> None:
        """每轮突出刚改/还缺/下一步；初稿预填后可连续采用建议推进。"""
        service = self.runtime.classroom.lesson_design
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        result = service.turn(
            design["design_id"],
            "设计一节胡焕庸线与中国人口分布课，40分钟，重点分析东南密集西北稀疏。",
            design["revision"],
        )
        focus = result["focus_summary"]
        self.assertIn("教学过程", focus["changed_labels"])
        self.assertIn("教学建议", focus["unverified_note"])
        revision = result["revision"]
        confirmed_seen = False
        for step in ("requirements", "analysis", "objectives", "core_questions",
                     "process", "question_matching", "capabilities"):
            resolved = self.runtime.classroom.resolve_lesson_design(design["design_id"], step, "accept", "", revision)
            revision = resolved["design"]["revision"]
            if resolved["focus_summary"]["confirmed_labels"]:
                confirmed_seen = True
        self.assertTrue(confirmed_seen)
        final_view = self.runtime.classroom.get_lesson_design(design["design_id"])
        self.assertEqual(final_view["current_step"], "rehearsal")
        required_confirmed = all(
            final_view["section_status"].get(key) == "confirmed" for key in
            ("requirements", "curriculum_interpretation", "objectives", "core_questions", "stages", "capabilities")
        )
        self.assertTrue(required_confirmed)
        self.assertTrue(final_view["focus_summary"]["next_confirm_sections"] is not None)

    def test_smart_optimize_uses_model_respects_confirmations_and_duration(self) -> None:
        """一键智能优化：模型结果只落非确认章节，环节分钟合计被校正回课时。"""
        service = self.runtime.classroom.lesson_design
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        first = service.turn(
            design["design_id"],
            "设计一节胡焕庸线与中国人口分布课，40分钟，重点分析东南密集西北稀疏。",
            design["revision"],
        )
        accepted = self.runtime.classroom.resolve_lesson_design(
            design["design_id"], "requirements", "accept", "", first["revision"]
        )["design"]
        requirements_before = json.dumps(accepted["draft"]["requirements"], ensure_ascii=False, sort_keys=True)
        optimized_objectives = ["优化目标一", "优化目标二", "优化目标三"]
        fake = {
            "reply": "重写了目标并强化了探究环节",
            "section_patch": {
                "requirements": {"raw": "越权改写已确认需求"},
                "objectives": optimized_objectives,
                "stages": [
                    {"stage_id": "s1", "title": "读图", "minutes": 10, "questions": [],
                     "material": "图", "question_chain": ["q"], "teacher_activities": ["t"],
                     "student_activities": ["s"], "knowledge_conclusion": "c", "design_intent": "i"},
                    {"stage_id": "s2", "title": "探究", "minutes": 15, "questions": [],
                     "material": "图", "question_chain": ["q"], "teacher_activities": ["t"],
                     "student_activities": ["s"], "knowledge_conclusion": "c", "design_intent": "i"},
                ],
            },
            "next_step": "", "source_refs": [], "capability_bindings": [], "suggestions": [],
        }
        with patch.object(service, "_ask_minimax", return_value=fake) as ask:
            result = service.turn(
                design["design_id"], "一键智能优化", accepted["revision"],
                accepted["current_step"],
            )
        ask.assert_called_once()
        self.assertIs(ask.call_args.kwargs.get("optimize"), True)
        self.assertEqual(result["generation_mode"], "model")
        self.assertEqual(result["draft"]["objectives"], optimized_objectives)
        # 已确认章节不被优化触碰
        self.assertEqual(result["section_status"].get("requirements"), "confirmed")
        self.assertEqual(
            json.dumps(result["draft"]["requirements"], ensure_ascii=False, sort_keys=True),
            requirements_before,
        )
        # 环节分钟合计被校正回课时
        self.assertEqual(sum(int(stage["minutes"]) for stage in result["draft"]["stages"]), 40)
        self.assertTrue(all(int(stage["minutes"]) >= 1 for stage in result["draft"]["stages"]))
        self.assertEqual(result["section_status"].get("objectives"), "proposed")
        self.assertIn("待确认", result["assistant_message"])
        self.assertIn("objectives", result["review_sections"])
        self.assertIn("教学过程", result["focus_summary"]["changed_labels"])

    def test_smart_optimize_falls_back_to_rules_fill_without_model(self) -> None:
        """AI 不可用时，一键优化退化为规则补齐：只填空缺，不改已有内容。"""
        service = self.runtime.classroom.lesson_design
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        first = service.turn(design["design_id"], "高一40分钟《人口分布》", design["revision"])
        before_topic = first["draft"]["topic"]
        with patch.object(service, "_ask_minimax", return_value=None):
            result = service.turn(
                design["design_id"], "一键智能优化", first["revision"], first["next_step"]
            )
        self.assertEqual(result["generation_mode"], "optimize_rules")
        self.assertIn("AI 未返回有效内容", result["assistant_message"])
        self.assertTrue(result["draft"]["stages"])
        self.assertTrue(result["draft"]["curriculum_interpretation"])
        self.assertTrue(result["draft"]["core_questions"]["core"])
        self.assertEqual(result["draft"]["topic"], before_topic)
        self.assertEqual(result["section_status"].get("stages"), "proposed")
        # 需求章节已有内容，规则兜底不覆盖
        self.assertNotIn("requirements", result["review_sections"])

    def test_smart_optimize_requires_existing_draft(self) -> None:
        service = self.runtime.classroom.lesson_design
        design = self.runtime.classroom.create_lesson_design(self.project, "local_admin")
        with self.assertRaisesRegex(ValueError, "生成完整初稿"):
            service.turn(design["design_id"], "一键智能优化", design["revision"])

    def test_optimize_prompt_has_full_draft_budget_and_guardrails(self) -> None:
        service = self.runtime.classroom.lesson_design
        design = service.create_or_resume(self.project, "local_admin")
        client = Mock()
        client.chat_completion.return_value = '{"section_patch": {}}'
        service.minimax_client = client
        service._ask_minimax(design, "confirmation", "一键智能优化", optimize=True)
        system_content = client.chat_completion.call_args.args[0][0]["content"]
        self.assertIn("一键智能优化", system_content)
        self.assertIn("不得改变教学环节的名称、顺序和分钟数", system_content)
        self.assertEqual(client.chat_completion.call_args.kwargs["extra_payload"]["max_completion_tokens"], 12288)
        self.assertEqual(client.chat_completion.call_args.kwargs["timeout"], 90.0)


if __name__ == "__main__":
    unittest.main()
