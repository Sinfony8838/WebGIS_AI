from __future__ import annotations

import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image

from backend.app.config import AppConfig
from backend.app.runtime import WebGISRuntime
from backend.app.store import RuntimeStore
from backend.app.services.lesson_docx_import import parse_lesson_docx


def _sample_docx_bytes(with_image: bool = True) -> bytes:
    doc = Document()
    doc.add_heading("人口迁移教案", level=1)
    doc.add_paragraph("课题：人口的空间变化")
    doc.add_paragraph("年级：高一")
    doc.add_paragraph("课时：45分钟")
    doc.add_paragraph("一、教学目标")
    doc.add_paragraph("描述人口迁移的方向")
    doc.add_paragraph("解释影响人口迁移的因素")
    doc.add_paragraph("二、教学重难点")
    doc.add_paragraph("教学重点：人口迁移的推拉因素")
    doc.add_paragraph("教学难点：迁移对区域发展的影响")
    doc.add_paragraph("三、教学方法")
    doc.add_paragraph("情境教学；小组合作探究")
    doc.add_paragraph("四、教学过程")
    table = doc.add_table(rows=4, cols=6)
    headers = ["环节", "时长(分钟)", "教师活动", "学生活动", "材料", "知识结论"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    rows = [
        ["情境导入", "5", "播放春运视频", "观察并描述现象", "春运新闻图片", "人口迁移普遍存在"],
        ["成因探究", "25", "引导分组讨论", "归纳推拉因素", "教材图1.8", "推拉因素共同作用"],
        ["总结迁移", "15", "梳理知识结构", "完成思维导图", "", "迁移受多重因素影响"],
    ]
    for row, values in zip(table.rows[1:], rows):
        for cell, text in zip(row.cells, values):
            cell.text = text
    doc.add_paragraph("五、板书设计")
    doc.add_paragraph("人口迁移：概念 → 因素 → 影响")
    doc.add_paragraph("古蜀文明起源于四川盆地的考古新发现，涉及三重证据链。")
    if with_image:
        image = Image.new("RGB", (3, 3), (200, 30, 30))
        buffer = BytesIO()
        image.save(buffer, format="PNG")
        buffer.seek(0)
        doc.add_picture(buffer, width=Inches(1))
        buffer.close()
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


class LessonDocxImportTest(unittest.TestCase):
    """Word 教案导入：段落/表格/图片解析 → 新建未确认教案草稿。"""

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        root_dir = Path(__file__).resolve().parents[2]
        config = AppConfig(root_dir=root_dir)
        config.data_dir = Path(self.temp_dir.name) / "backend" / "data"
        config.state_dir = config.data_dir / "state"
        config.uploads_dir = config.data_dir / "uploads"
        config.outputs_dir = config.data_dir / "outputs"
        config.state_file = config.state_dir / "runtime.json"
        config.minimax_api_key = ""
        config.minimax_token_plan_key = ""
        config.ensure_dirs()
        self.store = RuntimeStore(config.state_file)
        self.runtime = WebGISRuntime(config=config, store=self.store)
        self.project = self.runtime.create_project()["project_id"]
        self.service = self.runtime.classroom.lesson_design
        self.config = config
        self.addCleanup(self.temp_dir.cleanup)

    def test_parse_maps_paragraphs_tables_and_images(self) -> None:
        result = parse_lesson_docx(_sample_docx_bytes())
        draft = result["draft"]
        self.assertEqual(draft["title"], "人口的空间变化")
        self.assertEqual(draft["grade"], "高一")
        self.assertEqual(draft["duration_minutes"], 45)
        self.assertEqual(len(draft["objectives"]), 2)
        self.assertEqual(draft["key_difficulties"]["key"], ["人口迁移的推拉因素"])
        self.assertEqual(draft["key_difficulties"]["difficult"], ["迁移对区域发展的影响"])
        self.assertEqual(len(draft["methods"]), 2)
        self.assertEqual(draft["board_design"], "人口迁移：概念 → 因素 → 影响")
        self.assertEqual(len(draft["stages"]), 3)
        first = draft["stages"][0]
        self.assertEqual(first["title"], "情境导入")
        self.assertEqual(first["minutes"], 5)
        self.assertEqual(first["teacher_activities"], ["播放春运视频"])
        self.assertEqual(first["student_activities"], ["观察并描述现象"])
        self.assertEqual(first["material"], "春运新闻图片")
        self.assertEqual(first["knowledge_conclusion"], "人口迁移普遍存在")
        self.assertTrue(all(stage.get("stage_id") for stage in draft["stages"]))
        # 无法识别的段落保留
        unclassified_texts = [item["text"] for item in result["unclassified"] if item["kind"] == "text"]
        self.assertTrue(any("古蜀文明" in text for text in unclassified_texts))
        # 图片提取为待归类内容
        images = [item for item in result["unclassified"] if item["kind"] == "image"]
        self.assertEqual(len(images), 1)
        self.assertEqual(images[0]["content_type"], "image/png")
        # 映射记录可追溯
        fields = {item["field"] for item in result["mapping"]}
        self.assertIn("title", fields)
        self.assertIn("stages", fields)
        origins = [item["origin"] for item in result["mapping"] if item["field"] == "stages"]
        self.assertTrue(any("table" in origin for origin in origins))

    def test_parse_rejects_corrupted_file(self) -> None:
        with self.assertRaisesRegex(ValueError, "无法解析"):
            parse_lesson_docx(b"not a docx file")

    def test_import_creates_independent_unconfirmed_design(self) -> None:
        existing = self.service.create_or_resume(self.project, "local_admin")
        result = self.runtime.classroom.import_lesson_design_docx(
            self.project, "local_admin", _sample_docx_bytes(), "人口迁移教案.docx"
        )
        design = self.service.get(result["design"]["design_id"])
        self.assertNotEqual(design.design_id, existing.design_id)
        self.assertEqual(design.status, "active")
        self.assertEqual(design.draft["title"], "人口的空间变化")
        self.assertEqual(design.section_status["objectives"], "proposed")
        self.assertNotEqual(design.section_status["stages"], "confirmed")
        # 图片已存盘并可公开访问
        image_items = [item for item in result["unclassified"] if item["kind"] == "image"]
        self.assertTrue(image_items[0]["url"].startswith("/files/uploads/"))
        self.assertTrue((self.config.uploads_dir).exists())
        self.assertIn("待归类", result["summary"])

    def test_import_without_any_image_still_succeeds(self) -> None:
        result = self.runtime.classroom.import_lesson_design_docx(
            self.project, "local_admin", _sample_docx_bytes(with_image=False), "教案.docx"
        )
        self.assertFalse([item for item in result["unclassified"] if item["kind"] == "image"])


if __name__ == "__main__":
    unittest.main()
