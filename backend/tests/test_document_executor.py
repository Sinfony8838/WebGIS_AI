from __future__ import annotations

import asyncio
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
import os
import random
import re
from pathlib import Path
import sys
import threading
import time

import pytest

from backend.app.services.archive_safety import ArchiveLimitError, ArchiveValidationError
from backend.app.services.document_executor import (
    DocumentBusyError, DocumentCancelledError, DocumentExecutor, DocumentMemoryError,
    DocumentResultLimitError, DocumentTimeoutError, DocumentWorkerError,
)
from backend.app.services.question_bank_parser import DocxBudgetError, DocxValidationError, ParsedBank

FIXTURE = Path(__file__).parent / "fixtures" / "document_process_fixture.py"


def fixture_executor(tmp_path, **kwargs):
    return DocumentExecutor(temp_dir=tmp_path, _worker_command=[sys.executable, str(FIXTURE)], **kwargs)


def payload(**kwargs):
    return json.dumps({key: str(value) if isinstance(value, Path) else value for key, value in kwargs.items()}).encode()


def eventually(predicate, timeout=10):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if predicate():
            return
        time.sleep(0.01)
    assert predicate(), "condition was not reached"


def alive(pid):
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
        kernel.OpenProcess.restype = wintypes.HANDLE
        kernel.WaitForSingleObject.argtypes = [wintypes.HANDLE, wintypes.DWORD]
        kernel.WaitForSingleObject.restype = wintypes.DWORD
        kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        handle = kernel.OpenProcess(0x100000, False, pid)
        if not handle:
            return False
        try:
            return kernel.WaitForSingleObject(handle, 0) == 0x102
        finally:
            kernel.CloseHandle(handle)
    try:
        os.kill(pid, 0)
        status = Path(f"/proc/{pid}/status")
        return not status.exists() or "State:\tZ" not in status.read_text()
    except ProcessLookupError:
        return False


def assert_clean(executor, tmp_path):
    assert executor.status()["running"] == 0
    assert executor.status()["waiting"] == 0
    assert not list(tmp_path.glob("webgis-document-*"))


def make_docx():
    from docx import Document
    document = Document()
    document.add_paragraph("课题：人口分布")
    document.add_paragraph("考点01 人口")
    document.add_paragraph("（2025·湖南·高考真题）下图为人口分布图。")
    from PIL import Image
    picture = BytesIO()
    Image.frombytes("RGB", (64, 64), random.Random(7).randbytes(64 * 64 * 3)).save(picture, format="PNG")
    document.add_picture(BytesIO(picture.getvalue()))
    document.add_paragraph("1．人口主要分布于（  ）")
    document.add_paragraph("A．平原\tB．高原\tC．荒漠\tD．冰原")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "教学环节"
    table.cell(0, 1).text = "教师活动"
    table.cell(1, 0).text = "导入"
    table.cell(1, 1).text = "观察地图"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_default_budgets_and_unknown_operation(tmp_path):
    executor = DocumentExecutor(temp_dir=tmp_path)
    assert executor.max_workers == 2
    assert executor.max_waiting == 8
    assert executor.timeout_seconds == 120
    assert executor.memory_limit_bytes == 1024 ** 3
    with pytest.raises(ValueError, match="Unsupported"):
        executor.run("eval", b"evil")
    assert_clean(executor, tmp_path)


def test_real_operations_keep_existing_shapes_content_and_image_bytes(tmp_path):
    executor = DocumentExecutor(temp_dir=tmp_path)
    raw = make_docx()
    assert executor.run("office_validate", raw, kind="docx") is None
    lesson = executor.run("lesson_docx", raw)
    assert lesson["draft"]["title"] == "人口分布"
    from backend.app.services.lesson_docx_import import parse_lesson_docx
    from backend.app.services.question_bank_parser import QuestionBankDocxParser
    expected_lesson = parse_lesson_docx(raw)
    for actual, expected in zip(lesson["draft"]["stages"], expected_lesson["draft"]["stages"], strict=True):
        assert re.fullmatch(r"import_[0-9a-f]{6}_\d+", actual.pop("stage_id"))
        expected.pop("stage_id")
    assert lesson == expected_lesson
    assert any(isinstance(item.get("data"), bytes) for item in lesson["unclassified"])
    assert lesson["draft"]["stages"][0]["title"] == "导入"
    bank = executor.run("question_bank_docx", raw, filename="人口原卷版.docx")
    assert isinstance(bank, ParsedBank)
    assert bank == QuestionBankDocxParser("人口原卷版.docx", raw).parse()
    assert bank.groups[0].images[0].blob
    assert bank.groups[0].questions[0].stem.startswith("人口主要")
    assert "观察地图" in executor.run("timeline_text", raw, filename="课程.docx")
    assert executor.run("timeline_text", "你好".encode(), filename="课程.md") == "你好"
    assert_clean(executor, tmp_path)


def test_validation_errors_and_ipc_limits(tmp_path):
    executor = DocumentExecutor(temp_dir=tmp_path, max_input_bytes=100)
    with pytest.raises(ArchiveValidationError):
        executor.run("office_validate", b"bad", kind="docx")
    with pytest.raises(DocxValidationError):
        executor.run("question_bank_docx", b"bad", filename="bad.docx")
    with pytest.raises(DocumentResultLimitError):
        executor.run("timeline_text", b"a" * 101, filename="input.txt")
    fixture = fixture_executor(tmp_path, max_result_bytes=1024)
    with pytest.raises(DocumentResultLimitError):
        fixture.run("lesson_docx", payload(mode="large"))
    assert_clean(executor, tmp_path)
    assert_clean(fixture, tmp_path)


def test_child_env_and_no_runtime_or_model_import(tmp_path, monkeypatch):
    monkeypatch.setenv("WEBGIS_AI_MINIMAX_API_KEY", "must-not-enter-worker")
    monkeypatch.setenv("WEBGIS_AI_DATA_DIR", "must-not-enter-worker")
    executor = fixture_executor(tmp_path)
    result = executor.run("lesson_docx", payload())
    assert result["pid"] != os.getpid()
    assert result["memory_cap"] == 1024 ** 3
    assert result["secret"] is None
    assert result["loaded_forbidden"] == []
    assert not alive(result["pid"])
    assert_clean(executor, tmp_path)


@pytest.mark.parametrize("mode", ["sleep", "crash", "memory"])
def test_real_timeout_crash_memory_then_capacity_reuse(tmp_path, mode):
    executor = fixture_executor(tmp_path, timeout_seconds=2, memory_limit_bytes=128 * 1024 ** 2)
    pid_path = tmp_path / "pid"
    child_path = tmp_path / "child"
    expected = {"sleep": DocumentTimeoutError, "crash": DocumentWorkerError, "memory": DocumentMemoryError}[mode]
    with pytest.raises(expected):
        executor.run("lesson_docx", payload(mode=mode, seconds=10, pid_file=pid_path, child_pid_file=child_path))
    assert pid_path.exists() and child_path.exists()
    assert not alive(int(pid_path.read_text()))
    assert not alive(int(child_path.read_text()))
    assert executor.run("lesson_docx", payload(value="recovered"))["value"] == "recovered"
    assert_clean(executor, tmp_path)


def test_repeated_crashes_reap_descendants_before_capacity_is_reused(tmp_path):
    # Exercise the short parent-exit / descendant-termination race repeatedly.
    executor = fixture_executor(tmp_path, timeout_seconds=5)
    for index in range(8):
        child = tmp_path / f"child-{index}"
        with pytest.raises(DocumentWorkerError):
            executor.run("lesson_docx", payload(mode="crash", child_pid_file=child))
        assert not alive(int(child.read_text()))
        assert_clean(executor, tmp_path)
    assert executor.run("lesson_docx", payload(value="recovered"))["value"] == "recovered"


def test_exact_two_running_eight_waiting_busy_and_queued_cancel(tmp_path):
    executor = fixture_executor(tmp_path, timeout_seconds=20)
    gate = tmp_path / "go"
    cancel = threading.Event()
    with ThreadPoolExecutor(max_workers=12) as pool:
        running = [pool.submit(executor.run, "lesson_docx", payload(mode="gate", gate_file=gate,
                    pid_file=tmp_path / f"pid{i}")) for i in range(2)]
        eventually(lambda: all((tmp_path / f"pid{i}").exists() for i in range(2)))
        queued = [pool.submit(executor.run, "lesson_docx", payload(value=i),
                             cancel_event=cancel if i == 0 else None) for i in range(8)]
        eventually(lambda: executor.status()["waiting"] == 8)
        assert executor.status()["running"] == 2
        with pytest.raises(DocumentBusyError):
            executor.run("lesson_docx", payload())
        cancel.set()
        with pytest.raises(DocumentCancelledError):
            queued[0].result(timeout=5)
        assert executor.status()["waiting"] == 7
        replacement = pool.submit(executor.run, "lesson_docx", payload(value="replacement"))
        eventually(lambda: executor.status()["waiting"] == 8)
        gate.touch()
        for future in running + queued[1:] + [replacement]:
            future.result(timeout=10)
    assert_clean(executor, tmp_path)


def test_running_cancel_reaps_descendant_before_return(tmp_path):
    executor = fixture_executor(tmp_path, timeout_seconds=20)
    cancel = threading.Event()
    pid = tmp_path / "pid"
    child = tmp_path / "child"
    with ThreadPoolExecutor(max_workers=1) as pool:
        future = pool.submit(executor.run, "lesson_docx", payload(mode="sleep", seconds=10,
                             pid_file=pid, child_pid_file=child), cancel_event=cancel)
        eventually(child.exists)
        cancel.set()
        with pytest.raises(DocumentCancelledError):
            future.result(timeout=5)
    assert not alive(int(pid.read_text()))
    assert not alive(int(child.read_text()))
    assert_clean(executor, tmp_path)


def test_async_cancel_and_async_admission_are_bounded(tmp_path):
    executor = fixture_executor(tmp_path, max_workers=1, max_waiting=1, timeout_seconds=20)
    pid = tmp_path / "pid"
    child = tmp_path / "child"

    async def exercise():
        task = asyncio.create_task(executor.arun("lesson_docx", payload(mode="sleep", seconds=10,
                                   pid_file=pid, child_pid_file=child)))
        while not child.exists():
            await asyncio.sleep(0.01)
        queued = asyncio.create_task(executor.arun("lesson_docx", payload()))
        await asyncio.sleep(0)
        with pytest.raises(DocumentBusyError):
            await executor.arun("lesson_docx", payload())
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert not alive(int(pid.read_text()))
        assert not alive(int(child.read_text()))
        await queued
    asyncio.run(exercise())
    assert_clean(executor, tmp_path)


def test_shutdown_cancels_running_and_waiting(tmp_path):
    executor = fixture_executor(tmp_path, max_workers=1, max_waiting=1, timeout_seconds=20)
    pid = tmp_path / "pid"
    with ThreadPoolExecutor(max_workers=2) as pool:
        running = pool.submit(executor.run, "lesson_docx", payload(mode="sleep", seconds=10, pid_file=pid))
        eventually(pid.exists)
        queued = pool.submit(executor.run, "lesson_docx", payload())
        eventually(lambda: executor.status()["waiting"] == 1)
        executor.close()
        for future in (running, queued):
            with pytest.raises(DocumentCancelledError):
                future.result(timeout=5)
    assert not alive(int(pid.read_text()))
    with pytest.raises(DocumentCancelledError):
        executor.run("lesson_docx", payload())
    assert_clean(executor, tmp_path)


def test_validation_categories_remain_distinct(tmp_path):
    executor = fixture_executor(tmp_path)
    for mode, expected in [("archive_limit", ArchiveLimitError), ("docx_budget", DocxBudgetError)]:
        with pytest.raises(expected, match="fixture budget"):
            executor.run("lesson_docx", payload(mode=mode))
    assert_clean(executor, tmp_path)


@pytest.mark.parametrize("mode,expected", [
    ("oversized_result", DocumentResultLimitError), ("invalid_result", DocumentWorkerError),
    ("invalid_schema", DocumentWorkerError), ("unserializable", DocumentWorkerError),
])
def test_child_result_corruption_is_rejected_and_releases_capacity(tmp_path, mode, expected):
    executor = fixture_executor(tmp_path, max_result_bytes=1024)
    with pytest.raises(expected):
        executor.run("lesson_docx", payload(mode=mode))
    assert executor.run("lesson_docx", payload(value="ok"))["value"] == "ok"
    assert_clean(executor, tmp_path)


def test_result_written_before_timeout_is_never_accepted(tmp_path):
    executor = fixture_executor(tmp_path, timeout_seconds=1)
    pid = tmp_path / "pid"
    with pytest.raises(DocumentTimeoutError):
        executor.run("lesson_docx", payload(mode="late_result", seconds=5, pid_file=pid))
    assert not alive(int(pid.read_text()))
    assert_clean(executor, tmp_path)


def test_queue_time_counts_against_one_end_to_end_deadline(tmp_path):
    executor = fixture_executor(tmp_path, max_workers=1, max_waiting=1, timeout_seconds=2)
    gate = tmp_path / "go"
    first_pid = tmp_path / "first"
    second_pid = tmp_path / "second"
    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(executor.run, "lesson_docx", payload(mode="gate", gate_file=gate, pid_file=first_pid))
        eventually(first_pid.exists)
        started = time.monotonic()
        second = pool.submit(executor.run, "lesson_docx", payload(mode="sleep", seconds=1.6, pid_file=second_pid))
        eventually(lambda: executor.status()["waiting"] == 1)
        time.sleep(0.8)
        gate.touch()
        first.result(timeout=3)
        with pytest.raises(DocumentTimeoutError):
            second.result(timeout=3)
        assert time.monotonic() - started < 2.8
    assert not alive(int(second_pid.read_text()))
    assert_clean(executor, tmp_path)


def test_fifo_order_is_reserved_at_admission(tmp_path):
    executor = fixture_executor(tmp_path, max_workers=1, max_waiting=3, timeout_seconds=10)
    gate = tmp_path / "go"
    pid = tmp_path / "pid"
    order = tmp_path / "order"
    with ThreadPoolExecutor(max_workers=4) as pool:
        first = pool.submit(executor.run, "lesson_docx", payload(mode="gate", gate_file=gate, pid_file=pid))
        eventually(pid.exists)
        queued = []
        for value in range(3):
            queued.append(pool.submit(executor.run, "lesson_docx", payload(value=value, order_file=order)))
            eventually(lambda: executor.status()["waiting"] == value + 1)
        gate.touch()
        first.result(timeout=5)
        for future in queued:
            future.result(timeout=5)
    assert order.read_text().splitlines() == ["0", "1", "2"]
    assert_clean(executor, tmp_path)


def test_async_deadline_does_not_depend_on_busy_default_thread_pool(tmp_path):
    executor = fixture_executor(tmp_path, timeout_seconds=0.5)
    default_gate = threading.Event()

    async def exercise():
        loop = asyncio.get_running_loop()
        loop.set_default_executor(ThreadPoolExecutor(max_workers=1))
        occupied = loop.run_in_executor(None, default_gate.wait)
        try:
            with pytest.raises(DocumentTimeoutError):
                await asyncio.wait_for(executor.arun("lesson_docx", payload(mode="sleep", seconds=5)), timeout=2)
        finally:
            default_gate.set()
            await occupied
    try:
        asyncio.run(exercise())
    finally:
        executor.close()
    assert_clean(executor, tmp_path)


def test_success_also_kills_surviving_descendants(tmp_path):
    executor = fixture_executor(tmp_path)
    child = tmp_path / "child"
    result = executor.run("lesson_docx", payload(child_pid_file=child))
    assert not alive(result["pid"])
    assert not alive(int(child.read_text()))
    assert_clean(executor, tmp_path)


@pytest.mark.parametrize("action", ["timeout", "cancel", "failure"])
def test_cleanup_is_inside_deadline_and_cleanup_failure_disables_executor(tmp_path, monkeypatch, action):
    import backend.app.services.document_executor as module

    cancel = threading.Event()
    original = module.tempfile.TemporaryDirectory
    cleanup_called = threading.Event()

    class CleanupFixture(original):
        def cleanup(self):
            super().cleanup()
            cleanup_called.set()
            if action == "timeout":
                time.sleep(0.6)
            elif action == "cancel":
                cancel.set()
            else:
                raise PermissionError("simulated locked IPC directory")

    monkeypatch.setattr(module.tempfile, "TemporaryDirectory", CleanupFixture)
    executor = fixture_executor(tmp_path, timeout_seconds=0.5 if action == "timeout" else 10)
    expected = {"timeout": DocumentTimeoutError, "cancel": DocumentCancelledError,
                "failure": DocumentWorkerError}[action]
    with pytest.raises(expected):
        executor.run("lesson_docx", payload(), cancel_event=cancel)
    assert cleanup_called.is_set()
    if action == "failure":
        assert executor.status()["broken"] is True
        with pytest.raises(DocumentWorkerError, match="unavailable"):
            executor.run("lesson_docx", payload())
    assert_clean(executor, tmp_path)


@pytest.mark.skipif(not sys.platform.startswith("linux"), reason="requires Linux subreaper /proc")
@pytest.mark.parametrize("mode", ["echo", "sleep", "crash"])
def test_linux_descendants_are_reaped_not_zombies_before_capacity_reuse(tmp_path, mode):
    executor = fixture_executor(tmp_path, timeout_seconds=1)
    pid = tmp_path / "pid"
    child = tmp_path / "child"
    arguments = payload(mode=mode, seconds=5, pid_file=pid, child_pid_file=child)
    if mode == "echo":
        executor.run("lesson_docx", arguments)
    else:
        with pytest.raises(DocumentTimeoutError if mode == "sleep" else DocumentWorkerError):
            executor.run("lesson_docx", arguments)
    assert not Path(f"/proc/{pid.read_text()}").exists()
    assert not Path(f"/proc/{child.read_text()}").exists()
    executor.run("lesson_docx", payload())
    assert_clean(executor, tmp_path)
