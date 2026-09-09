from __future__ import annotations

import base64
import tempfile
import unittest
from unittest.mock import Mock
from datetime import datetime, timedelta, timezone
from pathlib import Path

from backend.app.config import AppConfig
from backend.app.runtime import WebGISRuntime
from backend.app.store import RuntimeStore


# 真实题库（私有资产，不入库）：目录不存在时跳过，而不是失败。
REAL_BANK_DIR = (
    Path(r"C:\Users\zcyxn\Desktop\WebGIS-AI") / "人口专题题库"
    / "专题08 人口（全国通用）-【好题汇编】五年（2016-2025）高考地理真题分类汇编"
)
REAL_BANK_ORIGINAL = REAL_BANK_DIR / "专题08 人口（全国通用）（原卷版）.docx"
REAL_BANK_ANALYSIS = REAL_BANK_DIR / "专题08 人口（全国通用）（解析版）.docx"

# 1x1 PNG，用于真实上传题图（练习卷嵌图路径）。
TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
    "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)

STUDENT_ONLY_MARKERS = ("参考答案", "参考解析", "教师参考", "课堂实测", "课堂速记：", "考点：", "正确选项")


def docx_text(path: str | Path) -> str:
    from docx import Document

    return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)


class PracticeExportTestBase(unittest.TestCase):
    def build_runtime(self) -> tuple[WebGISRuntime, RuntimeStore, str]:
        temp_dir = tempfile.TemporaryDirectory()
        root_dir = Path(__file__).resolve().parents[2]
        config = AppConfig(root_dir=root_dir)
        config.data_dir = Path(temp_dir.name) / "backend" / "data"
        config.state_dir = config.data_dir / "state"
        config.uploads_dir = config.data_dir / "uploads"
        config.outputs_dir = config.data_dir / "outputs"
        config.state_file = config.state_dir / "runtime.json"
        # 测试强制无 LLM key：导出与诊断都走规则路径，不发起真实请求。
        config.minimax_api_key = ""
        config.minimax_token_plan_key = ""
        config.ensure_dirs()
        store = RuntimeStore(config.state_file)
        runtime = WebGISRuntime(config=config, store=store)
        project_id = runtime.create_project()["project_id"]
        self.addCleanup(temp_dir.cleanup)
        return runtime, store, project_id

    def backdate_running_since(self, session_id: str, seconds: int) -> None:
        """把 running 段起点改到过去，模拟真实计时秒数（避免等待）。"""
        session = self.store.get_class_session(session_id)
        active = dict(session.active_question)
        active["timer"]["running_since"] = (
            datetime.now(timezone.utc) - timedelta(seconds=seconds)
        ).isoformat()
        self.store.set_active_question(session_id, active)


class PracticeExportLessonGoalsTest(PracticeExportTestBase):
    def test_export_selection_rejects_stale_empty_foreign_and_unlisted_items(self):
        runtime, store, project_id = self.build_runtime()
        lesson = store.get_lesson("lesson_builtin_population_distribution")
        lesson.plan = {"homework": {"basic": ["比较人口密度与人口总量。", "分析地形与人口分布。"]}}
        session = store.get_class_session(runtime.classroom.create_class_session(lesson.lesson_id, project_id)["session"]["session_id"])
        service = runtime.classroom.practice_export
        items, _, _ = service.collect_items(session, lesson)
        manifest = service.selection_manifest(session, items)
        selected = [manifest["item_ids"][0]]
        for selection in ({"token": manifest["token"], "selected_ids": []},
                          {"token": "stale", "selected_ids": selected},
                          {"token": manifest["token"], "selected_ids": ["foreign"]},
                          {"token": manifest["token"], "selected_ids": selected * 2}):
            with self.assertRaises(ValueError):
                service.export(session, lesson, selection)
        other = store.get_class_session(runtime.classroom.create_class_session(lesson.lesson_id, project_id)["session"]["session_id"])
        with self.assertRaisesRegex(ValueError, "已变化"):
            service.export(other, lesson, {"token": manifest["token"], "selected_ids": selected})
        result = service.export(session, lesson, {"token": manifest["token"], "selected_ids": selected})
        self.assertEqual(result["selected_ids"], selected)
        recovered = runtime.classroom.session_review_history(session.session_id)["practice"]["result"]
        self.assertEqual(recovered, result)
        self.assertEqual(recovered["selection_token"], manifest["token"])
        self.assertEqual(RuntimeStore(store.state_file).get_job(result["job_id"]).result, result)
        self.assertEqual(sum(v["count"] for v in result["selection_summary"]), 1)
        student = docx_text(result["student_artifact"]["path"])
        teacher = docx_text(result["teacher_artifact"]["path"])
        self.assertIn(items[0]["text"], student)
        self.assertNotIn(items[1]["text"], student)
        self.assertNotIn(items[1]["text"], teacher)
        lesson.plan["homework"]["basic"][0] += "（修订）"
        with self.assertRaisesRegex(ValueError, "已变化"):
            service.export(session, lesson, {"token": manifest["token"], "selected_ids": selected})


    def test_report_bank_preview_matches_paper_and_keeps_material(self):
        runtime, store, project_id = self.build_runtime()
        lesson = store.get_lesson("lesson_builtin_population_distribution")
        session = store.get_class_session(runtime.classroom.create_class_session(lesson.lesson_id, project_id)["session"]["session_id"])
        service = runtime.classroom.practice_export
        bank = Mock()
        bank.search.return_value = {"items": [{
            "question_id": "bank-q", "text": "人口分布差异的原因？", "material": "某地区人口材料",
            "options": ["甲", "乙"], "answer": "乙", "answer_complete": True,
            "year": "2025", "region": "河南", "bank_id": "bank-a", "explanation": "参考解析",
        }]}
        service.question_bank = bank
        recommendations, notes, manifest = service.report_bank_recommendations(session, lesson)
        items, _, _ = service.collect_items(session, lesson)
        self.assertEqual([r["question"]["question_id"] for r in recommendations],
                         [i["question"]["question_id"] for i in items if i["origin"] == "bank_core"])
        self.assertEqual(recommendations[0]["question"]["material"], "某地区人口材料")
        self.assertEqual(recommendations[0]["question"]["options"], ["甲", "乙"])
        self.assertIn("不代表学生答错", recommendations[0]["evidence_basis"])
        self.assertTrue(all(call.kwargs["project_id"] == project_id and call.kwargs["use_llm"] is False for call in bank.search.call_args_list))
        markdown = runtime.classroom.report_service.render_markdown({}, {"text": "", "generator": "rules"}, recommendations)
        self.assertIn("某地区人口材料", markdown)
        self.assertIn("B. 乙", markdown)


    def test_population_auto_selection_requires_relevant_complete_material(self):
        runtime, store, project_id = self.build_runtime()
        lesson = store.get_lesson("lesson_builtin_population_distribution")
        session = store.get_class_session(runtime.classroom.create_class_session(lesson.lesson_id, project_id)["session"]["session_id"])
        service = runtime.classroom.practice_export
        bank = Mock()
        def question(qid, text, **extra):
            return {"question_id": qid, "text": text, "answer": "示例答案", "answer_complete": True, **extra}
        bank.search.return_value = {"items": [
            question("transport", "铁路建设的主要目的是？"),
            question("no-image", "图中人口密度如何分布？"),
            question("no-answer", "人口密度如何计算？", answer_complete=False),
            question("unrelated", "人口分布", auto_selectable=False),
            question("good", "人口密度如何计算？"),
        ]}
        service.question_bank = bank
        items, _, notes = service.collect_items(session, lesson)
        self.assertEqual([item["question"]["question_id"] for item in items if item["kind"] == "question"], ["good"])
        self.assertIn("读图题缺少题图", "；".join(notes))
        self.assertIn("题干未直接考查人口分布", "；".join(notes))

    def test_builtin_without_plan_searches_goals_and_does_not_invent_misconceptions(self):
        runtime, store, project_id = self.build_runtime()
        lesson = store.get_lesson("lesson_builtin_population_distribution")
        lesson.plan = {}
        session = store.get_class_session(runtime.classroom.create_class_session(lesson.lesson_id, project_id)["session"]["session_id"])
        service = runtime.classroom.practice_export
        bank = Mock()
        bank.search.return_value = {"items": [{"question_id": "bank-q", "text": "影响人口分布的因素", "type": "open", "answer": "自然与人文因素", "answer_complete": True}]}
        service.question_bank = bank
        items, summary, notes = service.collect_items(session, lesson)
        self.assertEqual(bank.search.call_count, 1)
        self.assertEqual(bank.search.call_args.kwargs["project_id"], project_id)
        self.assertEqual(bank.search.call_args.kwargs["topic"], "人口分布")
        self.assertEqual(bank.search.call_args.kwargs["knowledge"], "人口分布")
        self.assertEqual(bank.search.call_args.kwargs["objectives"], lesson.objectives)
        self.assertEqual([item["origin"] for item in items], ["bank_core"])

    def test_empty_export_fails_before_writing_files_and_search_failure_is_distinct(self):
        runtime, store, project_id = self.build_runtime()
        lesson = runtime.classroom.lesson_service.create_lesson({"title": "无题库课程", "stages": []})
        session = store.get_class_session(runtime.classroom.create_class_session(lesson.lesson_id, project_id)["session"]["session_id"])
        service = runtime.classroom.practice_export
        bank = Mock()
        bank.search.return_value = {"items": []}
        service.question_bank = bank
        with self.assertRaisesRegex(ValueError, "未找到可用作业"):
            service.export(session, lesson)
        bank.search.side_effect = RuntimeError("offline")
        with self.assertRaisesRegex(ValueError, "题库检索失败"):
            service.export(session, lesson)


class PracticeExportDualPaperTest(PracticeExportTestBase):
    """合成题对 + 课堂事件：双卷内容门控（学生卷无答案、教师卷有）与证据纪律。"""

    def setUp(self) -> None:
        self.runtime, self.store, self.project_id = self.build_runtime()
        self.lesson = self.runtime.classroom.lesson_service.create_lesson(
            {
                "title": "人口分布课后练习课",
                "subject": "地理",
                "grade": "高一",
                "objectives": ["能描述人口分布格局"],
                "stages": [
                    {
                        "stage_id": "s1",
                        "title": "读图探究",
                        "minutes": 10,
                        "questions": [
                            {
                                "question_id": "q1",
                                "type": "choice",
                                "text": "我国人口分布的特点是？",
                                "options": ["均匀分布", "东多西少"],
                                "answer": "B",
                                "answer_index": 1,
                                "explanation": "胡焕庸线两侧人口密度差异显著。",
                                "knowledge_points": ["人口分布格局"],
                                "answer_complete": True,
                            },
                            {
                                "question_id": "q2",
                                "type": "open",
                                "text": "说说你的家乡人口密度大概多少。",
                                "options": [],
                                "answer_index": None,
                                "images": [
                                    {"url": "/files/uploads/missing_bank/missing.png", "width": 10, "height": 10, "content_type": "image/png"}
                                ],
                            },
                            {
                                "question_id": "q3",
                                "type": "choice",
                                "text": "影响人口分布的自然因素有哪些？",
                                "options": ["气候与地形", "户籍政策"],
                                "answer": "气候与地形",
                                "answer_index": 0,
                                "explanation": "自然因素提供基础条件。",
                                "knowledge_points": ["影响因素"],
                                "answer_complete": True,
                            },
                            {
                                "question_id": "q4",
                                "type": "choice",
                                "text": "全班都答对的课堂检测题是哪道？",
                                "options": ["甲", "乙"],
                                "answer": "甲",
                                "answer_index": 0,
                                "explanation": "课堂检测说明已掌握。",
                                "answer_complete": True,
                            },
                        ],
                    }
                ],
                "plan": {
                    "title": "人口分布课后练习课",
                    "topic": "人口分布",
                    "duration_minutes": 40,
                    "homework": {
                        "basic": ["完成课本第32页活动题。"],
                        "inquiry": ["查找胡焕庸线资料写一段短评。"],
                    },
                    "objectives": ["能描述人口分布格局"],
                    "core_questions": {"core": "人口分布格局如何形成？", "sub_questions": []},
                },
            },
            source="manual",
        )
        self.session_id = self.runtime.classroom.create_class_session(
            self.lesson.lesson_id, self.project_id
        )["session"]["session_id"]
        cw = self.runtime.classroom
        # q1：投屏、计时 150 秒、揭示、2 人作答（1 对 1 错）
        cw.launch_session_question(self.session_id, stage_id="s1", question_id="q1")
        with self.store.batch():
            self.store.add_student_response(self.session_id, "q1", {"nickname": "小李", "choice_index": 1})
            self.store.add_student_response(self.session_id, "q1", {"nickname": "小王", "choice_index": 0})
        cw.update_question_timer(self.session_id, "start")
        self.backdate_running_since(self.session_id, 150)
        cw.reveal_session_question(self.session_id)
        # q3：投屏、计时 40 秒、未揭示收口、无人作答
        cw.launch_session_question(self.session_id, stage_id="s1", question_id="q3")
        cw.update_question_timer(self.session_id, "start")
        self.backdate_running_since(self.session_id, 40)
        cw.close_session_question(self.session_id)
        # q4：投屏、2 人全对，教师未做观察（按选题规则不进练习卷）
        cw.launch_session_question(self.session_id, stage_id="s1", question_id="q4")
        with self.store.batch():
            self.store.add_student_response(self.session_id, "q4", {"nickname": "小李", "choice_index": 0})
            self.store.add_student_response(self.session_id, "q4", {"nickname": "小王", "choice_index": 0})
        cw.close_session_question(self.session_id)
        # 课堂速记：q1 误区、q2 部分掌握、q3 误区
        cw.add_session_observation(
            self.session_id,
            {"stage_id": "s1", "question_id": "q1", "verdict": "misconception", "tag": "混淆数量与密度", "note": "有学生认为总量大密度就大"},
        )
        cw.add_session_observation(
            self.session_id,
            {"stage_id": "s1", "question_id": "q2", "verdict": "partial", "tag": "密度概念不清", "note": "把总量当密度"},
        )
        cw.add_session_observation(
            self.session_id,
            {"stage_id": "s1", "question_id": "q3", "verdict": "misconception", "tag": "读图不细", "note": ""},
        )
        cw.end_class_session(self.session_id)

    def export(self) -> dict:
        return self.runtime.classroom.export_session_practice(self.session_id)

    def test_unknown_session_raises_keyerror(self) -> None:
        with self.assertRaises(KeyError):
            self.runtime.classroom.export_session_practice("session_missing")

    def test_dual_paper_selection_sources_and_counts(self) -> None:
        export = self.export()
        self.assertEqual(export["status"], "success")
        summary = {entry["origin"]: entry["count"] for entry in export["selection_summary"]}
        self.assertEqual(summary["lesson_homework_basic"], 1)
        self.assertEqual(summary["lesson_homework_inquiry"], 1)
        self.assertEqual(summary["class_observation"], 3)  # q1/q2/q3 有教师速记
        self.assertEqual(summary["observation_variant"], 0)  # 未导入题库，如实为 0
        self.assertEqual(summary["bank_core"], 0)
        # q4 未被教师标注：不属于三个选题来源，不进练习卷
        self.assertNotIn("全班都答对的课堂检测题是哪道", self.exported_text(export)[0])

    def exported_text(self, export: dict) -> tuple[str, str]:
        return (
            docx_text(export["student_artifact"]["path"]),
            docx_text(export["teacher_artifact"]["path"]),
        )

    def test_student_paper_never_contains_answers_or_teacher_blocks(self) -> None:
        export = self.export()
        student_text, teacher_text = self.exported_text(export)

        # 题面完整：题干、选项、作业任务都在学生卷
        self.assertIn("我国人口分布的特点是？", student_text)
        self.assertIn("A. 均匀分布", student_text)
        self.assertIn("B. 东多西少", student_text)
        self.assertIn("完成课本第32页活动题。", student_text)
        self.assertIn("查找胡焕庸线资料写一段短评。", student_text)

        # 答案隔离：解析、参考答案、教师参考块、课堂实测、正确选项标记一律不出现
        for marker in STUDENT_ONLY_MARKERS:
            self.assertNotIn(marker, student_text)
        self.assertNotIn("胡焕庸线两侧人口密度差异显著", student_text)
        self.assertNotIn("正确率", student_text)

        # 教师卷包含对应内容（同一份数据，另一侧存在）
        self.assertIn("参考答案", teacher_text)
        self.assertIn("胡焕庸线两侧人口密度差异显著", teacher_text)

    def test_teacher_paper_contains_answers_knowledge_and_real_classroom_data(self) -> None:
        export = self.export()
        student_text, teacher_text = self.exported_text(export)

        # q1：参考答案/解析/考点 + 真实课堂实测（2 人作答、正确率 50%、150 秒、超时 30 秒、已揭示）
        self.assertIn("参考答案：B", teacher_text)
        self.assertIn("参考解析：胡焕庸线两侧人口密度差异显著。", teacher_text)
        self.assertIn("考点：人口分布格局", teacher_text)
        self.assertIn("2 人作答", teacher_text)
        self.assertIn("正确率 50%", teacher_text)
        self.assertIn("投屏用时 150 秒（超时 30 秒），课堂上已揭示答案", teacher_text)
        self.assertIn("课堂速记：存在误区（标签：混淆数量与密度）——有学生认为总量大密度就大", teacher_text)
        self.assertIn("来源：教师课堂速记（原题回炉）", teacher_text)

        # q2：无参考答案如实标注；未投屏因此没有作答数据行
        self.assertIn("参考答案：本题未提供参考答案", teacher_text)
        self.assertIn("参考解析：本题未提供参考解析", teacher_text)
        self.assertIn("课堂速记：部分掌握（标签：密度概念不清）——把总量当密度", teacher_text)
        self.assertNotIn("参考答案：本题未提供参考答案", student_text)

        # q3：投屏但无人作答 → 如实写“未收集到作答数据”，绝不填 0% 或臆测正确率
        self.assertIn("未收集到作答数据", teacher_text)
        self.assertNotIn("正确率 0%", teacher_text)
        self.assertIn("投屏用时 40 秒，课堂上未揭示答案", teacher_text)

        # 缺失题图：保留题号并给占位说明，两卷一致
        self.assertIn("1 张题图未能嵌入", student_text)
        self.assertIn("1 张题图未能嵌入", teacher_text)

        # 选题说明如实标注来源与计数
        self.assertIn("教案课后作业（基础）：1 项", teacher_text)
        self.assertIn("误区变式（题库检索）：0 题", teacher_text)
        self.assertIn("核心目标巩固（题库检索）：0 题", teacher_text)

    def test_export_registers_project_gated_artifacts_and_job(self) -> None:
        export = self.export()
        student_artifact = export["student_artifact"]
        teacher_artifact = export["teacher_artifact"]
        self.assertEqual(student_artifact["artifact_type"], "practice_paper_student")
        self.assertEqual(teacher_artifact["artifact_type"], "practice_paper_teacher")
        for artifact in (student_artifact, teacher_artifact):
            self.assertTrue(Path(artifact["path"]).is_file())
            self.assertTrue(artifact["metadata"]["public_url"].startswith("/files/outputs/"))
            self.assertEqual(artifact["metadata"]["session_id"], self.session_id)
        outputs = self.store.list_outputs(self.project_id)
        types = [item["artifact_type"] for item in outputs]
        self.assertIn("practice_paper_student", types)
        self.assertIn("practice_paper_teacher", types)
        job = self.runtime.get_job(export["job_id"])
        self.assertEqual(job["status"], "success")
        self.assertEqual(job["workflow_type"], "practice_export")

    def test_printable_papers_share_group_material_keep_choices_and_separate_answer_space(self) -> None:
        from docx import Document
        from zipfile import ZipFile
        upload = self.runtime.upload_image_asset(self.project_id, "shared.png", TINY_PNG)
        image = {"url": upload["artifact"]["metadata"]["public_url"], "anchor": "group"}
        question = {"material": "共同读图材料", "images": [image], "options": ["甲", "乙"],
                    "answer": "B", "explanation": "只供教师查阅的解释", "group_key": "same", "bank_id": "bank"}
        items = [{"kind": "task", "origin": "lesson_homework_basic", "text": "比较人口密度并解释原因。",
                  "teacher_guidance": {"answer_points": ["不能混淆总量与密度"]}}]
        items += [{"kind": "question", "origin": "bank_core", "question": {**question, "text": f"读图问题{i}"}}
                  for i in (1, 2)]
        service = self.runtime.classroom.practice_export
        for index, item in enumerate(items):
            item["practice_id"] = f"fixture_{index}"
        service.collect_items = Mock(return_value=(items, [], []))
        export = self.export()
        for kind in ("student", "teacher"):
            path = export[f"{kind}_artifact"]["path"]
            doc = Document(path)
            text = docx_text(path)
            self.assertEqual(text.count("共同读图材料"), 1)
            self.assertEqual(len(doc.inline_shapes), 1)
            self.assertIn("第 2—3 题共用材料", text)
            self.assertIn("读图问题1", text)
            self.assertIn("读图问题2", text)
            self.assertTrue(next(p for p in doc.paragraphs if p.text == "A. 甲").paragraph_format.keep_with_next)
            answer_lines = [p for p in doc.paragraphs if p._p.xpath("./w:pPr/w:pBdr")]
            self.assertEqual(len(answer_lines), 6 if kind == "student" else 0)
            self.assertTrue(all(p._p.xpath("./w:pPr/w:pBdr/w:between") for p in answer_lines))
            self.assertFalse(doc.styles["Title"].element.xpath("./w:pPr/w:pBdr"))
            with ZipFile(path) as zf:
                self.assertIn(b'PAGE', zf.read("word/footer1.xml"))
            if kind == "student":
                self.assertNotIn("只供教师查阅的解释", text)
                self.assertNotIn("不能混淆总量与密度", text)
                self.assertNotIn("选题说明", text)
            else:
                self.assertIn("只供教师查阅的解释", text)
                self.assertIn("第 2 题 教师参考", text)
                self.assertIn("第 3 题 教师参考", text)
                self.assertGreater(text.index("参考答案与讲评"), text.index("读图问题2"))
        # Images anchored to individual questions must not be merged as shared material.
        image["anchor"] = "question"
        export = self.export()
        self.assertEqual(len(Document(export["student_artifact"]["path"]).inline_shapes), 2)

    def test_uploaded_image_is_embedded_into_both_papers(self) -> None:
        upload = self.runtime.upload_image_asset(
            self.project_id, "density_map.png", TINY_PNG, title="人口密度示意"
        )
        image_url = upload["artifact"]["metadata"]["public_url"]
        # 题图在开课前绑定，确保进入开课时刻的课时快照（真实课堂顺序）。
        lesson = self.store.get_lesson(self.lesson.lesson_id)
        lesson.stages[0]["questions"][1]["images"] = [
            {"url": image_url, "width": 640, "height": 480, "content_type": "image/png", "anchor": "question", "order": 1}
        ]
        self.store.upsert_lesson(lesson)
        session_id = self.runtime.classroom.create_class_session(
            self.lesson.lesson_id, self.project_id
        )["session"]["session_id"]
        cw = self.runtime.classroom
        cw.add_session_observation(
            session_id,
            {"stage_id": "s1", "question_id": "q2", "verdict": "partial", "tag": "密度概念不清", "note": ""},
        )

        export = cw.export_session_practice(session_id)
        from docx import Document

        for artifact in (export["student_artifact"], export["teacher_artifact"]):
            doc = Document(artifact["path"])
            self.assertEqual(len(doc.inline_shapes), 1)  # 真实上传的题图已嵌入
            self.assertNotIn("题图未能嵌入", docx_text(artifact["path"]))


@unittest.skipUnless(
    REAL_BANK_ORIGINAL.is_file() and REAL_BANK_ANALYSIS.is_file(),
    "真实题库目录不存在（私有资产不入库），跳过真实题库验收。",
)
class PracticeExportRealBankTest(PracticeExportTestBase):
    """真实题库验收：原卷+解析版导入 → 检索选题 → 双卷导出与内容核对。"""

    def test_real_bank_dual_paper_export(self) -> None:
        runtime, store, project_id = self.build_runtime()
        cw = runtime.classroom
        banks = cw.question_bank.import_files(
            project_id,
            "local_admin",
            [
                (REAL_BANK_ORIGINAL.name, REAL_BANK_ORIGINAL.read_bytes()),
                (REAL_BANK_ANALYSIS.name, REAL_BANK_ANALYSIS.read_bytes()),
            ],
        )
        self.assertGreaterEqual(banks[0]["question_count"], 1)

        search = cw.search_question_banks(project_id, topic="人口分布", knowledge="人口分布 自然因素", limit=5)
        candidates = [item for item in search["items"] if item.get("answer_complete")]
        self.assertTrue(candidates, "真实题库应能检索到答案完备的题目")
        snapshot = cw.question_bank.snapshot_question(candidates[0]["question_id"], project_id=project_id)

        lesson = cw.lesson_service.create_lesson(
            {
                "title": "人口分布真题练习课",
                "subject": "地理",
                "grade": "高一",
                "objectives": ["能描述人口分布格局"],
                "stages": [
                    {
                        "stage_id": "s1",
                        "title": "真题探究",
                        "minutes": 10,
                        "questions": [snapshot],
                    }
                ],
                "plan": {
                    "title": "人口分布真题练习课",
                    "topic": "人口分布",
                    "duration_minutes": 40,
                    "homework": {"basic": ["整理本次真题中的读图方法。"], "inquiry": []},
                    "objectives": ["能描述人口分布格局"],
                    "core_questions": {"core": "人口分布格局如何形成？", "sub_questions": []},
                },
            },
            source="manual",
        )
        session_id = cw.create_class_session(lesson.lesson_id, project_id)["session"]["session_id"]
        cw.launch_session_question(session_id, stage_id="s1", question_id=snapshot["question_id"])
        with store.batch():
            store.add_student_response(session_id, snapshot["question_id"], {"nickname": "小李", "choice_index": 0})
        cw.add_session_observation(
            session_id,
            {
                "stage_id": "s1",
                "question_id": snapshot["question_id"],
                "verdict": "partial",
                "tag": "人口分布",
                "note": "依据不完整",
            },
        )
        cw.end_class_session(session_id)

        export = cw.export_session_practice(session_id)
        self.assertEqual(export["status"], "success")
        summary = {entry["origin"]: entry["count"] for entry in export["selection_summary"]}
        self.assertGreaterEqual(summary["class_observation"], 1)

        student_text = docx_text(export["student_artifact"]["path"])
        teacher_text = docx_text(export["teacher_artifact"]["path"])
        stem_fragment = str(snapshot["text"])[:20]
        self.assertIn(stem_fragment, student_text)
        self.assertIn(stem_fragment, teacher_text)
        # 答案隔离：学生卷无任何教师参考块与参考答案标记
        for marker in STUDENT_ONLY_MARKERS:
            self.assertNotIn(marker, student_text)
        # 教师卷有参考答案/解析与真实作答数据
        self.assertIn("参考答案", teacher_text)
        self.assertIn("1 人作答", teacher_text)
        self.assertIn("来源：教师课堂速记（原题回炉）", teacher_text)
        # 工件经项目门控注册
        for artifact in (export["student_artifact"], export["teacher_artifact"]):
            self.assertTrue(Path(artifact["path"]).stat().st_size > 0)
            self.assertTrue(artifact["metadata"]["public_url"].startswith("/files/outputs/"))


if __name__ == "__main__":
    unittest.main()


def test_local_bank_selection_never_invokes_model_reranking():
    from backend.app.services.question_bank import QuestionBankService
    service = object.__new__(QuestionBankService)
    service._scored_candidates = Mock(return_value=[{"question_id": "q1", "answer_complete": True, "relevance": .9}])
    service._rerank = Mock(side_effect=AssertionError("must stay local"))
    result = service.search(project_id="p1", topic="人口分布", use_llm=False)
    assert result["generator"] == "rules"
    assert result["items"][0]["question_id"] == "q1"
    assert result["items"][0]["auto_selectable"] is True
    service._rerank.assert_not_called()
