"""Guided, population-first lesson-plan co-creation service.

The service deliberately keeps the teacher-facing conversation natural while
persisting a small, validated draft after every turn. The draft is portable
JSON, so refreshing the browser never loses the current design.
"""
from __future__ import annotations

import copy
import json
import logging
import re
from pathlib import Path
from typing import Any, Dict, List, Optional
from uuid import uuid4

from ..models import LessonDesignRecord, LessonRecord
from ..store import RuntimeStore
from .lessons import LessonService
from .minimax_client import MiniMaxClient
from .population_lesson_prep import ALLOWED_GLOBE_THEME_IDS
from .question_bank import QuestionBankService


logger = logging.getLogger(__name__)


STEP_KEYS = (
    "requirements", "analysis", "objectives", "core_questions", "process",
    "question_matching", "capabilities", "rehearsal", "confirmation",
)
STEP_LABELS = {
    "requirements": "需求确认", "analysis": "课标与学情", "objectives": "目标与重难点",
    "core_questions": "核心问题与问题链", "process": "教学过程", "question_matching": "题目匹配",
    "capabilities": "GIS/AI能力", "rehearsal": "预演检查", "confirmation": "确认发布",
}
SECTION_KEYS = (
    "requirements", "curriculum_interpretation", "student_analysis", "textbook_analysis",
    "objectives", "key_difficulties", "methods", "knowledge_structure", "core_questions",
    "stages", "board_design", "question_citations", "homework", "capabilities",
    "design_thinking", "references", "reflection",
)
STEP_SECTIONS = {
    "requirements": ("requirements",),
    "analysis": ("curriculum_interpretation", "student_analysis", "textbook_analysis"),
    "objectives": ("objectives", "key_difficulties", "methods", "knowledge_structure"),
    "core_questions": ("core_questions",),
    "process": ("stages", "board_design"),
    "question_matching": ("question_citations", "homework"),
    "capabilities": ("capabilities",),
    "confirmation": ("design_thinking", "reflection"),
}
REQUIRED_SECTIONS = (
    "requirements", "curriculum_interpretation", "student_analysis", "textbook_analysis",
    "objectives", "key_difficulties", "core_questions", "stages", "capabilities",
)
SECTION_LABELS = {
    "requirements": "教学需求", "curriculum_interpretation": "课标解读", "student_analysis": "学情分析",
    "textbook_analysis": "教材分析", "objectives": "教学目标", "key_difficulties": "教学重难点",
    "methods": "教学方法", "knowledge_structure": "知识结构", "core_questions": "核心问题与问题链",
    "stages": "教学过程", "board_design": "板书设计", "question_citations": "题库引用",
    "homework": "课后作业", "capabilities": "GIS/AI能力", "design_thinking": "设计思路",
    "references": "参考资料", "reflection": "教学反思",
}


def default_draft() -> Dict[str, Any]:
    return {
        "title": "", "subject": "地理", "grade": "", "duration_minutes": 40, "topic": "",
        "requirements": {}, "curriculum_interpretation": "", "student_analysis": "",
        "textbook_analysis": "", "design_thinking": "",
        "objectives": [], "key_difficulties": {"key": [], "difficult": []},
        "methods": [], "knowledge_structure": [],
        "core_questions": {"core": "", "sub_questions": []},
        "stages": [], "board_design": "",
        "question_citations": [], "homework": {"basic": [], "inquiry": []},
        "capabilities": [], "references": [], "reflection": "",
    }


class LessonDesignService:
    """Create, discuss, validate and publish teacher-owned lesson drafts."""

    def __init__(
        self, config: Any, store: RuntimeStore, lesson_service: LessonService,
        minimax_client: Optional[MiniMaxClient] = None, template_service: Any = None,
        catalog_service: Any = None, knowledge_base_service: Any = None,
        resource_search_service: Any = None, question_bank_service: Any = None,
    ):
        self.config = config
        self.store = store
        self.lesson_service = lesson_service
        self.minimax_client = minimax_client
        self.template_service = template_service
        self.catalog_service = catalog_service
        self.knowledge_base_service = knowledge_base_service
        self.resource_search_service = resource_search_service
        self.question_bank_service = question_bank_service

    @staticmethod
    def _normalize_text_lists(draft: Dict[str, Any]) -> None:
        """Canonical text fields stay renderable; keep rich model details separately."""
        fields = [(draft, "objectives", "objectives")]
        core = draft.get("core_questions")
        if isinstance(core, dict):
            fields.append((core, "sub_questions", "core_questions.sub_questions"))
        for parent, key, path in fields:
            if key not in parent:
                continue
            values = parent[key]
            if not isinstance(values, list):
                raise ValueError("教学目标和子问题必须按条目保存，不能使用整段对象。")
            normalized = []
            details = []
            for item in values:
                text = item if isinstance(item, str) else next(
                    (item[field] for field in ("statement", "text", "question", "description")
                     if isinstance(item.get(field), str) and item[field].strip()), None
                ) if isinstance(item, dict) else None
                if not isinstance(text, str) or not text.strip():
                    raise ValueError("教学目标或子问题缺少可读文字，请补充具体内容。")
                normalized.append(text.strip())
                details.append(copy.deepcopy(item) if isinstance(item, dict) else None)
            if any(item is not None for item in details):
                # Archival metadata only; current text and 1-based order remain authoritative.
                draft.setdefault("structured_text_originals", {})[path] = details
            parent[key] = normalized

    @staticmethod
    def _backfill_draft(draft: Dict[str, Any]) -> Dict[str, Any]:
        """旧教案/旧设计载入时补齐新章节默认值，保证字段完整可用。"""
        fresh = default_draft()
        for key, value in fresh.items():
            draft.setdefault(key, copy.deepcopy(value))
        if not isinstance(draft.get("core_questions"), dict):
            draft["core_questions"] = {"core": str(draft.get("core_questions") or ""), "sub_questions": []}
        draft["core_questions"].setdefault("core", "")
        draft["core_questions"].setdefault("sub_questions", [])
        if not isinstance(draft.get("homework"), dict):
            draft["homework"] = {"basic": [], "inquiry": []}
        draft["homework"].setdefault("basic", [])
        draft["homework"].setdefault("inquiry", [])
        if not isinstance(draft.get("question_citations"), list):
            draft["question_citations"] = []
        LessonDesignService._normalize_text_lists(draft)
        return draft

    def create_or_resume(
        self, project_id: str, owner_user_id: str, base_lesson_id: str = "",
        requirements: Optional[Dict[str, Any]] = None,
    ) -> LessonDesignRecord:
        existing = self.store.list_lesson_designs(project_id=project_id, owner_user_id=owner_user_id, active_only=True)
        for item in existing:
            if str(item.base_lesson_id or "") == str(base_lesson_id or ""):
                return self.get(item.design_id)
        draft = default_draft()
        if base_lesson_id:
            lesson = self.store.get_lesson(base_lesson_id)
            if lesson is None:
                raise KeyError("Unknown base lesson")
            lesson_project_id = str((lesson.metadata or {}).get("project_id") or "")
            if lesson.source != "builtin" and lesson_project_id and lesson_project_id != project_id:
                raise ValueError("关联课时不属于当前项目")
            draft.update(copy.deepcopy(lesson.plan or {}))
            draft["title"] = draft.get("title") or lesson.title
            draft["subject"] = draft.get("subject") or lesson.subject
            draft["grade"] = draft.get("grade") or lesson.grade
            draft["objectives"] = copy.deepcopy(lesson.objectives)
            draft["stages"] = copy.deepcopy(lesson.stages)
        self._backfill_draft(draft)
        req = {str(k): v for k, v in (requirements or {}).items()}
        if req:
            draft["requirements"] = {**draft.get("requirements", {}), **req}
        base_draft = copy.deepcopy(draft) if base_lesson_id else {}
        design = LessonDesignRecord.create(
            project_id=project_id, owner_user_id=owner_user_id, base_lesson_id=base_lesson_id,
            requirements=req, draft=draft, base_draft=base_draft,
        )
        design.section_status = {key: ("proposed" if base_lesson_id and self._section_has_content(draft.get(key)) else "pending") for key in SECTION_KEYS}
        design.diff_summary = self._build_diff_summary(design)
        return self.store.upsert_lesson_design(design)

    def get(self, design_id: str) -> LessonDesignRecord:
        design = self.store.get_lesson_design(design_id)
        if design is None:
            raise KeyError("Unknown lesson design")
        design = copy.deepcopy(design)
        self._backfill_draft(design.draft)
        return design

    def capability_catalog(self) -> List[Dict[str, Any]]:
        items: List[Dict[str, Any]] = [
            {"id": "map_2d", "label": "2D 地图场景", "kind": "map", "available": True},
            {"id": "globe_3d", "label": "3D 地球场景", "kind": "map", "available": True},
            {"id": "basemap_switch", "label": "底图切换", "kind": "map", "available": True},
            {"id": "population_sources", "label": "人口专题来源卡", "kind": "data", "available": True},
            {"id": "population_top20", "label": "人口指标 TOP20 查询", "kind": "analysis", "available": True},
            {"id": "annotation_measurement", "label": "地图标注与测量", "kind": "map", "available": True},
            {"id": "image_understanding", "label": "图片识别", "kind": "ai", "available": True},
            {"id": "knowledge_base", "label": "项目知识库", "kind": "knowledge", "available": bool(self.knowledge_base_service)},
            {"id": "class_records", "label": "课堂记录与报告", "kind": "classroom", "available": True},
        ]
        if self.template_service is not None:
            try:
                for item in self.template_service.list_templates().get("items", []):
                    if isinstance(item, dict) and item.get("template_id"):
                        items.append({
                            "id": str(item["template_id"]), "label": str(item.get("title") or item["template_id"]),
                            "kind": "template", "available": True,
                        })
            except Exception:
                pass
        if self.catalog_service is not None:
            try:
                for item in self.catalog_service.list_catalog().get("items", []):
                    if isinstance(item, dict) and item.get("id"):
                        items.append({
                            "id": str(item["id"]), "label": str(item.get("name") or item["id"]),
                            "kind": "dataset",
                            "available": str(item.get("status") or "available") not in {"missing", "unavailable"},
                            "source_year": str(item.get("source_year") or ""),
                        })
            except Exception:
                pass
        result, seen = [], set()
        for item in items:
            if item["id"] not in seen:
                seen.add(item["id"])
                result.append(item)
        return result

    def turn(self, design_id: str, message: str, expected_revision: Optional[int] = None, step: str = "") -> Dict[str, Any]:
        design = self.get(design_id)
        if expected_revision is not None and expected_revision != design.revision:
            raise ValueError("教案草稿已更新，请刷新后再继续。")
        message = str(message or "").strip()
        if not message:
            raise ValueError("请先告诉我你的教学想法或修改要求。")
        current_step = step if step in STEP_KEYS else design.current_step
        if any(token in message for token in ("返回上一步", "回到上一步")):
            # 「返回上一步」是流程控制指令：短路处理，不进 LLM/规则需求解析，
            # 否则会被当成课题文本（如“我理解你想做一节‘返回上一步’”）。
            current_index = STEP_KEYS.index(current_step) if current_step in STEP_KEYS else 0
            design.revision += 1
            if current_index <= 0:
                reply = (
                    "现在已在第一步「需求确认」，没有更早的步骤了。"
                    "可以直接输入你的教学想法；想退出教案设计时点「退出设计」即可，草稿已自动保存。"
                )
            else:
                current_step = STEP_KEYS[current_index - 1]
                design.current_step = current_step
                design.pending_next_step = ""
                reply = (
                    f"已回到「{STEP_LABELS.get(current_step, current_step)}」。"
                    "这一部分可以重新讨论或修改；确认无误后点「接受本节」继续。"
                )
            design.turns.append({
                "revision": design.revision, "step": current_step, "message": message,
                "reply": reply, "section_patch": {},
            })
            self.store.upsert_lesson_design(design)
            return {
                "status": "success", "assistant_message": reply,
                "next_step": design.current_step, "step_label": STEP_LABELS.get(design.current_step, design.current_step),
                "draft": copy.deepcopy(design.draft), "section_status": copy.deepcopy(design.section_status),
                "source_refs": copy.deepcopy(design.source_refs), "capability_bindings": copy.deepcopy(design.capability_bindings),
                "revision": design.revision, "suggestions": [],
                "retrieval": {"mode": "none", "used": False},
                "diff_summary": copy.deepcopy(design.diff_summary),
                "review_sections": [],
                "rehearsal_report": None,
                "plan_items": self._plan_items(design),
                "retrieval_candidates": [],
                "auto_bound_questions": [],
                "active_design_question": self._active_design_question(design),
            }
        schedule = self._requested_schedule(message) if current_step == "process" else []
        result = self._ask_minimax(design, current_step, message)
        generation_mode = "model" if result is not None else "rules"
        if current_step == "process" and schedule:
            stages = (result or {}).get("section_patch", {}).get("stages")
            if not self._matches_schedule(stages, schedule):
                logger.warning("Lesson generation rejected: step=process reason=%s", "schedule_mismatch" if result else "model_unavailable")
                # Never replace an explicit teacher plan with the generic three-stage template.
                raise ValueError("本次生成未能按你指定的环节和时长完成，原草稿已保留。请重试，或在教学过程里直接编辑。")
        if result is None:
            result = self._fallback_turn(design, current_step, message)
        if current_step == "requirements":
            result = self._normalize_requirements_result(result, message, design.draft)
        patch = result.get("section_patch") if isinstance(result, dict) else {}
        patch = patch if isinstance(patch, dict) else {}
        protected_questions = self._questions_to_preserve(design.draft, message)
        if "stages" in patch and not self._preserves_questions(patch["stages"], protected_questions):
            raise ValueError("本次生成改动或遗漏了需保留的题目，原草稿已保留。请重试；题库题和教师录入题请通过题目编辑入口修改。")
        # 已确认章节默认是稳定约束；只有教师明确提出修改/返回时才重新打开。
        reopen = any(token in message for token in ("修改", "调整", "返回", "重做", "换一种"))
        for key in list(patch):
            if design.section_status.get(key) == "confirmed" and not reopen:
                patch.pop(key, None)
            elif design.section_status.get(key) == "confirmed" and reopen:
                design.section_status[key] = "pending"
        self._merge_patch(design.draft, patch)
        next_step = str(result.get("next_step") or self._next_step(current_step))
        if next_step not in STEP_KEYS:
            next_step = self._next_step(current_step)
        for key in patch:
            if key in SECTION_KEYS:
                design.section_status[key] = "proposed"
        if patch:
            design.current_step = current_step
            design.pending_next_step = next_step
        else:
            design.current_step = next_step
            design.pending_next_step = ""
        design.revision += 1
        retrieval_mode, retrieved_refs = self._retrieve(message, design.owner_user_id)
        design.source_refs = self._merge_refs(design.source_refs, result.get("source_refs"))
        design.source_refs = self._merge_refs(design.source_refs, retrieved_refs)
        design.draft["references"] = copy.deepcopy(design.source_refs)
        design.capability_bindings = self._validate_bindings(result.get("capability_bindings") or design.draft.get("capabilities") or [])
        design.draft["capabilities"] = copy.deepcopy(design.capability_bindings)
        design.diff_summary = self._build_diff_summary(design)
        # 题目匹配步骤：按环节自动检索题库候选并快照可自动选用的题目。
        retrieval_candidates: List[Dict[str, Any]] = []
        auto_bound: List[Dict[str, Any]] = []
        if current_step == "question_matching":
            retrieval_candidates, auto_bound = self._auto_bind_questions(design)
        rehearsal_report = None
        if current_step in {"rehearsal", "confirmation"}:
            rehearsal_report = self.validate_plan(design.draft, design.source_refs)
            unconfirmed = [SECTION_LABELS.get(key, key) for key in REQUIRED_SECTIONS
                           if design.section_status.get(key) != "confirmed"]
            if rehearsal_report["errors"]:
                reply = "预演检查未通过：" + "；".join(rehearsal_report["errors"])
            else:
                reply = "教案结构检查通过。"
            if unconfirmed:
                reply += " 尚待确认：" + "、".join(unconfirmed) + "。"
            reply += " 当前尚未发布为课时草稿，也未生成 Word。"
            if rehearsal_report["warnings"]:
                reply += " 提醒：" + "；".join(rehearsal_report["warnings"])
            result["reply"] = reply
        if generation_mode == "rules":
            result["reply"] = "本轮 AI 未返回有效内容，以下为规则草稿与系统检查结果，需逐项核对。" + str(result.get("reply") or "")
        design.turns.append({
            "revision": design.revision, "step": current_step, "message": message,
            "reply": str(result.get("reply") or ""), "section_patch": copy.deepcopy(patch),
            "generation_mode": generation_mode,
        })
        self.store.upsert_lesson_design(design)
        return {
            "status": "success", "assistant_message": str(result.get("reply") or self._natural_prompt(next_step)),
            "generation_mode": generation_mode,
            "next_step": design.current_step, "step_label": STEP_LABELS.get(design.current_step, design.current_step),
            "draft": copy.deepcopy(design.draft), "section_status": copy.deepcopy(design.section_status),
            "source_refs": copy.deepcopy(design.source_refs), "capability_bindings": copy.deepcopy(design.capability_bindings),
            "revision": design.revision, "suggestions": [str(item) for item in result.get("suggestions") or []],
            "retrieval": {"mode": retrieval_mode, "used": bool(retrieved_refs)},
            "diff_summary": copy.deepcopy(design.diff_summary),
            "review_sections": [key for key in STEP_SECTIONS.get(current_step, ()) if design.section_status.get(key) == "proposed"],
            "rehearsal_report": rehearsal_report,
            "plan_items": self._plan_items(design),
            "retrieval_candidates": retrieval_candidates,
            "auto_bound_questions": auto_bound,
            "active_design_question": self._active_design_question(design),
        }

    def resolve(self, design_id: str, section_id: str, decision: str = "accept", teacher_note: str = "", expected_revision: Optional[int] = None, value: Any = None) -> Dict[str, Any]:
        design = self.get(design_id)
        if expected_revision is not None and expected_revision != design.revision:
            raise ValueError("教案草稿已更新，请刷新后再操作。")
        if section_id not in SECTION_KEYS and section_id not in STEP_SECTIONS and section_id not in {"title", "grade", "duration_minutes"}:
            raise ValueError("未知教案章节")
        section_ids = STEP_SECTIONS.get(section_id, (section_id,))
        normalized_decision = str(decision).lower()
        if normalized_decision in {"edit", "direct_edit", "直接编辑"}:
            patch_value = copy.deepcopy(value)
            # A section can share a step name (objectives/core_questions).
            # Single-section payloads and grouped step payloads are both supported.
            if section_id == "objectives" and isinstance(patch_value, list):
                patch_value = {"objectives": patch_value}
            elif section_id == "core_questions" and isinstance(patch_value, dict) and "core_questions" not in patch_value:
                patch_value = {"core_questions": patch_value}
            if section_id in STEP_SECTIONS:
                if not isinstance(patch_value, dict):
                    raise ValueError("当前步骤的直接编辑内容格式不正确")
                allowed_group = set(STEP_SECTIONS[section_id])
                for key, item in patch_value.items():
                    if key in allowed_group:
                        if key == "capabilities":
                            if not isinstance(item, list):
                                raise ValueError("系统能力配置格式不正确")
                            validated = self._validate_bindings(item)
                            requested_ids = {
                                str(binding if isinstance(binding, str) else binding.get("id") or binding.get("capability_id") or "")
                                for binding in item if isinstance(binding, (str, dict))
                            }
                            accepted_ids = {binding["id"] for binding in validated}
                            unknown = sorted(capability_id for capability_id in requested_ids - accepted_ids if capability_id)
                            if unknown:
                                raise ValueError("包含不可用系统能力：" + "、".join(unknown))
                            design.capability_bindings = validated
                            design.draft[key] = copy.deepcopy(validated)
                        else:
                            design.draft[key] = copy.deepcopy(item)
                        design.section_status[key] = "proposed"
                if not any(key in allowed_group for key in patch_value):
                    raise ValueError("没有可保存的当前步骤内容")
                self._normalize_text_lists(design.draft)
                design.revision += 1
                design.diff_summary = self._build_diff_summary(design)
                self.store.upsert_lesson_design(design)
                return {"status": "success", "message": "已保存当前步骤的直接编辑内容。", "design": design.to_dict(), **self.session_view(design)}
            if section_id == "duration_minutes":
                try:
                    patch_value = max(1, int(value))
                except (TypeError, ValueError) as exc:
                    raise ValueError("课时时长必须是正整数") from exc
            if section_id in {"title", "grade"}:
                patch_value = str(value or "").strip()
                if not patch_value:
                    raise ValueError("直接编辑内容不能为空")
            design.draft[section_id] = patch_value
            self._normalize_text_lists(design.draft)
            if section_id in SECTION_KEYS:
                design.section_status[section_id] = "proposed"
            design.revision += 1
            design.diff_summary = self._build_diff_summary(design)
            self.store.upsert_lesson_design(design)
            return {"status": "success", "message": "已保存直接编辑内容。", "design": design.to_dict(), **self.session_view(design)}
        if normalized_decision in {"accept", "accepted", "确认", "接受"}:
            accepted = []
            for key in section_ids:
                if self._section_has_content(design.draft.get(key)):
                    design.section_status[key] = "confirmed"
                    accepted.append(key)
            if not accepted:
                raise ValueError("这一部分还没有可确认的内容")
            design.current_step = design.pending_next_step or self._next_step(design.current_step)
            design.pending_next_step = ""
            design.revision += 1
            message = "已接受这一部分。"
        else:
            for key in section_ids:
                if design.section_status.get(key) == "proposed":
                    design.section_status[key] = "pending"
            if teacher_note:
                design.turns.append({"revision": design.revision, "step": design.current_step, "message": teacher_note, "reply": ""})
            message = "好的，我们保留现稿，按你的补充继续修改。"
        design.diff_summary = self._build_diff_summary(design)
        self.store.upsert_lesson_design(design)
        return {"status": "success", "message": message, "design": design.to_dict(), **self.session_view(design)}

    # ------------------------------------------------------------------
    # 题目匹配：题库快照绑定 / 手动录入 / 替换 / 移除
    # ------------------------------------------------------------------

    def bind_question(
        self,
        design_id: str,
        stage_id: str,
        question_id: str = "",
        manual: Optional[Dict[str, Any]] = None,
        action: str = "add",
        position: Optional[int] = None,
        expected_revision: Optional[int] = None,
    ) -> Dict[str, Any]:
        """把题库题目快照或教师手动题目写入教案环节（快照不可变）。

        - ``action=add`` + ``question_id``：题库题目快照入库；
        - ``action=add`` + ``manual``：教师手动录题（source=teacher_manual），
          若要用于真实课堂/课后练习，必须带答案与解析；
        - ``action=remove``：按 question_id 移除环节中的题目与引用。
        """
        design = self.get(design_id)
        if expected_revision is not None and expected_revision != design.revision:
            raise ValueError("教案草稿已更新，请刷新后再操作。")
        stages = design.draft.get("stages") or []
        stage = next((item for item in stages if str(item.get("stage_id")) == stage_id), None)
        if stage is None:
            raise KeyError(f"Unknown stage: {stage_id}")
        questions = stage.setdefault("questions", [])
        citations = design.draft.setdefault("question_citations", [])

        if action == "remove":
            target = str(question_id)
            if not target:
                raise ValueError("移除题目需要 question_id。")
            before = len(questions)
            questions[:] = [item for item in questions if str(item.get("question_id")) != target]
            citations[:] = [
                item for item in citations
                if not (str(item.get("question_id")) == target and str(item.get("stage_id")) == stage_id)
            ]
            if len(questions) == before:
                raise KeyError(f"环节 {stage_id} 中没有题目 {target}")
        elif question_id:
            if any(str(item.get("question_id")) == str(question_id) for item in questions):
                raise ValueError("这道题已经在本环节中。")
            snapshot = self._snapshot_bank_question(question_id, design.project_id)
            if position is not None and isinstance(position, int) and 0 <= position <= len(questions):
                questions.insert(position, snapshot)
            else:
                questions.append(snapshot)
            citations.append({
                "question_id": snapshot["question_id"], "bank_id": snapshot.get("bank_id", ""),
                "group_key": snapshot.get("group_key", ""), "stage_id": stage_id,
                "source": "question_bank", "number": snapshot.get("number", ""),
                "year": snapshot.get("year", ""), "region": snapshot.get("region", ""),
                "source_paper": snapshot.get("source_paper", ""),
                "bound_at_revision": design.revision + 1,
                "selection_reason": "教师手动选用",
            })
        elif manual is not None:
            snapshot = self._normalize_manual_question(manual, stage_id, len(questions) + 1)
            if position is not None and isinstance(position, int) and 0 <= position <= len(questions):
                questions.insert(position, snapshot)
            else:
                questions.append(snapshot)
            citations.append({
                "question_id": snapshot["question_id"], "bank_id": "", "group_key": "",
                "stage_id": stage_id, "source": "teacher_manual", "number": "",
                "bound_at_revision": design.revision + 1,
                "selection_reason": "教师手动录入",
            })
        else:
            raise ValueError("绑定题目需要 question_id 或 manual 内容。")

        for key in ("question_citations", "homework"):
            if design.section_status.get(key) == "confirmed":
                design.section_status[key] = "proposed"
        design.revision += 1
        design.diff_summary = self._build_diff_summary(design)
        self.store.upsert_lesson_design(design)
        return {
            "status": "success",
            "message": "题目已更新。" if action == "add" else "题目已移除。",
            "design": design.to_dict(),
            "stage": copy.deepcopy(stage),
        }

    def _snapshot_bank_question(self, question_id: str, project_id: str) -> Dict[str, Any]:
        """从题库取题并做成不可变快照（含材料/题图/答案/解析）。"""
        if self.question_bank_service is None:
            raise ValueError("题库服务不可用。")
        return self.question_bank_service.snapshot_question(question_id, project_id=project_id)

    def _normalize_snapshot_question(self, question: Dict[str, Any]) -> Dict[str, Any]:
        """压平为题目快照；逻辑收敛在题库服务，教案设计与模拟测试共用。"""
        return QuestionBankService.normalize_snapshot_question(question)

    def _normalize_manual_question(self, manual: Dict[str, Any], stage_id: str, index: int) -> Dict[str, Any]:
        """手动题目规范化；逻辑收敛在题库服务，教案设计与模拟测试共用。"""
        return QuestionBankService.build_manual_question(manual, stage_id, index)

    def session_view(self, design: LessonDesignRecord) -> Dict[str, Any]:
        """会话响应的扩展字段（Plan 卡片与当前推进问题），create/get/turn 共用。"""
        return {
            "plan_items": self._plan_items(design),
            "retrieval_candidates": [],
            "auto_bound_questions": [],
            "active_design_question": self._active_design_question(design),
        }

    def _auto_bind_questions(self, design: LessonDesignRecord) -> tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        """题目匹配步骤：按环节检索题库，快照答案完备且相关度达标的题目。

        返回 (候选列表, 本轮自动绑定列表)。已绑定过的题目自动排除；没有
        题库或检索不到达标题目时返回空列表，由教师手动挑选。
        """
        if self.question_bank_service is None:
            return [], []
        draft = design.draft
        stages = draft.get("stages") or []
        if not stages:
            return [], []
        citations = draft.setdefault("question_citations", [])
        bound_ids = {str(item.get("question_id")) for item in citations}
        candidates: List[Dict[str, Any]] = []
        auto_bound: List[Dict[str, Any]] = []
        for stage in stages:
            query_parts = [
                str(draft.get("topic") or draft.get("title") or ""),
                str(stage.get("knowledge_point") or ""),
                str(stage.get("knowledge_unit") or ""),
            ]
            query = " ".join(part for part in query_parts if part)
            if not query:
                continue
            try:
                result = self.question_bank_service.search(
                    project_id=design.project_id,
                    topic=str(draft.get("topic") or ""),
                    knowledge=query,
                    objectives=[str(item) for item in draft.get("objectives") or []],
                    exclude_ids=sorted(bound_ids),
                    limit=5,
                )
            except Exception:
                continue
            for item in result.get("items") or []:
                entry = {
                    "stage_id": str(stage.get("stage_id") or ""),
                    "stage_title": str(stage.get("title") or ""),
                    "question_id": str(item.get("question_id") or ""),
                    "bank_id": str(item.get("bank_id") or ""),
                    "group_key": str(item.get("group_key") or ""),
                    "number": str(item.get("number") or ""),
                    "type": str(item.get("type") or "open"),
                    "stem": str(item.get("stem") or "")[:120],
                    "material": str(item.get("material") or "")[:120],
                    "year": str(item.get("year") or ""),
                    "region": str(item.get("region") or ""),
                    "source_paper": str(item.get("source_paper") or ""),
                    "answer_complete": bool(item.get("answer_complete")),
                    "relevance": float(item.get("relevance") or 0.0),
                    "auto_selectable": bool(item.get("auto_selectable")),
                    "selection_reason": str(item.get("selection_reason") or ""),
                    "image_count": len(item.get("images") or []),
                }
                if entry not in candidates:
                    candidates.append(entry)
            stage_bank_questions = [
                item for item in (stage.get("questions") or [])
                if str(item.get("source")) == "question_bank"
            ]
            if stage_bank_questions:
                continue
            for item in result.get("items") or []:
                if not item.get("auto_selectable"):
                    continue
                snapshot = self._normalize_snapshot_question(item)
                stage.setdefault("questions", []).append(snapshot)
                citations.append({
                    "question_id": snapshot["question_id"], "bank_id": snapshot.get("bank_id", ""),
                    "group_key": snapshot.get("group_key", ""),
                    "stage_id": str(stage.get("stage_id") or ""),
                    "source": "question_bank", "number": snapshot.get("number", ""),
                    "year": snapshot.get("year", ""), "region": snapshot.get("region", ""),
                    "source_paper": snapshot.get("source_paper", ""),
                    "bound_at_revision": design.revision + 1,
                    "relevance": float(item.get("relevance") or 0.0),
                    "selection_reason": str(item.get("selection_reason") or "相关度达标自动选用"),
                })
                bound_ids.add(snapshot["question_id"])
                auto_bound.append({
                    "stage_id": str(stage.get("stage_id") or ""),
                    "question_id": snapshot["question_id"],
                    "number": snapshot.get("number", ""),
                    "stem": snapshot["text"][:80],
                    "relevance": float(item.get("relevance") or 0.0),
                    "selection_reason": str(item.get("selection_reason") or "相关度达标自动选用"),
                })
                break
        return candidates, auto_bound

    def _plan_items(self, design: LessonDesignRecord) -> List[Dict[str, Any]]:
        """当前步骤的可编辑内容卡片（Plan 式三栏布局的数据源）。"""
        step = design.current_step
        items: List[Dict[str, Any]] = []
        for key in STEP_SECTIONS.get(step, ()):
            value = design.draft.get(key)
            if key == "stages":
                value = [
                    {
                        "stage_id": stage.get("stage_id", ""),
                        "title": stage.get("title", ""),
                        "minutes": stage.get("minutes", 0),
                        "knowledge_point": stage.get("knowledge_point", ""),
                        "question_count": len(stage.get("questions") or []),
                    }
                    for stage in (value or [])
                ]
            items.append({
                "section_key": key,
                "label": SECTION_LABELS.get(key, key),
                "value": copy.deepcopy(value),
                "status": design.section_status.get(key, "pending"),
            })
        return items

    def _active_design_question(self, design: LessonDesignRecord) -> str:
        """当前步骤的下一个推进问题——每轮只提一个。"""
        draft = design.draft
        step = design.current_step
        if step == "requirements":
            if not str(draft.get("grade") or "").strip():
                return "这节课面向哪个年级？"
            if not str(draft.get("title") or draft.get("topic") or "").strip():
                return "课题名称定为什么？"
            duration = int(draft.get("duration_minutes") or 0)
            if duration != 40:
                return "正式课堂固定为 40 分钟，按 40 分钟设计可以吗？"
        elif step == "analysis":
            if not str(draft.get("curriculum_interpretation") or "").strip():
                return "这节课对应的课标条目原文是什么？"
            if not str(draft.get("student_analysis") or "").strip():
                return "学生目前对这部分内容的已有基础和常见误区是什么？"
        elif step == "objectives":
            objectives = draft.get("objectives") or []
            if not objectives:
                return "这节课最希望学生带走的 3 个可观察目标是什么？"
            if len(objectives) < 3:
                return "目标现在只有 {} 个，能再补充到 3 个吗？".format(len(objectives))
            if not (draft.get("key_difficulties") or {}).get("difficult"):
                return "这节课最难突破的一个点是什么？"
        elif step == "core_questions":
            core = draft.get("core_questions") or {}
            if not str(core.get("core") or "").strip():
                return "贯穿这节课的核心问题用一句话怎么说？"
            subs = core.get("sub_questions") or []
            if len(subs) < 2:
                return "核心问题可以拆成哪 2-4 个递进的子问题？"
        elif step == "process":
            stages = draft.get("stages") or []
            if not stages:
                return "教学过程从哪个情境或现象切入？"
            for stage in stages:
                if not (stage.get("student_activities") or stage.get("activities")):
                    return f"环节“{stage.get('title', '')}”里学生具体做什么？"
                if not str(stage.get("knowledge_conclusion") or "").strip():
                    return f"环节“{stage.get('title', '')}”要落下的知识结论是什么？"
        elif step == "question_matching":
            citations = draft.get("question_citations") or []
            if not citations:
                return "还没有从题库选题，先看自动检索的候选还是你指定一组真题？"
        elif step == "capabilities":
            if not (draft.get("capabilities") or []):
                return "这节课哪些环节需要地图、数据或 AI 能力？"
        elif step == "rehearsal":
            return "预演检查发现的问题要现在调整，还是回到对应环节修改？"
        elif step == "confirmation":
            if not str(draft.get("design_thinking") or "").strip():
                return "确认发布前，先用 100-150 字概括这节课的设计思路好吗？"
        return self._natural_prompt(step)

    def rehearse(self, design_id: str) -> Dict[str, Any]:
        design = self.get(design_id)
        return self.validate_plan(design.draft, design.source_refs)

    def validate_plan(
        self, draft: Dict[str, Any], source_refs: List[Dict[str, Any]]
    ) -> Dict[str, Any]:
        """预演校验（教案设计与模拟测试共用）：draft 为完整教案数据。"""
        errors, warnings = [], []
        if not str(draft.get("title") or draft.get("topic") or "").strip():
            errors.append("还缺少课题名称。")
        if not str(draft.get("grade") or "").strip():
            warnings.append("尚未填写年级，课堂难度需要教师确认。")
        objectives = draft.get("objectives") or []
        if not isinstance(objectives, list) or not objectives:
            errors.append("至少需要一个可观察的教学目标。")
        elif len(objectives) > 4:
            warnings.append(f"教学目标有 {len(objectives)} 个，建议精简到 3-4 个可观察目标。")
        # 核心问题与递进问题链
        core = draft.get("core_questions") or {}
        if not isinstance(core, dict):
            core = {}
        if not str(core.get("core") or "").strip():
            errors.append("还缺少一个贯穿课堂的核心问题。")
        chain = [str(item).strip() for item in (core.get("sub_questions") or []) if str(item).strip()]
        if len(chain) < 2:
            errors.append("核心问题需要拆成至少 2 个递进子问题。")
        elif len(chain) > 4:
            errors.append(f"递进子问题有 {len(chain)} 个，请精简到 2-4 个。")
        # 设计思路 100-150 字
        thinking = str(draft.get("design_thinking") or "").strip()
        if not thinking:
            errors.append("还缺少设计思路（100-150 字）。")
        elif not (100 <= len(thinking) <= 150):
            warnings.append(f"设计思路现在 {len(thinking)} 字，建议控制在 100-150 字。")
        # 板书设计 / 作业 / 预设反思
        if not str(draft.get("board_design") or "").strip():
            warnings.append("还没有板书设计，确认发布前建议补充。")
        homework = draft.get("homework") or {}
        if not isinstance(homework, dict):
            homework = {}
        if not [item for item in homework.get("basic") or [] if str(item).strip()]:
            errors.append("还缺少基础作业。")
        if not [item for item in homework.get("inquiry") or [] if str(item).strip()]:
            warnings.append("还没有探究作业，建议补充一道开放探究任务。")
        if not str(draft.get("reflection") or "").strip():
            warnings.append("还没有预设教学反思，确认发布前建议补充。")
        stages = draft.get("stages") or []
        if not isinstance(stages, list) or not stages:
            errors.append("还没有教学过程环节。")
        total = 0
        capability_items = self.capability_catalog()
        valid_ids = {item["id"] for item in capability_items if item.get("available", True)}
        dataset_ids = {item["id"] for item in capability_items if item.get("kind") == "dataset" and item.get("available", True)}
        objective_refs_seen: set = set()
        for index, stage in enumerate(stages, 1):
            if not isinstance(stage, dict):
                errors.append(f"第{index}个环节格式不完整。")
                continue
            try:
                minutes = int(stage.get("minutes") or 0)
            except (TypeError, ValueError):
                minutes = 0
            total += max(0, minutes)
            stage_label = stage.get('title') or f'环节{index}'
            if not str(stage.get("title") or "").strip():
                errors.append(f"第{index}个环节缺少名称。")
            if not str(stage.get("design_intent") or stage.get("content") or "").strip():
                warnings.append(f"“{stage_label}”还可以补充设计意图。")
            if not str(stage.get("knowledge_conclusion") or "").strip():
                errors.append(f"环节“{stage_label}”还没有知识结论。")
            student_activities = stage.get("student_activities") or stage.get("activities") or []
            if not student_activities:
                errors.append(f"环节“{stage_label}”还没有学生活动。")
            if not (str(stage.get("material") or "").strip() or (stage.get("question_chain") or [])):
                warnings.append(f"环节“{stage_label}”还没有材料或问题链，建议补充其一。")
            point = str(stage.get("knowledge_point") or "").strip()
            if point and not student_activities:
                errors.append(f"核心知识点“{point}”还没有学生活动。")
            for ref in stage.get("objective_refs") or []:
                try:
                    objective_refs_seen.add(int(ref))
                except (TypeError, ValueError):
                    continue
            questions = stage.get("questions") or []
            if not isinstance(questions, list) or not any(isinstance(item, dict) and str(item.get("text") or "").strip() for item in questions):
                errors.append(f"环节“{stage_label}”至少需要一个明确问题。")
            for question in questions:
                if not isinstance(question, dict):
                    continue
                source = str(question.get("source") or "")
                label = str(question.get("number") or "") or str(question.get("text") or "")[:16]
                if source == "question_bank":
                    if not question.get("answer_complete"):
                        errors.append(f"环节“{stage_label}”的题库题目（{label}）答案不完备，不能进入真实课堂。")
                    for image in question.get("images") or []:
                        if isinstance(image, dict) and not str(image.get("url") or "").strip():
                            errors.append(f"环节“{stage_label}”的题目（{label}）存在缺失题图，发布前需要补全。")
                elif source == "teacher_manual":
                    if not str(question.get("answer") or "").strip() and question.get("answer_index") is None and not any(str(sub.get("answer") or "").strip() for sub in question.get("sub_questions") or []):
                        errors.append(f"环节“{stage_label}”的手动题目（{label}）还没有答案；用于真实课堂或课后练习必须补齐。")
                    if not str(question.get("explanation") or "").strip():
                        errors.append(f"环节“{stage_label}”的手动题目（{label}）还没有解析；用于真实课堂或课后练习必须补齐。")
            for template_id in (stage.get("scene") or {}).get("templates") or []:
                if template_id not in valid_ids:
                    errors.append(f"环节“{stage_label}”引用了不可用能力 {template_id}。")
            for dataset_id in (stage.get("scene") or {}).get("catalog_layers") or []:
                if dataset_id not in dataset_ids:
                    errors.append(f"环节“{stage_label}”引用了不可用数据 {dataset_id}。")
            globe = (stage.get("scene") or {}).get("globe") or {}
            if isinstance(globe, dict) and globe.get("enabled"):
                unknown_themes = [theme for theme in globe.get("themes") or [] if str(theme) not in ALLOWED_GLOBE_THEME_IDS]
                if unknown_themes:
                    errors.append(f"环节“{stage_label}”引用了不可用三维主题：{'、'.join(map(str, unknown_themes))}。")
        # 每个教学目标都有活动支撑
        if isinstance(objectives, list) and objectives:
            refs = {ref - 1 for ref in objective_refs_seen}
            if objective_refs_seen:
                uncovered = [
                    str(objectives[position])[:24]
                    for position in range(len(objectives))
                    if position not in refs
                ]
                if uncovered:
                    errors.append("以下教学目标还没有对应环节活动：" + "；".join(uncovered))
            else:
                warnings.append("各环节尚未标注 objective_refs，无法核对目标-活动对应关系。")
        duration = int(draft.get("duration_minutes") or 0)
        if duration > 0 and total != duration:
            errors.append(
                f"各环节合计 {total} 分钟，与课堂时长 {duration} 分钟不一致；"
                "请先调整环节时长，再定稿。"
            )
        if not source_refs:
            warnings.append("目前没有引用资料；如需课标或年份核验，请在下一轮明确提出。")
        for source in source_refs:
            if isinstance(source, dict) and source.get("dataset_id") and not str(source.get("year") or source.get("source_year") or "").strip():
                warnings.append(f"资料“{source.get('title') or source.get('dataset_id')}”尚未标明年份。")
        return {"ready": not errors, "errors": errors, "warnings": warnings, "total_minutes": total, "duration_minutes": duration, "capabilities": self.capability_catalog()}

    def finalize(self, design_id: str, expected_revision: Optional[int] = None, apply_base: bool = False) -> Dict[str, Any]:
        design = self.get(design_id)
        if expected_revision is not None and expected_revision != design.revision:
            raise ValueError("教案草稿已更新，请刷新后再确认。")
        report = self.rehearse(design_id)
        if not report["ready"]:
            raise ValueError("教案还不能保存：" + " ".join(report["errors"]))
        unconfirmed = [key for key in REQUIRED_SECTIONS if design.section_status.get(key) != "confirmed"]
        if unconfirmed:
            labels = "、".join(SECTION_LABELS.get(key, key) for key in unconfirmed)
            raise ValueError("请先逐项确认这些章节：" + labels)
        if design.base_lesson_id and apply_base:
            base = self.store.get_lesson(design.base_lesson_id)
            if base is None or base.source == "builtin":
                raise ValueError("内置课时不能直接覆盖，请保存为新的教师草稿。")
        draft = copy.deepcopy(design.draft)
        lesson = self.lesson_service.create_lesson({
            "title": draft.get("title") or draft.get("topic") or "人口地理教案",
            "subject": draft.get("subject") or "地理", "grade": draft.get("grade") or "",
            "objectives": draft.get("objectives") or [], "stages": draft.get("stages") or [],
            "metadata": {
                "design_id": design.design_id, "project_id": design.project_id,
                "duration_minutes": draft.get("duration_minutes", 40),
                "lesson_version": 1, "ready_for_class": False,
                "created_from": "lesson_design",
                "question_citation_count": len(draft.get("question_citations") or []),
            },
            "plan": draft,
        }, source="assistant_draft", owner_user_id=design.owner_user_id)
        design.status, design.final_lesson_id, design.current_step = "finalized", lesson.lesson_id, "confirmation"
        design.revision += 1
        self.store.upsert_lesson_design(design)
        # 草稿完成即产出第一版 Word；模拟测试通过后重新生成新版本。
        export_result: Dict[str, Any] = {}
        try:
            export_result = self.export_docx(lesson.lesson_id, design.project_id, design.design_id)
        except Exception:
            export_result = {"status": "skipped", "message": "第一版 Word 导出失败，可稍后在课时详情重新导出。"}
        return {
            "status": "success", "lesson": lesson.to_dict(), "design": {**design.to_dict(), **self.session_view(design)},
            "capability_report": report, "export": export_result,
            "ready_for_class": False,
            "next_hint": "教案草稿已生成。进入「模拟测试」试讲一遍，通过后即可发布为正式课堂。",
        }

    def export_docx(self, lesson_id: str, project_id: str, design_id: str = "") -> Dict[str, Any]:
        lesson = self.store.get_lesson(lesson_id)
        if lesson is None:
            raise KeyError("Unknown lesson")
        lesson_project_id = str((lesson.metadata or {}).get("project_id") or "")
        if lesson.source != "builtin" and lesson_project_id and lesson_project_id != project_id:
            raise ValueError("课时不属于当前项目")
        if design_id:
            design = self.get(design_id)
            if design.project_id != project_id or (design.final_lesson_id and design.final_lesson_id != lesson_id):
                raise ValueError("教案会话与当前项目或课时不匹配")
        output_dir = self.config.project_output_dir(project_id) / "lesson_plans"
        output_dir.mkdir(parents=True, exist_ok=True)
        path = output_dir / f"lesson_plan_{lesson.lesson_id}_{uuid4().hex[:8]}.docx"
        self._write_docx(path, lesson)
        job = self.store.create_job(project_id=project_id, job_type="lesson_plan_export", title=f"导出教案：{lesson.title}", workflow_type="lesson_plan_export", request={"lesson_id": lesson_id, "design_id": design_id})
        artifact = self.store.register_artifact(project_id=project_id, job_id=job.job_id, artifact_type="lesson_plan_docx", title=f"{lesson.title} 教案", path=str(path), metadata={"lesson_id": lesson_id, "design_id": design_id, "public_url": self.config.public_url_for_path(path)})
        self.store.set_job_status(job.job_id, "success", {"artifact": artifact.to_dict()})
        return {"status": "success", "artifact": artifact.to_dict(), "job_id": job.job_id}

    def _retrieve(self, message: str, owner_user_id: str) -> tuple[str, List[Dict[str, Any]]]:
        lowered = str(message or "").lower()
        local = any(token in lowered for token in ("教材", "课程资料", "知识库", "课标", "项目资料"))
        web = any(token in lowered for token in ("最新", "当前数据", "今年", "核实", "来源", "在线"))
        mode = "local+web" if local and web else "local" if local else "web" if web else "none"
        refs: List[Dict[str, Any]] = []
        if local and self.knowledge_base_service is not None:
            try:
                payload = self.knowledge_base_service.search(query=message, limit=3, owner_user_id=owner_user_id)
                for item in payload.get("items", []):
                    if isinstance(item, dict):
                        refs.append({"id": "kb:" + str(item.get("id") or ""), "title": str(item.get("title") or "项目知识库资料")})
            except Exception:
                pass
        if web and self.resource_search_service is not None:
            try:
                payload = self.resource_search_service.search(query=message, scope="web", limit=3)
                for item in payload.get("items", []):
                    if isinstance(item, dict):
                        refs.append({"id": str(item.get("id") or ""), "title": str(item.get("title") or "在线资料"), "url": str(item.get("url") or "")})
            except Exception:
                pass
        return mode, refs

    def _dataset_facts(self, draft: Dict[str, Any], message: str) -> List[Dict[str, Any]]:
        """Bounded scalar excerpts from referenced local datasets, never geometry or remote fetches."""
        if self.catalog_service is None:
            return []
        identifiers = []
        for stage in draft.get("stages") or []:
            scene = stage.get("scene") if isinstance(stage, dict) else None
            for identifier in scene.get("catalog_layers") or [] if isinstance(scene, dict) else []:
                if isinstance(identifier, str) and identifier not in identifiers:
                    identifiers.append(identifier)
        context = message + json.dumps({k: v for k, v in draft.items() if k != "structured_text_originals"}, ensure_ascii=False)
        excerpts = []
        for identifier in identifiers[:8]:
            try:
                item = self.catalog_service.get_item(identifier)
                path = self.catalog_service.resolve_item_path(item)
                if path.suffix.lower() not in {".json", ".geojson"} or path.stat().st_size > 4_000_000:
                    continue
                data = json.loads(path.read_text(encoding="utf-8"))
                features = data.get("features") if isinstance(data, dict) else None
                if not isinstance(features, list):
                    continue
                records = []
                keys = list(dict.fromkeys(["name", "short_name", "source_year", *item.get("fields", [])]))
                for feature in features:
                    props = feature.get("properties") if isinstance(feature, dict) else None
                    if not isinstance(props, dict) or not any(isinstance(props.get(key), (int, float)) and not isinstance(props.get(key), bool) for key in item.get("fields", []) if key not in {"adcode", "region_code"}):
                        continue
                    record = {key: props[key] for key in keys if isinstance(props.get(key), (str, int, float, bool)) and (not isinstance(props[key], str) or len(props[key]) <= 160)}
                    records.append(record)
                if not records:
                    continue
                records.sort(key=lambda row: not any(isinstance(row.get(key), str) and row[key] and row[key] in context for key in ("name", "short_name")))
                excerpts.append({"dataset_id": identifier, "source_name": item.get("source_name", ""),
                                 "source_url": item.get("source_url", ""), "source_year": item.get("source_year", ""),
                                 "sampled": len(records) > 12, "record_count": len(records), "records": records[:12]})
            except (OSError, ValueError, KeyError, TypeError):
                logger.warning("Lesson dataset excerpt unavailable")
        return excerpts

    def _ask_minimax(self, design: LessonDesignRecord, step: str, message: str) -> Optional[Dict[str, Any]]:
        if self.minimax_client is None:
            return None
        system = (
            "你是高中地理教案共创助手。默认简体中文，每轮只推进一个步骤，先复述教师意图，再给可修改建议，"
            "回复末尾只提一个推进问题。"
            "只输出 JSON，字段为 reply、section_patch、next_step、source_refs、capability_bindings、suggestions。"
            'section_patch 必须是以章节名为键的对象，不是 JSON Patch 操作数组，不需要改动的章节请省略，不要填 null。'
            '格式示例：{"reply":"本轮建议","section_patch":{"stages":[{"stage_id":"s1","title":"环节名","minutes":5,"material":"材料与来源或待补充说明","question_chain":["问题"],"teacher_activities":["教师操作"],"student_activities":["学生任务"],"knowledge_conclusion":"结论","design_intent":"意图","objective_refs":[1],"system_steps":["真实操作"],"scene":{"templates":[],"catalog_layers":[]},"questions":[]}],"board_design":"板书文字"},"next_step":"question_matching","source_refs":[],"capability_bindings":[],"suggestions":[]}。'
            '当前为教学过程时只输出 stages 和 board_design，保持其他章节不变；每个活动字段简洁具体，避免重复长段落。'
            "不要输出内部轨迹。当前步骤：" + STEP_LABELS.get(step, step) +
            "。九个步骤依次是：需求确认→课标与学情→目标与重难点→核心问题与问题链→教学过程→题目匹配→GIS/AI能力→预演检查→确认发布。"
            "核心章节要求：设计思路100-150字；3-4个可观察教学目标；1个核心问题+2-4个递进子问题；"
            "每环节包含 material/question_chain/teacher_activities/student_activities/knowledge_conclusion/"
            "design_intent/minutes/system_steps/objective_refs（1-based目标序号）；板书设计；基础作业+探究作业；预设教学反思。"
            'objectives 必须是字符串数组；core_questions 必须为 {"core":"核心问题文字","sub_questions":["子问题文字"]}，不要把条目写成对象。'
            "教师明确指定的教学环节名称、顺序和每环节分钟数是硬约束，必须逐一原样保留，不能合并或改成通用模板。"
            "你的回复只说明本轮草稿修改，不得声称已发布、已生成文件、预演通过或全部步骤完成；这些状态由系统核验。"
            "题目匹配只能引用题库检索给出的题目，不得编造题目内容。"
            "草稿：" + json.dumps(design.draft, ensure_ascii=False)[:12000] +
            "。真实能力目录：" + json.dumps(self.capability_catalog(), ensure_ascii=False)[:8000] +
            "。当前引用图层的本地统计摘录（只是本地数据，不代表已独立核验来源；sampled=true 时只是部分记录）：" +
            json.dumps(self._dataset_facts(design.draft, message) if step == "process" else [], ensure_ascii=False) +
            "。涉及数字、大小关系或排名时必须与摘录一致；摘录未覆盖或单位未明确时说明需核对，不要补造数值或宣称统计口径相同。"
        )
        messages = [{"role": "system", "content": system}, {"role": "user", "content": message[:6000]}]
        # One bounded correction for model formatting/content errors, never for network failures.
        for attempt in range(2):
            try:
                content = self.minimax_client.chat_completion(
                    messages, temperature=0.2,
                    extra_payload={"max_completion_tokens": 12288 if step == "process" else 2400},
                    timeout=90.0 if step == "process" else 45.0,
                )
            except Exception as exc:
                logger.warning("Lesson generation failed: step=%s error_type=%s", step, type(exc).__name__)
                return None
            reason = "invalid_payload"
            try:
                payload = self._extract_json(content)
                result = self._validate_model_payload(payload)
            except (json.JSONDecodeError, ValueError, TypeError):
                result, reason = None, "invalid_json"
            schedule = self._requested_schedule(message) if step == "process" else []
            if result is not None and schedule and not self._matches_schedule(result["section_patch"].get("stages"), schedule):
                result, reason = None, "schedule_mismatch"
            protected_questions = self._questions_to_preserve(design.draft, message)
            if result is not None and "stages" in result["section_patch"] and not self._preserves_questions(result["section_patch"]["stages"], protected_questions):
                result, reason = None, "protected_question_changed"
            if result is not None:
                return result
            logger.warning("Lesson generation rejected: step=%s reason=%s attempt=%s", step, reason, attempt + 1)
            if attempt == 0:
                # Regenerate from the original teacher request, without replaying malformed output.
                correction = "上次输出未通过系统校验。请重新生成完整 JSON 对象：section_patch 是章节字典，不是数组；不改动的字段省略，不填 null。"
                if schedule:
                    correction += "必须保留这些环节名称、顺序与分钟数：" + json.dumps(schedule, ensure_ascii=False)
                if protected_questions:
                    correction += "这些原题对象必须完整保留且各出现一次，可按教学意图移动到不同环节，不得改写或删减：" + json.dumps(protected_questions, ensure_ascii=False)
                messages.append({"role": "user", "content": correction + "只返回简洁的完整结果，不要解释格式错误。"})
        return None

    @staticmethod
    def _questions_to_preserve(draft: Dict[str, Any], message: str) -> List[Dict[str, Any]]:
        preserve_open = bool(re.search(r"保留.{0,40}(?:题|原文)", message))
        return [copy.deepcopy(question)
                for stage in draft.get("stages") or [] if isinstance(stage, dict)
                for question in stage.get("questions") or [] if isinstance(question, dict)
                if preserve_open or question.get("source") in {"question_bank", "teacher_manual"}]

    @staticmethod
    def _preserves_questions(stages: Any, required: List[Dict[str, Any]]) -> bool:
        if not required:
            return True
        if not isinstance(stages, list):
            return False
        questions = [question for stage in stages if isinstance(stage, dict)
                     for question in stage.get("questions") or [] if isinstance(question, dict)]
        for original in required:
            identifier = original.get("question_id")
            matches = [question for question in questions
                       if question.get("question_id") == identifier] if identifier else [question for question in questions if question == original]
            if len(matches) != 1 or matches[0] != original:
                return False
        return True

    @staticmethod
    def _requested_schedule(message: str) -> List[Dict[str, Any]]:
        """Read explicit named time slots; do not interpret the lesson's total as a stage."""
        content = re.split(r"环节安排(?:为|是)?[：:\s]*", message, maxsplit=1)[-1]
        slots = []
        for part in re.split(r"[、，,；;。\n]", content):
            match = re.fullmatch(r"\s*(.+?)\s*(\d+)\s*分钟\s*", part)
            if not match:
                continue
            title = re.sub(r"^\s*(?:\d+[.．、)]|第[一二三四五六七八九十\d]+环节[：:]?)\s*", "", match[1]).strip()
            if not title or any(word in title for word in ("课时", "总计", "总共", "教案", "设计", "合计")):
                continue
            slots.append({"title": title, "minutes": int(match[2])})
        return slots if len(slots) >= 2 else []

    @staticmethod
    def _matches_schedule(stages: Any, schedule: List[Dict[str, Any]]) -> bool:
        if not isinstance(stages, list) or len(stages) != len(schedule):
            return False
        for stage, slot in zip(stages, schedule):
            if not isinstance(stage, dict) or str(stage.get("title") or "").strip() != slot["title"]:
                return False
            # Numeric strings from a model are tolerated; fractional and boolean values are not.
            if str(stage.get("minutes")) != str(slot["minutes"]):
                return False
        return True

    def _fallback_turn(self, design: LessonDesignRecord, step: str, message: str) -> Dict[str, Any]:
        draft, clean = design.draft, message.strip()
        rename = re.search(r"(?:课题|标题)\s*(?:改为|调整为|改成|是)[:：\s]*(.+)", clean)
        if rename:
            title = rename.group(1).strip("。 ")
            return {
                "reply": f"好的，我把课题改为“{title}”，其他章节暂时保持不变。",
                "section_patch": {"title": title, "topic": title},
                "next_step": step,
                "source_refs": [],
                "capability_bindings": [],
                "suggestions": [],
            }
        if step == "requirements":
            duration_match = re.search(r"(\d+)\s*分钟", clean)
            grade_match = re.search(r"(高[一二三]|初[一二三]|七年级|八年级|九年级)", clean)
            topic = self._extract_topic(clean) or str(draft.get("topic") or draft.get("title") or "人口地理专题课")
            patch = {"title": topic or draft.get("title") or "人口地理专题课", "topic": topic or draft.get("topic") or "人口地理", "grade": grade_match.group(1) if grade_match else draft.get("grade") or "", "duration_minutes": int(duration_match.group(1)) if duration_match else int(draft.get("duration_minutes") or 40), "requirements": {"raw": clean}}
            reply = f"我理解你想做一节“{patch['title']}”。目前先按 {patch['duration_minutes']} 分钟、{patch['grade'] or '年级待定'}来搭框架。接下来我们先确认学生最需要带走的核心认识，可以吗？"
            next_step = "analysis"
        elif step == "analysis":
            patch = {"curriculum_interpretation": clean, "student_analysis": "学生已有基础知识，但需要通过地图和案例把分布格局与成因联系起来。", "textbook_analysis": "围绕人口分布、人口迁移或区域案例组织由图到理的学习任务。"}
            reply, next_step = "我先把课标、学情和教材放在同一条逻辑线上：从地图观察出发，再用案例解释空间差异。你可以直接改写其中任意一段。", "objectives"
        elif step == "objectives":
            items = [line.strip(" -•；;") for line in re.split(r"[；;\n。]", clean) if line.strip()] or ["能够描述人口分布的空间差异", "能够结合地图和案例解释差异成因"]
            patch = {
                "objectives": items[:5],
                "key_difficulties": {"key": [items[0]], "difficult": [items[-1]]},
                "methods": ["地图观察", "案例探究", "问题链"],
                "knowledge_structure": ["地图观察", "空间格局描述", "区域差异解释", "规律迁移"],
            }
            reply, next_step = "我把你的想法整理成了可观察的学习目标，并标出一个重点和一个难点。下一步我们把它们收束成一个核心问题。", "core_questions"
        elif step == "core_questions":
            topic = str(draft.get("topic") or draft.get("title") or "人口分布")
            patch = {
                "core_questions": {
                    "core": f"{topic}的空间格局是怎样形成的，为什么在不同区域存在差异？",
                    "sub_questions": [
                        f"从地图上观察，{topic}呈现出怎样的空间格局？",
                        "哪些自然因素塑造了这种格局？",
                        "人文因素又如何改变或强化了它？",
                    ],
                }
            }
            reply, next_step = "我先把核心问题与三条递进子问题搭出来了（观察描述→自然归因→人文归因）。你可以直接改写核心问题的表述。", "process"
        elif step == "process":
            patch, next_step = {"stages": self._make_stages(draft.get("topic") or "人口分布", draft.get("duration_minutes") or 40), "board_design": self._make_board_design(draft.get("topic") or "人口分布")}, "question_matching"
            reply = "我先给出一版可执行的教学过程：每个环节都写清材料、问题链、教师活动、学生活动、知识结论、设计意图、时间和系统操作。你可以指出要改哪一个环节。"
        elif step == "question_matching":
            report = self.rehearse(design.design_id)
            citations = draft.get("question_citations") or []
            patch: Dict[str, Any] = {}
            homework = draft.get("homework") if isinstance(draft.get("homework"), dict) else {}
            if not [item for item in homework.get("basic") or [] if str(item).strip()]:
                title = str(draft.get("title") or draft.get("topic") or "本课")
                homework["basic"] = [f"完成《{title}》配套基础练习，巩固本课基础知识与读图方法。"]
                if not [item for item in homework.get("inquiry") or [] if str(item).strip()]:
                    homework["inquiry"] = ["选择一个你感兴趣的区域或案例，查阅资料分析其特点，用本课方法说明结论与依据。"]
                patch["homework"] = homework
            if citations:
                summary = "、".join(f"{item.get('number') or item.get('question_id', '')[:12]}" for item in citations[:6])
                reply = f"当前已引用 {len(citations)} 道题（{summary}）。你可以让我替换某道题，或直接补充手动题目。"
            else:
                reply = "我先按环节检索题库候选，答案完备且相关度达标的会自动选用；同时补了一版课后作业草稿。"
            return {"reply": reply, "section_patch": patch, "next_step": "capabilities", "suggestions": report["warnings"][:3]}
        elif step == "capabilities":
            bindings = [{"id": "map_2d", "reason": "展示人口分布空间格局"}, {"id": "knowledge_base", "reason": "必要时核对项目资料"}]
            if any(word in clean for word in ("三维", "3D", "地球")):
                bindings.append({"id": "globe_3d", "reason": "从全球尺度观察人口分布"})
            patch, next_step = {"capabilities": bindings}, "rehearsal"
            reply = "结合现有注册能力，我建议先用 2D 地图完成分布观察；如果你明确需要全球尺度，再加入 3D 地球。当前没有假造数据或图层。"
        elif step == "rehearsal":
            report = self.rehearse(design.design_id)
            reply = "预演检查已完成。" + (f"目前有：{'；'.join(report['errors'])}" if report["errors"] else "主要结构已经齐全。") + " 你可以先接受这一版，或者告诉我希望调整的环节。"
            return {"reply": reply, "section_patch": {}, "next_step": "confirmation", "suggestions": report["warnings"]}
        else:
            patch: Dict[str, Any] = {}
            if not str(draft.get("design_thinking") or "").strip():
                patch["design_thinking"] = self._make_design_thinking(draft)
            if not str(draft.get("reflection") or "").strip():
                patch["reflection"] = "预设反思：关注学生对空间格局的描述是否规范、成因解释是否出现单因素归因；课后依据课堂记录补充。"
            next_step, reply = "confirmation", "我把设计思路与预设反思补了一版草稿，你可以直接改写。确认无误后即可发布为课时草稿（模拟测试通过后进入正式课堂）。"
        return {"reply": reply, "section_patch": patch, "next_step": next_step, "source_refs": [], "capability_bindings": patch.get("capabilities", []), "suggestions": []}

    @staticmethod
    def _extract_topic(message: str) -> str:
        """Extract a short lesson topic from a natural multi-part requirement."""
        clean = str(message or "").strip()
        book_title = re.search(r"《\s*([^》\r\n]{1,80}?)\s*》", clean)
        if book_title:
            return book_title.group(1).strip()
        explicit = re.search(
            r"(?:课题|标题)\s*(?:改为|调整为|改成|为|是)?\s*[:：]?\s*[\"“]?([^，,。；;\n”\"]{1,80})",
            clean,
        )
        if explicit:
            return explicit.group(1).strip()
        simplified = re.sub(r"(高[一二三]|初[一二三]|七年级|八年级|九年级|\d+\s*分钟|单课时)", "", clean)
        simplified = simplified.replace("我想上", "").replace("请设计", "").replace("共创", "")
        parts = [item.strip(" ：:，,。、") for item in re.split(r"[，,。；;\n、]", simplified)]
        parts = [item for item in parts if item and len(item) <= 40]
        return parts[0] if parts else simplified.strip(" ：:，,。、")[:80]

    @classmethod
    def _normalize_requirements_result(
        cls, result: Dict[str, Any], message: str, draft: Dict[str, Any]
    ) -> Dict[str, Any]:
        """Keep model output usable when one sentence contains topic plus constraints."""
        normalized = copy.deepcopy(result) if isinstance(result, dict) else {}
        patch = normalized.get("section_patch")
        patch = patch if isinstance(patch, dict) else {}
        topic = cls._extract_topic(message)
        duration_match = re.search(r"(\d+)\s*分钟", message)
        grade_match = re.search(r"(高[一二三]|初[一二三]|七年级|八年级|九年级)", message)
        original_title = str(patch.get("title") or patch.get("topic") or "").strip()
        if topic:
            patch["title"] = topic
            patch["topic"] = topic
        if duration_match:
            patch["duration_minutes"] = int(duration_match.group(1))
        if grade_match:
            patch["grade"] = grade_match.group(1)
        requirements = patch.get("requirements")
        if not isinstance(requirements, dict):
            requirements = {}
        requirements["raw"] = str(message or "").strip()
        patch["requirements"] = requirements
        normalized["section_patch"] = patch
        if topic and (len(original_title) > 60 or topic not in original_title):
            grade = str(patch.get("grade") or draft.get("grade") or "年级待定")
            duration = int(patch.get("duration_minutes") or draft.get("duration_minutes") or 40)
            normalized["reply"] = (
                f"我已记录：{grade}《{topic}》，{duration} 分钟；其余内容作为学情、证据主线和课堂任务要求保存。"
                "下一步我们确认课标与学情分析，可以吗？"
            )
        return normalized

    @staticmethod
    def _make_stages(topic: str, duration: int) -> List[Dict[str, Any]]:
        """按教研培训模板生成教学过程：材料/问题链/教师活动/学生活动/知识结论/设计意图/时间/系统操作。"""
        first, middle = max(5, min(10, duration // 5)), max(10, min(18, duration // 2))
        last = max(5, duration - first - middle)
        return [
            {
                "stage_id": "s1", "title": "情境导入与地图观察", "minutes": first,
                "knowledge_unit": topic, "knowledge_point": f"{topic}格局",
                "material": f"{topic}专题地图（叠加人口分布图层）",
                "question_chain": [f"{topic}在空间上呈现怎样的差异？", "这种差异是否具有稳定方向？"],
                "teacher_activities": ["展示专题地图并引导读图", "追问差异的方向性"],
                "student_activities": ["学生观察地图并圈出集中区与稀疏区", "用自己的语言描述空间格局"],
                "knowledge_conclusion": f"{topic}分布不均衡，呈现明显的空间集中与稀疏格局。",
                "design_intent": "从可观察的空间证据进入问题。", "objective_refs": [1],
                "system_steps": ["打开2D地图", "叠加已注册的人口专题图层"],
                "scene": {"templates": ["population_distribution"], "catalog_layers": []},
                "script": [],
                "questions": [{"question_id": "s1q1", "type": "open", "text": f"{topic}在空间上呈现怎样的差异？", "options": [], "answer_index": None, "expected_points": [], "misconceptions": []}],
                "assistant_prompts": [],
            },
            {
                "stage_id": "s2", "title": "案例探究与成因解释", "minutes": middle,
                "knowledge_unit": topic, "knowledge_point": f"{topic}的影响因素",
                "material": "两个典型区域的对比案例（图文材料）",
                "question_chain": ["两个区域的条件有何不同？", "哪些因素导致了人口集聚/稀疏？"],
                "teacher_activities": ["组织小组对比", "提供证据材料并追问因果"],
                "student_activities": ["小组比较两个区域", "用因果链说明自然与社会因素"],
                "knowledge_conclusion": "自然条件提供基础，社会经济因素强化或改变人口分布格局。",
                "design_intent": "让学生完成由描述到解释的认知跃迁。", "objective_refs": [2],
                "system_steps": ["使用地图标注", "调用已注册的案例或知识库资料"],
                "scene": {"templates": [], "catalog_layers": []},
                "script": [],
                "questions": [{"question_id": "s2q1", "type": "open", "text": "哪些自然和社会条件共同造成了这种区域差异？", "options": [], "answer_index": None, "expected_points": ["自然条件", "社会经济条件", "因果联系"], "misconceptions": []}],
                "assistant_prompts": [],
            },
            {
                "stage_id": "s3", "title": "归纳迁移与课堂小结", "minutes": last,
                "knowledge_unit": topic, "knowledge_point": f"{topic}规律与迁移",
                "material": "本节课生成的地图标注与板书要点",
                "question_chain": ["用一句话概括本课规律？", "换一个区域你会如何分析？"],
                "teacher_activities": ["根据回答收束规律", "布置迁移任务"],
                "student_activities": ["个人完成结论卡片", "口头迁移到新区域"],
                "knowledge_conclusion": "读图描述格局 → 归因自然与人文因素 → 迁移到新区域的分析方法。",
                "design_intent": "检查目标达成并形成可迁移的方法。", "objective_refs": [3],
                "system_steps": ["课堂记录学生回答", "生成课后复习提示"],
                "scene": {"templates": [], "catalog_layers": []},
                "script": [],
                "questions": [{"question_id": "s3q1", "type": "open", "text": "换到另一个区域时，你会按什么顺序完成读图与解释？", "options": [], "answer_index": None, "expected_points": ["读图例", "描述格局", "解释原因"], "misconceptions": []}],
                "assistant_prompts": [],
            },
        ]

    @staticmethod
    def _make_board_design(topic: str) -> str:
        return (
            f"【{topic}】\n"
            "左：空间格局（学生圈画）｜中：成因链（自然←→人文）｜右：方法（读图→归因→迁移）"
        )

    @staticmethod
    def _make_design_thinking(draft: Dict[str, Any]) -> str:
        topic = str(draft.get("topic") or draft.get("title") or "人口地理")
        objectives = [str(item)[:20] for item in (draft.get("objectives") or []) if str(item).strip()]
        methods = [str(item) for item in (draft.get("methods") or []) if str(item).strip()] or ["地图观察", "问题链"]
        difficult = [str(item)[:20] for item in ((draft.get("key_difficulties") or {}).get("difficult") or [])[:2]]
        core = str((draft.get("core_questions") or {}).get("core") or f"{topic}的空间格局及其成因")[:40]
        objective_text = "、".join(objectives[:3]) if objectives else "空间格局观察与成因解释"
        parts = [
            f"本课围绕核心问题“{core}”展开。",
            f"以{objective_text}为目标，",
            f"通过{'、'.join(methods[:3])}组织学习，",
            f"重点突破{'、'.join(difficult) if difficult else '成因解释'}，"
            "让学生在地图证据与真实真题情境中完成由描述到解释再到迁移的思维进阶。",
        ]
        text = "".join(parts)
        if len(text) > 150:
            text = text[:147] + "……"
        return text

    @staticmethod
    def _section_has_content(value: Any) -> bool:
        if isinstance(value, dict):
            return any(LessonDesignService._section_has_content(item) for item in value.values())
        if isinstance(value, (list, tuple, set)):
            return any(LessonDesignService._section_has_content(item) for item in value)
        return bool(str(value or "").strip())

    @staticmethod
    def _build_diff_summary(design: LessonDesignRecord) -> List[Dict[str, Any]]:
        if not design.base_draft:
            return []
        result: List[Dict[str, Any]] = []
        for key in SECTION_KEYS:
            before = design.base_draft.get(key)
            after = design.draft.get(key)
            if before != after:
                result.append({
                    "section": key,
                    "label": SECTION_LABELS.get(key, key),
                    "changed": True,
                })
        return result

    @staticmethod
    def _validate_model_payload(payload: Any) -> Optional[Dict[str, Any]]:
        if not isinstance(payload, dict):
            logger.warning("Lesson payload shape: root_type=%s", type(payload).__name__)
            return None
        patch = payload.get("section_patch", {})
        if not isinstance(patch, dict):
            logger.warning("Lesson payload shape: patch_type=%s", type(patch).__name__)
            return None
        allowed = set(SECTION_KEYS) | {"title", "subject", "grade", "topic", "duration_minutes"}
        normalized_patch = {str(key): copy.deepcopy(value) for key, value in patch.items() if str(key) in allowed}
        try:
            LessonDesignService._normalize_text_lists(copy.deepcopy(normalized_patch))
        except ValueError:
            core = normalized_patch.get("core_questions")
            logger.warning("Lesson payload shape: objectives_type=%s sub_questions_type=%s", type(normalized_patch.get("objectives")).__name__, type(core.get("sub_questions") if isinstance(core, dict) else None).__name__)
            return None
        next_step = str(payload.get("next_step") or "")
        source_refs = payload.get("source_refs") if isinstance(payload.get("source_refs"), list) else []
        bindings = payload.get("capability_bindings") if isinstance(payload.get("capability_bindings"), list) else []
        suggestions = payload.get("suggestions") if isinstance(payload.get("suggestions"), list) else []
        return {
            "reply": str(payload.get("reply") or "").strip(),
            "section_patch": normalized_patch,
            "next_step": next_step if next_step in STEP_KEYS else "",
            "source_refs": [item for item in source_refs if isinstance(item, dict)],
            "capability_bindings": [item for item in bindings if isinstance(item, dict)],
            "suggestions": [str(item) for item in suggestions if str(item).strip()],
        }

    @staticmethod
    def _merge_patch(draft: Dict[str, Any], patch: Dict[str, Any]) -> None:
        allowed = set(SECTION_KEYS) | {"title", "subject", "grade", "topic", "duration_minutes"}
        for key, value in patch.items():
            if key not in allowed:
                continue
            if isinstance(value, dict) and isinstance(draft.get(key), dict):
                draft[key] = {**draft[key], **copy.deepcopy(value)}
            else:
                draft[key] = copy.deepcopy(value)
        LessonDesignService._normalize_text_lists(draft)

    @staticmethod
    def _next_step(step: str) -> str:
        try:
            return STEP_KEYS[min(STEP_KEYS.index(step) + 1, len(STEP_KEYS) - 1)]
        except ValueError:
            return STEP_KEYS[0]

    @staticmethod
    def _step_for_section(section_id: str) -> str:
        return {"requirements": "analysis", "objectives": "process", "stages": "capabilities"}.get(section_id, "confirmation")

    @staticmethod
    def _natural_prompt(step: str) -> str:
        return {
            "requirements": "先告诉我年级、课题和大致课时，我会帮你搭起第一版框架。",
            "analysis": "接下来一起确认课标重点、学生基础和教材位置。",
            "objectives": "这节课最希望学生学会什么？可以先说一两点。",
            "core_questions": "现在把目标收束成一个核心问题，再拆成 2-4 个递进子问题。",
            "process": "我们把目标落到具体环节，每轮只设计一个关键环节。",
            "question_matching": "接下来从题库里匹配真题；答案不完备的题目我不会自动选用。",
            "capabilities": "最后核对哪些地图、数据和课堂记录能力真的可用。",
            "rehearsal": "现在可以做一次时长、问题链、题目和能力检查。",
            "confirmation": "补上设计思路与预设反思后，就可以确认发布并导出 Word。",
        }.get(step, "我们从课题、年级和课时开始吧。")

    def _validate_bindings(self, bindings: Any) -> List[Dict[str, Any]]:
        catalog = {item["id"]: item for item in self.capability_catalog()}
        result = []
        for item in bindings if isinstance(bindings, list) else []:
            if isinstance(item, str):
                item = {"id": item}
            if not isinstance(item, dict):
                continue
            capability_id = str(item.get("id") or item.get("capability_id") or "")
            if capability_id in catalog:
                result.append({"id": capability_id, "label": catalog[capability_id]["label"], "reason": str(item.get("reason") or "")})
        return result

    @staticmethod
    def _merge_refs(existing: List[Dict[str, Any]], refs: Any) -> List[Dict[str, Any]]:
        result, seen = list(existing), {str(item.get("id") or item.get("source_id") or "") for item in existing}
        for item in refs if isinstance(refs, list) else []:
            if isinstance(item, str):
                item = {"id": item, "title": item}
            if not isinstance(item, dict):
                continue
            key = str(item.get("id") or item.get("source_id") or item.get("title") or "")
            if key and key not in seen:
                result.append({str(k): v for k, v in item.items() if k in {"id", "source_id", "dataset_id", "title", "year", "source_year", "url"}})
                seen.add(key)
        return result

    @staticmethod
    def _extract_json(content: str) -> Any:
        cleaned = str(content or "").strip()
        fence = chr(96) * 3
        if cleaned.startswith(fence):
            cleaned = re.sub(r"^" + re.escape(fence) + r"(?:json)?", "", cleaned, flags=re.I).strip()
            cleaned = re.sub(re.escape(fence) + r"$", "", cleaned).strip()
        try:
            return json.loads(cleaned)
        except json.JSONDecodeError:
            start, end = cleaned.find("{"), cleaned.rfind("}")
            if start >= 0 and end > start:
                return json.loads(cleaned[start:end + 1])
            raise

    @staticmethod
    def _set_cell(cell: Any, value: Any, size: int = 9) -> None:
        from docx.shared import Pt
        from docx.oxml.ns import qn
        cell.text = str(value or "")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.name = "宋体"
                if run._element.rPr is not None:
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                run.font.size = Pt(size)

    def _write_docx(self, path: Path, lesson: LessonRecord) -> None:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt

        template_path = (
            self.config.root_dir
            / "backend"
            / "app"
            / "data"
            / "builtin"
            / "lesson_plan_template"
            / "lesson_plan_template.docx"
        )
        doc = Document(str(template_path)) if template_path.is_file() else Document()
        body = doc._element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        section = doc.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.5)
        section.left_margin = section.right_margin = Inches(0.5)
        normal = doc.styles["Normal"]
        normal.font.name = "宋体"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        normal.font.size = Pt(10)
        doc.core_properties.title = lesson.title
        doc.core_properties.subject = "WebGIS-AI 教案导出"
        doc.core_properties.author = "WebGIS-AI"
        doc.core_properties.last_modified_by = "WebGIS-AI"

        def set_table_geometry(table: Any, widths: list[int]) -> None:
            table.autofit = False
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table_properties = table._tbl.tblPr
            layout = table_properties.find(qn("w:tblLayout"))
            if layout is None:
                layout = OxmlElement("w:tblLayout")
                table_properties.append(layout)
            layout.set(qn("w:type"), "fixed")
            table_width = table_properties.find(qn("w:tblW"))
            if table_width is None:
                table_width = OxmlElement("w:tblW")
                table_properties.append(table_width)
            table_width.set(qn("w:type"), "dxa")
            table_width.set(qn("w:w"), str(sum(widths)))
            grid = table._tbl.tblGrid
            for node in list(grid):
                grid.remove(node)
            for width in widths:
                grid_col = OxmlElement("w:gridCol")
                grid_col.set(qn("w:w"), str(width))
                grid.append(grid_col)
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.width = width
                    cell_properties = cell._tc.get_or_add_tcPr()
                    cell_width = cell_properties.find(qn("w:tcW"))
                    if cell_width is None:
                        cell_width = OxmlElement("w:tcW")
                        cell_properties.append(cell_width)
                    cell_width.set(qn("w:type"), "dxa")
                    cell_width.set(qn("w:w"), str(width))
                    margins = cell_properties.find(qn("w:tcMar"))
                    if margins is None:
                        margins = OxmlElement("w:tcMar")
                        cell_properties.append(margins)
                    for side, value in (("top", 80), ("start", 100), ("bottom", 80), ("end", 100)):
                        margin = margins.find(qn(f"w:{side}"))
                        if margin is None:
                            margin = OxmlElement(f"w:{side}")
                            margins.append(margin)
                        margin.set(qn("w:w"), str(value))
                        margin.set(qn("w:type"), "dxa")

        def mark_header(row: Any) -> None:
            row_properties = row._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            row_properties.append(header)
            for cell in row.cells:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "D9EAF7")
                cell._tc.get_or_add_tcPr().append(shade)
                for run in cell.paragraphs[0].runs:
                    run.bold = True

        title_style = "Lesson Title" if "Lesson Title" in doc.styles else None
        title = doc.add_paragraph(style=title_style)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.add_run(lesson.title).bold = True
        plan = lesson.plan or {}
        info = doc.add_table(rows=4, cols=4)
        info.style = "Table Grid"
        rows = [
            ("学科", lesson.subject, "年级", lesson.grade),
            ("课题", lesson.title, "课时", f"{plan.get('duration_minutes', lesson.metadata.get('duration_minutes', 40))}分钟"),
            ("课型", plan.get("lesson_type", "专题探究课"), "平台", "WebGIS-AI"),
            ("设计来源", "教师与智能体共创", "版本", "当前草稿"),
        ]
        for row, values in zip(info.rows, rows):
            for cell, value in zip(row.cells, values):
                self._set_cell(cell, value)
        set_table_geometry(info, [1200, 3900, 1200, 4169])
        for row in info.rows:
            for index in (0, 2):
                for run in row.cells[index].paragraphs[0].runs:
                    run.bold = True

        heading_style = "Lesson Heading" if "Lesson Heading" in doc.styles else "Heading 2"
        doc.add_paragraph("一、课标与前置分析", style=heading_style)
        for heading, text in (("课标解读", plan.get("curriculum_interpretation", "")), ("学情分析", plan.get("student_analysis", "")), ("教材分析", plan.get("textbook_analysis", ""))):
            p = doc.add_paragraph()
            p.add_run(heading + "：").bold = True
            p.add_run(str(text or "待教师补充。"))
        doc.add_paragraph("二、教学目标与重难点", style=heading_style)
        for objective in plan.get("objectives") or lesson.objectives:
            doc.add_paragraph(str(objective), style="List Bullet")
        difficulties = plan.get("key_difficulties") or {}
        p = doc.add_paragraph()
        p.add_run("重点：").bold = True
        p.add_run("；".join(str(x) for x in difficulties.get("key") or []) or "待确认")
        p.add_run("    难点：").bold = True
        p.add_run("；".join(str(x) for x in difficulties.get("difficult") or []) or "待确认")
        methods = plan.get("methods") or []
        if methods:
            p = doc.add_paragraph()
            p.add_run("教学方法：").bold = True
            p.add_run("；".join(str(item) for item in methods))
        doc.add_paragraph("三、知识结构", style=heading_style)
        knowledge = plan.get("knowledge_structure") or []
        structure = " → ".join(str(item) for item in knowledge) if knowledge else f"{lesson.title} → 空间分布观察 → 成因解释 → 规律迁移"
        structure_paragraph = doc.add_paragraph(structure)
        structure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("四、教学过程", style=heading_style)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for cell, header in zip(table.rows[0].cells, ["环节", "知识单元", "知识点", "具体内容/活动", "设计意图"]):
            self._set_cell(cell, header)
        mark_header(table.rows[0])
        for stage in lesson.stages:
            row = table.add_row()
            activities = stage.get("activities") or []
            content = stage.get("content") or "；".join(str(x) for x in activities) or "待补充"
            if stage.get("system_steps"):
                content += "\n系统操作：" + "；".join(str(x) for x in stage["system_steps"])
            values = [
                f"{stage.get('title', '')}（{stage.get('minutes', 0)}分钟）", stage.get("knowledge_unit", ""),
                stage.get("knowledge_point", ""), content, stage.get("design_intent", ""),
            ]
            for cell, value in zip(row.cells, values):
                self._set_cell(cell, value)
        set_table_geometry(table, [1300, 1500, 1500, 3569, 2600])
        doc.add_paragraph("五、参考资料与教学反思", style=heading_style)
        refs = plan.get("references") or []
        for ref in refs or ["本教案未强制附加默认来源；需要核验时由教师指定资料。"]:
            doc.add_paragraph(str(ref.get("title") if isinstance(ref, dict) else ref), style="List Bullet")
        doc.add_paragraph("教学反思：" + str(plan.get("reflection") or "课后根据课堂记录补充。"))
        doc.save(path)
