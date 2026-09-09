"""After-class practice paper export: student paper + teacher paper (DOCX).

Question selection priority (each level falls back honestly and the source is
labelled on the paper):
1. Homework assigned in the published lesson plan (homework basic/inquiry tasks);
2. Classroom observation marks: re-drill the original question for partial /
   misconception verdicts, and retrieve bank variants for misconception tags;
3. Bank questions strongly related to the lesson's core objectives.

Evidence discipline: classroom measurements printed on the teacher paper come
only from real events/responses/observations. Questions without answer data are
labelled "未收集到作答数据" instead of an invented 0% rate. The student paper is
generated from a body writer that never reads answer/explanation fields, so
answers cannot leak into it.
"""

from __future__ import annotations

import re
import hashlib
import json
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from ..models import ClassSessionRecord, LessonRecord
from .question_bank import QuestionBankService
from .lesson_homework import homework_guidance

ORIGIN_LABELS = {
    "lesson_homework_basic": "教案课后作业（基础）",
    "lesson_homework_inquiry": "教案课后作业（探究）",
    "class_observation": "教师课堂速记（原题回炉）",
    "observation_variant": "误区变式（题库检索）",
    "bank_core": "核心目标巩固（题库检索）",
}

ORIGIN_LEVELS = {
    "lesson_homework_basic": "基础必做",
    "lesson_homework_inquiry": "拓展选做",
    "class_observation": "课堂巩固",
    "observation_variant": "课堂巩固",
    "bank_core": "课后巩固",
}

VERDICT_LABELS = {"partial": "部分掌握", "misconception": "存在误区"}

VARIANT_TAG_LIMIT = 3
VARIANT_LIMIT = 4
CORE_LIMIT = 6


class PracticeExportService:
    def __init__(self, config: Any, store: Any, question_bank_service: Optional[QuestionBankService] = None):
        self.config = config
        self.store = store
        self.question_bank = question_bank_service

    # ------------------------------------------------------------------
    # Export entry (mirrors lesson_design.export_docx: job + artifacts)
    # ------------------------------------------------------------------

    def export(self, session: ClassSessionRecord, lesson: Optional[LessonRecord],
               selection: Optional[Dict[str, Any]] = None, job_id: Optional[str] = None) -> Dict[str, Any]:
        lesson_title = str(
            (lesson.title if lesson else "") or (session.metadata or {}).get("lesson_title") or "本课"
        )
        items, summary, notes = self.collect_items(session, lesson)
        if selection is not None:
            manifest = self.selection_manifest(session, items)
            selected = selection.get("selected_ids")
            if not isinstance(selected, list) or not selected or any(not isinstance(v, str) for v in selected):
                raise ValueError("请至少勾选一道题目或一项作业；未生成空白试卷。")
            if selection.get("token") != manifest["token"]:
                raise ValueError("候选内容或课堂记录已变化，请重新生成报告并确认选题。")
            if len(set(selected)) != len(selected) or not set(selected).issubset(manifest["item_ids"]):
                raise ValueError("选题包含重复或不可用内容，请重新确认。")
            items = [item for item in items if item["practice_id"] in selected]
            summary = self._selection_summary(items)
            notes.append(f"本卷包含教师勾选的 {len(items)} 项作业与题目。")

        if not items:
            raise ValueError("未找到可用作业内容。请检查当前项目题库、课时目标或先添加作业任务；未生成空白试卷。")

        job = self.store.get_job(job_id) if job_id else None
        if job is None:
            job = self.store.create_job(
                project_id=session.project_id,
                job_type="practice_export",
                title=f"导出课后练习卷：{lesson_title}",
                workflow_type="practice_export",
                request={"session_id": session.session_id, "selection": selection},
            )
        self.store.set_job_status(job.job_id, "running")
        try:
            output_dir = self.config.project_output_dir(session.project_id)
            student_path = self.config.unique_path(output_dir, f"practice_student_{session.session_id[:12]}.docx")
            teacher_path = self.config.unique_path(output_dir, f"practice_teacher_{session.session_id[:12]}.docx")
            self._write_paper(student_path, lesson_title, items, summary, notes, teacher=False)
            self._write_paper(teacher_path, lesson_title, items, summary, notes, teacher=True)

            student_artifact = self.store.register_artifact(
                project_id=session.project_id,
                job_id=job.job_id,
                artifact_type="practice_paper_student",
                title=f"{lesson_title} 课后练习卷（学生卷）",
                path=str(student_path),
                metadata={
                    "public_url": self.config.public_url_for_path(student_path),
                    "session_id": session.session_id,
                    "format": "docx",
                },
            )
            teacher_artifact = self.store.register_artifact(
                project_id=session.project_id,
                job_id=job.job_id,
                artifact_type="practice_paper_teacher",
                title=f"{lesson_title} 课后练习卷（教师卷）",
                path=str(teacher_path),
                metadata={
                    "public_url": self.config.public_url_for_path(teacher_path),
                    "session_id": session.session_id,
                    "format": "docx",
                },
            )
            result = {
                "status": "success",
                "job_id": job.job_id,
                "session_id": session.session_id,
                "student_artifact": student_artifact.to_dict(),
                "teacher_artifact": teacher_artifact.to_dict(),
                "selection_summary": summary,
                "selected_ids": [item["practice_id"] for item in items],
                "notes": notes,
                "selection_token": selection.get("token") if selection else None,
            }
            self.store.set_job_status(job.job_id, "success", result)
            return result
        except Exception as exc:
            self.store.set_job_status(job.job_id, "failed", error=str(exc))
            raise

    # ------------------------------------------------------------------
    # Selection (priority levels with honest fallback)
    # ------------------------------------------------------------------

    def collect_items(
        self, session: ClassSessionRecord, lesson: Optional[LessonRecord]
    ) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], List[str]]:
        plan = dict((lesson.plan if lesson is not None else {}) or {})
        # Built-in/legacy lessons may store goals at the top level, without a plan.
        if lesson is not None:
            plan["title"] = plan.get("title") or lesson.title
            plan["topic"] = plan.get("topic") or lesson.title
            plan["objectives"] = plan.get("objectives") or list(lesson.objectives or [])
        lesson_questions = self._lesson_question_index(lesson)
        evidence = self._classroom_evidence(session)
        observations = self._practice_observations(session)

        items: List[Dict[str, Any]] = []
        included_ids: set = set()
        skipped: List[str] = []

        # ① 教案指定的课后作业（homework basic/inquiry 文本任务）
        homework = plan.get("homework") if isinstance(plan.get("homework"), dict) else {}
        for origin in ("lesson_homework_basic", "lesson_homework_inquiry"):
            for index, text in enumerate(homework.get("basic" if origin == "lesson_homework_basic" else "inquiry") or [], start=1):
                text = str(text).strip()
                if text:
                    items.append({"kind": "task", "origin": origin, "text": text,
                                  "practice_id": f"{origin}_{index}",
                                  "teacher_guidance": homework_guidance(homework, text)})

        # ② 教师课堂标注的部分掌握/误区：先回炉原题，再按误区标签与知识点检索变式题
        variant_queries: List[str] = []
        for observation in observations:
            question_id = observation["question_id"]
            if question_id and question_id in lesson_questions and question_id not in included_ids:
                included_ids.add(question_id)
                items.append(
                    self._question_item(
                        "class_observation",
                        lesson_questions[question_id]["question"],
                        str(lesson_questions[question_id].get("stage_title") or ""),
                        session,
                        observation=observation,
                        evidence=evidence,
                    )
                )
                for point in lesson_questions[question_id]["question"].get("knowledge_points") or []:
                    point = str(point).strip()
                    if point and point not in variant_queries:
                        variant_queries.append(point)
            if observation["tag"] and observation["tag"] not in variant_queries:
                variant_queries.append(observation["tag"])

        for snapshot in self._search_bank(
            session, plan, knowledge_queries=variant_queries[:VARIANT_TAG_LIMIT], limit=VARIANT_LIMIT if variant_queries else 0,
            exclude_ids=included_ids, skipped=skipped,
        ):
            included_ids.add(snapshot["question_id"])
            items.append(
                self._question_item(
                    "observation_variant", snapshot, "", session, evidence=evidence,
                    selection_reason=str(snapshot.get("selection_reason") or ""),
                )
            )

        # ③ 围绕课时核心目标从题库检索强关联题
        for snapshot in self._search_bank(
            session, plan, knowledge_queries=[], limit=CORE_LIMIT, exclude_ids=included_ids, skipped=skipped
        ):
            included_ids.add(snapshot["question_id"])
            items.append(
                self._question_item(
                    "bank_core", snapshot, "", session, evidence=evidence,
                    selection_reason=str(snapshot.get("selection_reason") or ""),
                )
            )

        summary = self._selection_summary(items)
        notes = self._selection_notes(items, lesson, session)
        if skipped:
            notes.append("自动选题已跳过：" + "；".join(dict.fromkeys(skipped)) + "。未为凑题数补入不完整题目。")
        return items, summary, notes

    def report_bank_recommendations(
        self, session: ClassSessionRecord, lesson: Optional[LessonRecord]
    ) -> Tuple[List[Dict[str, Any]], List[str], Dict[str, Any]]:
        items, _, notes = self.collect_items(session, lesson)
        result = []
        for item in items:
            if item["origin"] not in {"bank_core", "observation_variant"}:
                continue
            q = item["question"]
            source = " · ".join(dict.fromkeys(str(q.get(k) or "").strip()
                                  for k in ("year", "region", "source_paper") if q.get(k)))
            answer = q.get("answer") or q.get("answer_letter") or ""
            points = [str(value) for value in (answer, q.get("explanation")) if value]
            for sub in q.get("sub_questions") or []:
                details = "；".join(str(value) for value in (sub.get("answer"), sub.get("explanation")) if value)
                if details:
                    points.append(f"({sub['index']}) {details}")
            basis = ("依据教师速记对应考点检索，需核对具体观察；不代表全班存在同一误区。"
                     if item["origin"] == "observation_variant" else
                     "按本次教案主题与目标进行本地题库匹配；作为巩固候选，不代表学生答错。")
            result.append({
                "practice_id": "bank_" + q["question_id"], "level": ORIGIN_LEVELS[item["origin"]],
                "title": source or "题库巩固题", "suggested_minutes": None,
                "prompt": q["text"] or q["task_text"], "answer_points": points,
                "evidence_basis": basis, "question": q,
            })
        notes.append("题库候选与练习卷使用同一套本地选题规则；教案、题库或课堂记录变化后请重新生成复盘。")
        return result, notes, self.selection_manifest(session, items)

    @staticmethod
    def selection_manifest(session: ClassSessionRecord, items: List[Dict[str, Any]]) -> Dict[str, Any]:
        payload = json.dumps({"session_id": session.session_id, "project_id": session.project_id, "items": items},
                             sort_keys=True, ensure_ascii=False, separators=(",", ":"))
        return {"token": hashlib.sha256(payload.encode("utf-8")).hexdigest(),
                "item_ids": [item["practice_id"] for item in items]}

    @staticmethod
    def _lesson_question_index(lesson: Optional[LessonRecord]) -> Dict[str, Dict[str, Any]]:
        index: Dict[str, Dict[str, Any]] = {}
        if lesson is None:
            return index
        for stage in lesson.stages or []:
            if not isinstance(stage, dict):
                continue
            for question in stage.get("questions") or []:
                if isinstance(question, dict) and str(question.get("question_id") or ""):
                    index[str(question["question_id"])] = {
                        "question": question,
                        "stage_title": str(stage.get("title") or ""),
                    }
        return index

    @staticmethod
    def _practice_observations(session: ClassSessionRecord) -> List[Dict[str, str]]:
        result: List[Dict[str, str]] = []
        for event in session.events:
            if event.get("type") != "teacher_observation":
                continue
            payload = event.get("payload") or {}
            verdict = str(payload.get("verdict") or "")
            if verdict not in {"partial", "misconception"}:
                continue
            result.append(
                {
                    "question_id": str(payload.get("question_id") or ""),
                    "verdict": verdict,
                    "tag": str(payload.get("tag") or "").strip(),
                    "note": str(payload.get("note") or "").strip(),
                }
            )
        return result

    @staticmethod
    def _classroom_evidence(session: ClassSessionRecord) -> Dict[str, Dict[str, Any]]:
        """从真实事件聚合每道题的课堂实测（绝不推断缺失字段）。"""
        evidence: Dict[str, Dict[str, Any]] = {}
        for event in session.events:
            event_type = str(event.get("type") or "")
            payload = event.get("payload") or {}
            question_id = str(payload.get("question_id") or "")
            if not question_id:
                continue
            entry = evidence.setdefault(
                question_id,
                {
                    "launched": False,
                    "response_count": None,
                    "correct_rate": None,
                    "revealed": None,
                    "actual_seconds": None,
                    "overtime_seconds": None,
                },
            )
            if event_type in {"question_launched", "teacher_question_presented"}:
                entry["launched"] = True
            elif event_type in {"question_revealed", "question_closed"}:
                if "revealed" in payload:
                    entry["revealed"] = bool(payload.get("revealed"))
                if payload.get("actual_seconds") is not None:
                    entry["actual_seconds"] = int(payload.get("actual_seconds"))
                if payload.get("overtime_seconds") is not None:
                    entry["overtime_seconds"] = int(payload.get("overtime_seconds"))
        for question_id, responses in (session.responses or {}).items():
            entry = evidence.setdefault(
                str(question_id),
                {
                    "launched": False,
                    "response_count": None,
                    "correct_rate": None,
                    "revealed": None,
                    "actual_seconds": None,
                    "overtime_seconds": None,
                },
            )
            entry["launched"] = True
            entry["response_count"] = len(responses or [])
        return evidence

    def _question_item(
        self,
        origin: str,
        question: Dict[str, Any],
        stage_title: str,
        session: ClassSessionRecord,
        observation: Optional[Dict[str, str]] = None,
        evidence: Optional[Dict[str, Dict[str, Any]]] = None,
        selection_reason: str = "",
    ) -> Dict[str, Any]:
        snapshot = QuestionBankService.normalize_snapshot_question(question)
        question_id = snapshot["question_id"]
        entry = dict((evidence or {}).get(question_id) or {})
        response_count = entry.get("response_count")
        if response_count is None:
            response_count = len((session.responses or {}).get(question_id) or [])
        entry["response_count"] = int(response_count or 0)
        entry["correct_rate"] = self._correct_rate(snapshot, session.responses.get(question_id) or [])
        return {
            "kind": "question",
            "practice_id": f"{'bank' if origin in {'bank_core', 'observation_variant'} else 'class_question'}_{question_id}",
            "origin": origin,
            "question": snapshot,
            "stage_title": stage_title,
            "classroom": entry,
            "observation": dict(observation) if observation else None,
            "selection_reason": selection_reason,
        }

    @staticmethod
    def _correct_rate(snapshot: Dict[str, Any], responses: List[Dict[str, Any]]) -> Optional[float]:
        """正确率只由真实作答与官方答案索引计算；条件不足时返回 None。"""
        answer_index = snapshot.get("answer_index")
        options = snapshot.get("options") or []
        if not isinstance(answer_index, int) or not (0 <= answer_index < len(options)) or not responses:
            return None
        correct = sum(
            1 for item in responses if isinstance(item.get("choice_index"), int) and item["choice_index"] == answer_index
        )
        return round(correct / len(responses), 4)

    def _search_bank(
        self,
        session: ClassSessionRecord,
        plan: Dict[str, Any],
        knowledge_queries: List[str],
        limit: int,
        exclude_ids: set,
        skipped: Optional[List[str]] = None,
    ) -> List[Dict[str, Any]]:
        if self.question_bank is None or limit <= 0:
            return []
        topic = str(plan.get("topic") or plan.get("title") or "")
        if "人口分布" in topic:
            topic = "人口分布"
        core = (plan.get("core_questions") or {}) if isinstance(plan.get("core_questions"), dict) else {}
        core_question = str(core.get("core") or "")
        objectives = [str(item) for item in plan.get("objectives") or [] if str(item).strip()]
        found: List[Dict[str, Any]] = []
        excluded = set(exclude_ids)
        queries = list(knowledge_queries)
        if not queries:
            # Long prose diluted the topic score and rejected even direct matches.
            # Objectives are already passed separately to the bank ranker.
            queries.append(topic or core_question)
        for query in queries:
            if not query.strip() or len(found) >= limit:
                continue
            try:
                result = self.question_bank.search(
                    project_id=session.project_id,
                    topic=topic,
                    knowledge=query,
                    objectives=objectives,
                    exclude_ids=sorted(excluded),
                    limit=10,
                    use_llm=False,  # Report preview and paper use the same reproducible local ranking.
                )
            except Exception as exc:
                raise ValueError("题库检索失败，请重试；未将失败当作无题库或生成空卷。") from exc
            for item in result.get("items") or []:
                if not isinstance(item, dict):
                    continue
                snapshot = QuestionBankService.normalize_snapshot_question(item)
                question_id = snapshot["question_id"]
                if not question_id or question_id in excluded:
                    continue
                reason = ""
                stem = snapshot["text"] + " " + snapshot["task_text"]
                if not snapshot.get("answer_complete"):
                    reason = "答案不完整"
                elif item.get("auto_selectable") is False:
                    reason = "关联度不足"
                elif "人口分布" in topic and not re.search(r"人口|人类.{0,5}居住|聚落", stem):
                    reason = "题干未直接考查人口分布"
                elif re.search(r"图示|图中|下图|如图|图为|读图", stem + snapshot["material"]) and not snapshot["images"]:
                    reason = "读图题缺少题图"
                elif any(self._resolve_image_path(image["url"]) is None for image in snapshot["images"]):
                    reason = "题图文件不可用"
                if reason:
                    if skipped is not None:
                        skipped.append(reason)
                    excluded.add(question_id)
                    continue
                snapshot["selection_reason"] = str(item.get("selection_reason") or "")
                found.append(snapshot)
                excluded.add(question_id)
                if len(found) >= limit:
                    break
        return found

    @staticmethod
    def _selection_summary(items: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        counts: Dict[str, int] = {origin: 0 for origin in ORIGIN_LABELS}
        for item in items:
            counts[item["origin"]] += 1
        return [
            {"origin": origin, "label": ORIGIN_LABELS[origin], "level": ORIGIN_LEVELS[origin], "count": count}
            for origin, count in counts.items()
        ]

    @staticmethod
    def _selection_notes(items: List[Dict[str, Any]], lesson: Optional[LessonRecord], session: ClassSessionRecord) -> List[str]:
        notes: List[str] = []
        snapshot = (session.metadata or {}).get("lesson_snapshot")
        if isinstance(snapshot, dict) and snapshot.get("lesson_id") == session.lesson_id:
            notes.append("选题依据开课时保存的教案；未使用课后修改替换课堂内容。")
        else:
            notes.append("本会话缺少开课教案快照，选题使用关联教案；布置前请核对课后修改。")
        if not any(item["kind"] == "question" for item in items):
            notes.append("没有选到符合当前目标且材料完整的题库题，也没有可回炉的课堂题；本卷仅包含教案作业任务。")
        return notes

    # ------------------------------------------------------------------
    # DOCX writing
    # ------------------------------------------------------------------

    def _write_paper(
        self,
        path: Path,
        lesson_title: str,
        items: List[Dict[str, Any]],
        summary: List[Dict[str, Any]],
        notes: List[str],
        teacher: bool,
    ) -> None:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        doc = Document()
        section = doc.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.7)
        section.left_margin = section.right_margin = Inches(0.8)
        for style_name, size in (("Normal", 11), ("Title", 16), ("Heading 2", 12)):
            style = doc.styles[style_name]
            style.font.name = "宋体"
            style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor(0, 0, 0)
            for border in style._element.xpath("./w:pPr/w:pBdr"):
                border.getparent().remove(border)
        normal = doc.styles["Normal"].paragraph_format
        normal.space_after = Pt(6)
        normal.line_spacing = 1.15
        normal.widow_control = True
        edition = "教师卷" if teacher else "学生卷"
        doc.core_properties.title = f"{lesson_title} 课后练习卷 {edition}"
        doc.core_properties.author = "WebGIS-AI"
        title = doc.add_paragraph(doc.core_properties.title, style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.keep_with_next = True
        for border in title._p.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.add_run("含参考答案、教案评分参考与已采集的课堂记录，仅供教师使用。" if teacher else
                         "班级 __________  姓名 __________  日期 __________")
        self._add_page_number(section.footer.paragraphs[0], edition)

        if teacher:
            doc.add_paragraph("选题说明", style="Heading 2")
            for entry in summary:
                unit = "项" if entry["origin"].startswith("lesson_homework") else "题"
                doc.add_paragraph(f"{entry['label']}：{entry['count']} {unit}", style="List Bullet")
            doc.add_paragraph("参考答案和解析来自题库或课时原题，未另行核验为官方发布版本。")
            doc.add_paragraph("课堂实测只统计本次已采集的记录；无作答记录时标注未收集，不推测正确率。")
            for note in notes:
                doc.add_paragraph(note, style="List Bullet")
        else:
            doc.add_paragraph("选择题填写选项；综合题写出读图依据与推理过程。拓展题选做，图示可另附。")
        doc.add_paragraph("练习内容", style="Heading 2")
        if not items:
            doc.add_paragraph("本次课堂没有可导出的课后练习内容。")

        shared_until = -1
        for index, item in enumerate(items):
            number = index + 1
            if item["kind"] == "task":
                start = len(doc.paragraphs)
                paragraph = doc.add_paragraph()
                paragraph.add_run(f"{number}. 【{ORIGIN_LEVELS[item['origin']]}】").bold = True
                paragraph.add_run(str(item["text"]))
                if not teacher:
                    self._answer_space(doc, 12 if item["origin"] == "lesson_homework_inquiry" else 6)
                self._keep_block(doc.paragraphs[start:])
                continue
            key = self._shared_material_key(item)
            if index > shared_until and key:
                end = index
                while end + 1 < len(items) and self._shared_material_key(items[end + 1]) == key:
                    end += 1
                if end > index:
                    shared_until = end
                    start = len(doc.paragraphs)
                    doc.add_paragraph(f"第 {number}—{end + 1} 题共用材料", style="Heading 2")
                    self._add_material(doc, item["question"], teacher)
                    self._keep_block(doc.paragraphs[start:], continue_next=True)
            self._add_question(doc, number, item, teacher, shared_material=index <= shared_until)
        if teacher:
            heading = doc.add_paragraph("参考答案与讲评", style="Heading 2")
            heading.paragraph_format.page_break_before = True
            for number, item in enumerate(items, 1):
                if item["kind"] == "task":
                    start = len(doc.paragraphs)
                    doc.add_paragraph(f"第 {number} 题 教师评分参考", style="Heading 2")
                    for point in item.get("teacher_guidance", {}).get("answer_points", []):
                        doc.add_paragraph("教师评分参考：" + point)
                    if len(doc.paragraphs) == start + 1:
                        doc.add_paragraph("本项未提供评分参考，请教师结合教学目标评阅。")
                    self._keep_block(doc.paragraphs[start:])
                else:
                    self._add_reference(doc, number, item)
        doc.save(path)

    @staticmethod
    def _keep_block(paragraphs: List[Any], continue_next: bool = False) -> None:
        for index, paragraph in enumerate(paragraphs):
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = continue_next or index < len(paragraphs) - 1

    @staticmethod
    def _answer_space(doc: Any, lines: int) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt
        for _ in range(lines):
            paragraph = doc.add_paragraph(" ")
            paragraph.paragraph_format.line_spacing = Pt(20)
            paragraph.paragraph_format.space_after = Pt(4)
            borders = OxmlElement("w:pBdr")
            for side in ("bottom", "between"):
                border = OxmlElement(f"w:{side}")
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn(f"w:{key}"), value)
                borders.append(border)
            paragraph._p.get_or_add_pPr().append(borders)

    @staticmethod
    def _add_page_number(paragraph: Any, edition: str) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(f"{edition}  第 ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        paragraph._p.append(field)
        paragraph.add_run(" 页")

    @staticmethod
    def _shared_material_key(item: Dict[str, Any]) -> Optional[Tuple[Any, ...]]:
        if item.get("kind") != "question":
            return None
        question = item["question"]
        material = str(question.get("material") or "").strip()
        images = question.get("images") or []
        if not material or any(image.get("anchor", "group") != "group" for image in images):
            return None
        return (question.get("bank_id"), question.get("group_key"), material,
                tuple(image.get("url") for image in images))

    def _add_material(self, doc: Any, question: Dict[str, Any], teacher: bool) -> None:
        material = str(question.get("material") or "").strip()
        if material:
            doc.add_paragraph(f"材料：{material}")
        self._add_images(doc, question, teacher)

    def _add_question(self, doc: Any, number: int, item: Dict[str, Any], teacher: bool, shared_material: bool = False) -> None:
        question = item["question"]
        body_start = len(doc.paragraphs)
        header = doc.add_paragraph()
        header_run = header.add_run(f"{number}. 【{ORIGIN_LEVELS[item['origin']]}】")
        header_run.bold = True
        if item.get("stage_title"):
            header.add_run(f"（课堂环节：{item['stage_title']}）")

        if not shared_material:
            self._add_material(doc, question, teacher)
        exam_source = " ".join(str(question.get(key) or "").strip() for key in ("year", "region")).strip()
        if exam_source:
            header.add_run(f"  {exam_source} 题库选题")
        task_text = str(question.get("task_text") or "").strip()
        if task_text:
            doc.add_paragraph(task_text)
        stem = str(question.get("text") or "").strip()
        if stem:
            doc.add_paragraph(stem)
        for option_index, option in enumerate(question.get("options") or []):
            doc.add_paragraph(f"{chr(65 + option_index)}. {self._option_text(option_index, option)}")
        for sub in question.get("sub_questions") or []:
            if not isinstance(sub, dict):
                continue
            sub_lines = [f"（{sub.get('index') or ''}）{str(sub.get('text') or '').strip()}"]
            for option_index, option in enumerate(sub.get("options") or []):
                sub_lines.append(f"{chr(65 + option_index)}. {self._option_text(option_index, option)}")
            doc.add_paragraph("　".join(part for part in sub_lines if part))

        # 学生卷到此为止：以下内容只在教师卷生成，学生卷绝不读取答案字段。
        if not teacher:
            if question.get("options"):
                doc.add_paragraph("作答：________")
            else:
                self._answer_space(doc, 5)
            self._keep_block(doc.paragraphs[body_start:])
            return
        self._keep_block(doc.paragraphs[body_start:])

    def _add_reference(self, doc: Any, number: int, item: Dict[str, Any]) -> None:
        question = item["question"]
        reference_start = len(doc.paragraphs)
        reference = doc.add_paragraph()
        reference_run = reference.add_run(f"第 {number} 题 教师参考")
        reference_run.bold = True

        official_answer = self._official_answer_text(question)
        self._labelled_line(doc, "参考答案", official_answer)
        explanation = str(question.get("explanation") or "").strip()
        if explanation:
            self._labelled_line(doc, "参考解析", explanation)
        else:
            self._labelled_line(doc, "参考解析", "本题未提供参考解析")
        sub_answers = self._sub_answer_lines(question)
        for line in sub_answers:
            doc.add_paragraph(line, style="List Bullet")
        knowledge_points = [str(item) for item in question.get("knowledge_points") or [] if str(item).strip()]
        if knowledge_points:
            self._labelled_line(doc, "考点", "、".join(knowledge_points))
        for line in self._classroom_lines(item):
            doc.add_paragraph(line, style="List Bullet")
        source_line = ORIGIN_LABELS[item["origin"]]
        if item.get("selection_reason"):
            source_line += f"（检索依据：{item['selection_reason']}）"
        self._labelled_line(doc, "来源", source_line)
        self._keep_block(doc.paragraphs[reference_start:])

    @staticmethod
    def _option_text(option_index: int, option: Any) -> str:
        """题库选项常自带“A．/A.”字母前缀，渲染前去重，避免出现“A. A．降水”。"""
        letter = chr(65 + option_index)
        text = str(option or "").strip()
        for separator in ("．", ".", "、"):
            prefix = f"{letter}{separator}"
            if text.startswith(prefix):
                return text[len(prefix):].strip()
        return text

    @staticmethod
    def _official_answer_text(question: Dict[str, Any]) -> str:
        answer = str(question.get("answer") or "").strip()
        if answer:
            return answer
        letter = str(question.get("answer_letter") or "").strip()
        if letter:
            return f"正确选项：{letter}"
        answer_index = question.get("answer_index")
        options = question.get("options") or []
        if isinstance(answer_index, int) and 0 <= answer_index < len(options):
            return f"正确选项：{chr(65 + answer_index)}. {options[answer_index]}"
        return "本题未提供参考答案"

    @staticmethod
    def _sub_answer_lines(question: Dict[str, Any]) -> List[str]:
        lines: List[str] = []
        for sub in question.get("sub_questions") or []:
            if not isinstance(sub, dict):
                continue
            answer = str(sub.get("answer") or "").strip()
            if not answer:
                answer_index = sub.get("answer_index")
                options = sub.get("options") or []
                if isinstance(answer_index, int) and 0 <= answer_index < len(options):
                    answer = f"正确选项：{chr(65 + answer_index)}. {options[answer_index]}"
            if not answer:
                continue
            line = f"小题（{sub.get('index') or ''}）参考答案：{answer}"
            explanation = str(sub.get("explanation") or "").strip()
            if explanation:
                line += f"｜解析：{explanation}"
            lines.append(line)
        return lines

    @staticmethod
    def _classroom_lines(item: Dict[str, Any]) -> List[str]:
        lines: List[str] = []
        classroom = item.get("classroom") or {}
        observation = item.get("observation") or {}
        if classroom.get("launched"):
            response_count = int(classroom.get("response_count") or 0)
            if response_count > 0:
                parts = [f"{response_count} 人作答"]
                correct_rate = classroom.get("correct_rate")
                if correct_rate is not None:
                    parts.append(f"正确率 {correct_rate:.0%}")
                lines.append("课堂实测：" + "，".join(parts))
            else:
                lines.append("课堂实测：本题已投屏，未收集到作答数据")
            if classroom.get("actual_seconds") is not None:
                timing = f"投屏用时 {int(classroom['actual_seconds'])} 秒"
                overtime = int(classroom.get("overtime_seconds") or 0)
                if overtime > 0:
                    timing += f"（超时 {overtime} 秒）"
                if classroom.get("revealed"):
                    timing += "，课堂上已揭示答案"
                else:
                    timing += "，课堂上未揭示答案"
                lines.append("课堂实测：" + timing)
        elif classroom.get("response_count"):
            lines.append(f"课堂实测：{int(classroom['response_count'])} 人作答（未投屏记录）")
        if observation.get("verdict"):
            detail = VERDICT_LABELS.get(observation["verdict"], observation["verdict"])
            if observation.get("tag"):
                detail += f"（标签：{observation['tag']}）"
            if observation.get("note"):
                detail += f"——{observation['note']}"
            lines.append(f"课堂速记：{detail}")
        return lines

    @staticmethod
    def _labelled_line(doc: Any, label: str, value: str) -> None:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(f"{label}：")
        run.bold = True
        paragraph.add_run(str(value))

    def _add_images(self, doc: Any, question: Dict[str, Any], teacher: bool) -> None:
        """嵌入题图；无法解析的题图保留占位说明，绝不静默丢弃。"""
        from docx.shared import Inches

        images = [image for image in question.get("images") or [] if isinstance(image, dict)]
        if not images:
            return
        embedded = 0
        for image in images:
            local_path = self._resolve_image_path(str(image.get("url") or ""))
            if local_path is None:
                continue
            try:
                doc.add_picture(str(local_path), width=Inches(3.5))
                embedded += 1
            except Exception:
                continue
        if embedded < len(images):
            doc.add_paragraph(f"（本题有 {len(images) - embedded} 张题图未能嵌入，请对照题库原卷查看。）")

    def _resolve_image_path(self, url: str) -> Optional[Path]:
        prefix = "/files/"
        if not url.startswith(prefix):
            return None
        try:
            path = self.config.resolve_public_path(url[len(prefix):])
        except (ValueError, OSError):
            return None
        if not path.is_file():
            return None
        return path
