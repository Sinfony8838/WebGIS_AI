"""Pure, whitelisted document operations and a bounded JSON subprocess protocol.

Run only via DocumentExecutor. No runtime, configuration, store or model imports.
The parent assigns a Windows Job Object before granting the stdin start token;
POSIX address-space limits are installed before any parser imports or input read.
"""
from __future__ import annotations

import base64
import dataclasses
import io
import json
import os
import signal
import time
from pathlib import Path
import sys
from typing import Any, Callable

OPERATIONS = frozenset({"office_validate", "lesson_docx", "question_bank_docx", "timeline_text"})


class ResultLimitError(ValueError):
    pass


class WireEncoder(json.JSONEncoder):
    def default(self, value: Any) -> Any:
        if isinstance(value, bytes):
            return {"$bytes": base64.b64encode(value).decode("ascii")}
        if dataclasses.is_dataclass(value) and not isinstance(value, type):
            return {field.name: getattr(value, field.name) for field in dataclasses.fields(value)}
        return super().default(value)


def wire_object(value: dict) -> Any:
    if set(value) == {"$bytes"}:
        return base64.b64decode(value["$bytes"], validate=True)
    return value


def restore_result(operation: str, value: Any) -> Any:
    """Restore only the fixed question-bank data schema; never unpickle input."""
    expected = {"office_validate": type(None), "lesson_docx": dict, "timeline_text": str}
    if operation in expected:
        if not isinstance(value, expected[operation]):
            raise ValueError("Invalid document result type")
        return value
    if operation != "question_bank_docx" or not isinstance(value, dict):
        raise ValueError("Invalid document operation or result type")
    from .question_bank_parser import ParsedBank, ParsedGroup, ParsedImage, ParsedQuestion

    groups = []
    for raw_group in value["groups"]:
        group = dict(raw_group)
        group["questions"] = [ParsedQuestion(**item) for item in group["questions"]]
        group["images"] = [ParsedImage(**item) for item in group["images"]]
        groups.append(ParsedGroup(**group))
    bank = dict(value)
    bank["groups"] = groups
    bank["unbound_images"] = [ParsedImage(**item) for item in bank["unbound_images"]]
    return ParsedBank(**bank)


def extract_timeline_text(filename: str, raw: bytes) -> str:
    """The deterministic extraction half of TimelineService, without LLM import."""
    from .archive_safety import validate_office_bytes

    lower = filename.lower()
    if lower.endswith(".pptx"):
        validate_office_bytes(raw, "pptx")
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ValueError("python-pptx is not installed") from exc
        return "\n".join(shape.text.strip() for slide in Presentation(io.BytesIO(raw)).slides
                         for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip())
    if lower.endswith(".pdf"):
        try:
            import fitz
        except ImportError as exc:
            raise ValueError("PyMuPDF is not installed") from exc
        with fitz.open(stream=raw, filetype="pdf") as document:
            return "\n".join(page.get_text() for page in document)
    if lower.endswith(".docx"):
        validate_office_bytes(raw, "docx")
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("python-docx is not installed") from exc
        try:
            document = Document(io.BytesIO(raw))
        except Exception as exc:
            raise ValueError(f"Unable to read Word document: {exc}") from exc
        texts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    texts.append("\t".join(cells))
        return "\n".join(texts)
    if lower.endswith(".doc"):
        raise ValueError("Legacy .doc files are not supported. Please save the lesson plan as .docx.")
    if lower.endswith((".txt", ".md")):
        return raw.decode("utf-8", errors="replace")
    raise ValueError(f"不支持的文件格式: {filename}")


def execute_operation(operation: str, raw: bytes, filename: str, kind: str) -> Any:
    if operation not in OPERATIONS:
        raise ValueError("Unsupported document operation")
    if operation == "office_validate":
        from .archive_safety import validate_office_bytes
        validate_office_bytes(raw, kind)
        return None
    if operation == "lesson_docx":
        from .lesson_docx_import import parse_lesson_docx
        return parse_lesson_docx(raw)
    if operation == "question_bank_docx":
        from .question_bank_parser import QuestionBankDocxParser
        return QuestionBankDocxParser(filename, raw).parse()
    return extract_timeline_text(filename, raw)


def _write_response(path: Path, value: dict, limit: int) -> None:
    partial = path.with_suffix(".partial")
    total = 0
    with partial.open("wb") as target:
        for part in WireEncoder(ensure_ascii=False, separators=(",", ":"), allow_nan=False).iterencode(value):
            chunk = part.encode("utf-8")
            total += len(chunk)
            if total > limit:
                raise ResultLimitError("Document result exceeds the IPC size budget")
            target.write(chunk)
    os.replace(partial, path)


def _linux_supervise(directory: Path, run: Callable[[], int]) -> int:
    """Own orphaned parser descendants without changing the backend's prctl state.

    The parser gets a separate process group. SIGTERM only asks this tiny
    supervisor to stop; it kills that group and waitpid-reaps every descendant
    before exiting. Trusted parsers must not deliberately escape the group.
    """
    import ctypes

    libc = ctypes.CDLL(None, use_errno=True)
    if libc.prctl(36, 1, 0, 0, 0) != 0:  # PR_SET_CHILD_SUBREAPER
        return 74
    stopping = False

    def stop(_signum, _frame):
        nonlocal stopping
        stopping = True

    signal.signal(signal.SIGTERM, stop)
    parser_pid = os.fork()
    if parser_pid == 0:
        signal.signal(signal.SIGTERM, signal.SIG_DFL)
        os.setpgid(0, 0)
        try:
            code = run()
        except BaseException:
            code = 75
        os._exit(code)
    status = None
    try:
        # Both sides set the group to remove the fork/setup signal race.
        try:
            os.setpgid(parser_pid, parser_pid)
        except ProcessLookupError:
            pass
        (directory / "parser-group.pid").write_text(str(parser_pid), encoding="ascii")
        while not stopping:
            child, child_status = os.waitpid(parser_pid, os.WNOHANG)
            if child:
                status = child_status
                break
            time.sleep(0.01)
    finally:
        try:
            os.killpg(parser_pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        # Linux reparents orphans to this dedicated subreaper. No unrelated
        # backend child can be reaped here, including concurrent document jobs.
        while True:
            try:
                os.waitpid(-1, 0)
            except ChildProcessError:
                break
        (directory / "process-tree-reaped").touch()
    if stopping or status is None:
        return 73
    return os.waitstatus_to_exitcode(status) if os.WIFEXITED(status) else 76


def worker_main(dispatch: Callable = execute_operation) -> int:
    """Fixed file names in a parent-created private directory; no user paths."""
    directory = Path(sys.argv[1])
    memory_limit = int(sys.argv[2])
    result_limit = int(sys.argv[3])
    input_limit = int(sys.argv[4])
    if sys.platform.startswith("linux"):
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (memory_limit, memory_limit))
        resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
    elif os.name != "nt":
        return 70  # Unsupported platforms fail closed.
    if sys.stdin.buffer.read(1) != b"G":
        return 71
    if sys.platform.startswith("linux"):
        return _linux_supervise(directory, lambda: _run_request(directory, input_limit, result_limit, dispatch))
    return _run_request(directory, input_limit, result_limit, dispatch)


def _run_request(directory: Path, input_limit: int, result_limit: int, dispatch: Callable) -> int:
    result_path = directory / "result.json"
    # Preallocate recovery space so ordinary allocation failures can be reported.
    reserve = bytearray(1024 * 1024)
    try:
        request = json.loads((directory / "request.json").read_text(encoding="utf-8"))
        path = directory / "input.bin"
        if path.stat().st_size > input_limit:
            raise ResultLimitError("Document input exceeds the IPC size budget")
        raw = path.read_bytes()
        value = dispatch(request["operation"], raw, request["filename"], request["kind"])
        _write_response(result_path, {"ok": True, "value": value}, result_limit)
        return 0
    except MemoryError:
        del reserve
        response = {"ok": False, "error": "memory", "message": "Document process exceeded its memory budget"}
    except Exception as exc:
        # Preserve the established validation categories, not arbitrary classes.
        from .archive_safety import ArchiveLimitError, ArchiveValidationError
        from .question_bank_parser import DocxBudgetError, DocxValidationError
        if isinstance(exc, DocxBudgetError):
            category = "docx_budget"
        elif isinstance(exc, DocxValidationError):
            category = "docx_validation"
        elif isinstance(exc, ArchiveLimitError):
            category = "archive_limit"
        elif isinstance(exc, ArchiveValidationError):
            category = "archive_validation"
        elif isinstance(exc, ResultLimitError):
            category = "result_limit"
        elif isinstance(exc, ValueError):
            category = "validation"
        else:
            category = "worker"
        response = {"ok": False, "error": category,
                    "message": str(exc)[:2000] if category != "worker" else "Document parser failed"}
    try:
        _write_response(result_path, response, result_limit)
        return 0
    except Exception:
        return 72


if __name__ == "__main__":
    raise SystemExit(worker_main())
