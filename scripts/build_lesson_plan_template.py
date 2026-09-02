from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "backend" / "app" / "data" / "builtin" / "lesson_plan_template" / "lesson_plan_template.docx"


def set_font(style, western: str, east_asia: str, size: float, bold: bool = False) -> None:
    style.font.name = western
    style.font.size = Pt(size)
    style.font.bold = bold
    style._element.rPr.rFonts.set(qn("w:ascii"), western)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), western)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end): run._r.append(node)
    paragraph.add_run(" 页")


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Mm(210); section.page_height = Mm(297)
    section.top_margin = Inches(0.5); section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5); section.right_margin = Inches(0.5)
    section.header_distance = Inches(0.3); section.footer_distance = Inches(0.3)
    normal = doc.styles["Normal"]
    set_font(normal, "Times New Roman", "宋体", 10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    title = doc.styles.add_style("Lesson Title", 1)
    set_font(title, "Arial", "黑体", 20, True)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading = doc.styles.add_style("Lesson Heading", 1)
    set_font(heading, "Arial", "黑体", 12, True)
    heading.font.color.rgb = RGBColor(31, 77, 120)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(4)
    emphasis = doc.styles.add_style("Lesson System Step", 1)
    set_font(emphasis, "Arial", "黑体", 9, True)
    emphasis.font.color.rgb = RGBColor(31, 77, 120)
    doc.core_properties.title = "高中地理教案模板"
    doc.core_properties.subject = "WebGIS-AI 教案导出"
    doc.core_properties.author = "WebGIS-AI"
    doc.core_properties.last_modified_by = "WebGIS-AI"
    doc.add_paragraph("高中地理教学设计", style="Lesson Title")
    doc.add_paragraph("此文件是去个人信息的版式模板，运行时会替换为教师确认后的教案内容。")
    add_page_field(section.footer.paragraphs[0])
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
