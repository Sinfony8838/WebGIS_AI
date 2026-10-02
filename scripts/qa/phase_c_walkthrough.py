"""Phase C browser acceptance, using real mouse/keyboard input in an isolated context.

Start the C frontend separately with VITE_API_BASE_URL pointing to the QA backend.
Set WEBGIS_QA_EMAIL / WEBGIS_QA_PASSWORD; no production credentials or runtime state.
Evidence is written to ignored scratch/phase-c. --fixture-only needs no running service.
"""
from __future__ import annotations

import argparse
import json
import os
from pathlib import Path

from playwright.sync_api import expect, sync_playwright
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
EVIDENCE = ROOT / "scratch" / "phase-c"
RESULTS: list[dict] = []


def fixture() -> Path:
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    ignore = EVIDENCE / ".gitignore"
    if not ignore.exists():
        ignore.write_text("*\n", encoding="utf-8")
    path = EVIDENCE / "phase-c-acceptance.pptx"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(5.625)
    for number, text in [(1, "Population observation"), (2, "Terrain comparison")]:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.7), Inches(0.6), Inches(8.6), Inches(1))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = f"Phase C - page {number}: {text}"
        paragraph.font.size = Pt(28)
        box = slide.shapes.add_textbox(Inches(0.7), Inches(2), Inches(8.6), Inches(2))
        box.text_frame.text = "Draw here. Switch pages. Return and undo."
        box.text_frame.paragraphs[0].font.size = Pt(24)
    deck.save(path)
    return path


def record(name: str, ok: bool, detail=None):
    RESULTS.append({"step": name, "ok": ok, "detail": detail})
    print(("PASS " if ok else "FAIL ") + name + (f" | {detail}" if detail else ""), flush=True)


def screenshot(page, name: str):
    page.screenshot(path=str(EVIDENCE / f"{name}.png"), full_page=False)


def step(page, name: str, action):
    try:
        detail = action()
        record(name, True, detail)
        return True
    except Exception as exc:
        record(name, False, str(exc)[:1800])
        screenshot(page, f"failure-{len(RESULTS):02}")
        return False


def draw(page, locator, offset=0):
    box = locator.bounding_box()
    assert box
    x, y = box["x"] + box["width"] * 0.30, box["y"] + box["height"] * 0.65 + offset
    page.mouse.move(x, y)
    page.mouse.down()
    page.mouse.move(x + 140, y - 50, steps=12)
    page.mouse.move(x + 210, y + 20, steps=8)
    page.mouse.up()


def ink(locator) -> int:
    return locator.evaluate("canvas => { const a = canvas.getContext('2d').getImageData(0,0,canvas.width,canvas.height).data; let n=0; for(let i=3;i<a.length;i+=4) if(a[i]) n++; return n; }")


def wait_for_ink(page, locator, expected: int) -> int:
    for _ in range(60):
        actual = ink(locator)
        if actual == expected:
            return actual
        page.wait_for_timeout(50)
    raise AssertionError({"expected_pixels": expected, "actual_pixels": actual})


def overlap(first, second) -> bool:
    a, b = first.bounding_box(), second.bounding_box()
    assert a and b
    return min(a["x"]+a["width"], b["x"]+b["width"]) > max(a["x"], b["x"]) and min(a["y"]+a["height"], b["y"]+b["height"]) > max(a["y"], b["y"])


def move_profile(page, *, lower_right: bool):
    window = page.locator("[data-testid^='profile-window-pw_']").first
    expect(window).to_be_visible()
    box, head = window.bounding_box(), window.locator(".profile-window-head").bounding_box()
    assert box and head
    viewport = page.viewport_size
    left = viewport["width"]-box["width"]-8 if lower_right else 96
    top = viewport["height"]-box["height"]-8 if lower_right else 120
    x, y = head["x"]+80, head["y"]+16
    page.mouse.move(x, y)
    page.mouse.down()
    page.mouse.move(x+left-box["x"], y+top-box["y"], steps=16)
    page.mouse.up()
    return window


def verify_clearance(page, foreground, obstacles):
    """Geometric check for chrome controls (teaching canvas may be overlaid intentionally)."""
    box = foreground.bounding_box()
    assert box
    viewport = page.viewport_size
    assert 0 <= box["x"] and 0 <= box["y"]
    assert box["x"] + box["width"] <= viewport["width"] + 1
    assert box["y"] + box["height"] <= viewport["height"] + 1
    overlap = []
    for selector in obstacles:
        for target in page.locator(selector).all():
            if not target.is_visible():
                continue
            other = target.bounding_box()
            if other and min(box["x"]+box["width"], other["x"]+other["width"]) > max(box["x"],other["x"]) and min(box["y"]+box["height"],other["y"]+other["height"]) > max(box["y"],other["y"]):
                overlap.append(selector)
    assert not overlap, {"box": box, "overlap": overlap}
    return box


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--front", default="http://127.0.0.1:5826")
    parser.add_argument("--fixture-only", action="store_true")
    args = parser.parse_args()
    ppt = fixture()
    if args.fixture_only:
        print(ppt)
        return 0
    email, password = os.environ["WEBGIS_QA_EMAIL"], os.environ["WEBGIS_QA_PASSWORD"]
    errors: list[dict] = []
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        context = browser.new_context(viewport={"width": 1366, "height": 768})
        page = context.new_page()
        page.on("pageerror", lambda error: errors.append({"type": "pageerror", "message": str(error)[:500]}))
        page.on("console", lambda message: errors.append({"type": "console", "message": message.text[:500]}) if message.type == "error" else None)
        page.on("response", lambda response: errors.append({"type": "http", "status": response.status, "url": response.url}) if response.status >= 400 else None)
        try:
            def login():
                page.goto(args.front, wait_until="domcontentloaded")
                page.locator("input[type=email]").fill(email)
                page.locator("input[type=password]").fill(password)
                page.get_by_role("button", name="登录", exact=True).last.click()
                expect(page.locator("input[type=email]")).to_have_count(0, timeout=30000)
                page.wait_for_function("Object.keys(localStorage).some(k=>k.startsWith('webgis_ai_project_id:'))", timeout=30000)
                page.get_by_role("button", name="展开地图工具与可视化地图").click()
                page.wait_for_timeout(2000)
                project_ids = page.evaluate("Object.keys(localStorage).filter(k=>k.startsWith('webgis_ai_project_id:')).map(k=>localStorage.getItem(k))")
                return {"isolated_context_project_ids": project_ids}
            if not step(page, "独立上下文登录并自动建立 QA 项目", login):
                return 1

            def map_navigation():
                toggle = page.get_by_test_id("map-mode-toggle")
                expect(toggle).to_be_visible()
                if toggle.get_attribute("aria-pressed") == "true":
                    toggle.click()
                expect(page.locator(".ol-viewport").first).to_be_visible()
                visual = page.get_by_role("button", name="可视化地图", exact=True)
                assert visual.locator("svg").count() >= 1
                visual.focus()
                visual.press("Enter")
                expect(visual).to_have_attribute("aria-expanded", "true")
                visual.press("Space")
                expect(visual).to_have_attribute("aria-expanded", "false")
                rail_buttons = page.get_by_test_id("map-tool-rail").locator("button")
                assert all(button.locator("svg").count() for button in rail_buttons.all())
            step(page, "2D/3D 切换与 SVG 导航键盘操作", map_navigation)

            def map_measure():
                page.get_by_role("button", name="测距模式 · 快捷键 M").click()
                help_button = page.get_by_role("button", name="测距模式操作帮助")
                expect(page.locator(".map-operation-help")).to_have_count(0)
                help_button.click()
                expect(page.locator(".map-operation-help")).to_be_visible()
                for width, height in [(1366,768), (1920,1080)]:
                    page.set_viewport_size({"width":width,"height":height})
                    verify_clearance(page, page.locator(".map-operation-controls"), [".right-rail", ".header-search"])
                    screenshot(page, f"map-help-{width}")
                help_button.click()
                viewport = page.locator(".ol-viewport").first.bounding_box()
                cx, cy = viewport["x"]+viewport["width"]/2, viewport["y"]+viewport["height"]/2
                page.mouse.move(cx,cy)
                for _ in range(5):
                    page.mouse.wheel(0,-300)
                    page.wait_for_timeout(180)
                page.mouse.click(cx-140,cy-35)
                page.mouse.click(cx+120,cy+45)
                page.get_by_role("button", name="完成", exact=True).click()
                expect(page.get_by_test_id("profile-windows-bar")).to_be_visible()
                expect(page.get_by_test_id("profile-manager-toggle")).to_have_attribute("aria-expanded", "false")
                page.get_by_test_id("profile-manager-toggle").click()
                expect(page.locator(".profile-manager-details")).to_be_visible()
                first_control = page.locator(".profile-manager-details button:not(:disabled)").first
                expect(first_control).to_be_focused()
                page.locator("[data-testid^='profile-line-open-population-']").first.click()
                page.keyboard.press("Escape")
                expect(page.locator(".profile-manager-details")).to_be_hidden()
                expect(page.get_by_test_id("profile-manager-toggle")).to_be_focused()
                window = move_profile(page, lower_right=True)
                assert overlap(window, page.get_by_test_id("profiles-collapse-all"))
                screenshot(page,"profile-controls-overlapping-window")
                page.get_by_test_id("profiles-collapse-all").click()
                expect(page.get_by_test_id("profiles-collapse-all")).to_have_attribute("aria-pressed","true")
                page.get_by_test_id("profiles-collapse-all").click()
                move_profile(page, lower_right=False)
                page.get_by_role("button", name="测距模式 · 快捷键 M").click()
                page.get_by_role("button", name="取消", exact=True).click()
                expect(page.locator(".map-operation-controls")).to_have_count(0)
            step(page, "地图帮助按需展示、完成/取消与剖面管理可达", map_measure)

            def map_brush():
                page.get_by_role("button",name="画笔模式 · 快捷键 P").click()
                expect(page.get_by_role("button",name="展开地图画笔设置")).to_be_visible()
                page.get_by_role("button",name="展开地图画笔设置").click()
                for width,height in [(1366,768),(1920,1080)]:
                    page.set_viewport_size({"width":width,"height":height})
                    verify_clearance(page,page.locator(".map-brush-panel"),[".right-rail", ".map-operation-controls", ".profile-manager-bar"])
                    screenshot(page,f"map-brush-{width}")
                page.locator(".map-brush-panel").get_by_role("button",name="画笔",exact=True).click()
                expect(page.get_by_role("button",name="展开地图画笔设置")).to_be_visible()
                canvas=page.locator(".brush-overlay--active")
                draw(page,canvas)
                before_clear=ink(canvas)
                assert before_clear > 0
                page.get_by_role("button",name="展开地图画笔设置").click()
                page.locator(".map-brush-panel").get_by_role("button",name="清除",exact=True).click()
                wait_for_ink(page,canvas,0)
                expect(page.locator(".map-brush-panel").get_by_role("button",name="撤销",exact=True)).to_be_enabled()
                expect(page.locator(".map-brush-panel").get_by_role("button",name="清除",exact=True)).to_be_disabled()
                page.locator(".map-brush-panel").get_by_role("button",name="撤销",exact=True).click()
                wait_for_ink(page,canvas,before_clear)
                screenshot(page,"map-clear-undo-restored")
                page.get_by_role("button",name="取消",exact=True).click()
                expect(page.locator(".map-brush-panel")).to_have_count(0)
            step(page,"地图画笔侧边折叠、真实绘制与桌面尺寸检查",map_brush)

            def upload_ppt():
                # Opening a deck from map drawing must leave the drawing owners isolated.
                page.get_by_role("button",name="画笔模式 · 快捷键 P").click()
                with page.expect_file_chooser() as chooser:
                    page.get_by_role("button",name="导入 PPT",exact=True).click()
                chooser.value.set_files(str(ppt))
                expect(page.locator(".ppt-viewer-counter")).to_have_text("1 / 2",timeout=120000)
                expect(page.get_by_test_id("map-brush-overlay")).to_have_count(0)
                expect(page.locator(".map-operation-controls")).to_have_count(0)
                expect(page.locator(".ppt-brush-float")).to_have_count(0)
                page.wait_for_function("[...document.querySelectorAll('.ppt-viewer-slide-image')].every(img=>img.complete && img.naturalWidth>0)")
                assert page.locator(".ppt-viewer-slide").inner_text() or page.locator(".ppt-viewer-slide-image").count()
                screenshot(page,"ppt-upload")
                return {"render_mode": "image" if page.locator(".ppt-viewer-slide-image").count() else "fallback_html"}
            if step(page,"PPT 文件选择上传并可用两页",upload_ppt):
                def brush_flow():
                    page.locator(".ppt-viewer-stage").click(position={"x":10,"y":10})
                    page.keyboard.press("p")
                    expect(page.get_by_test_id("ppt-brush-float")).to_be_visible()
                    expect(page.get_by_test_id("map-brush-overlay")).to_have_count(0)
                    expect(page.locator(".map-operation-controls")).to_have_count(0)
                    for width,height in [(1366,768),(1920,1080)]:
                        page.set_viewport_size({"width":width,"height":height})
                        verify_clearance(page,page.get_by_test_id("ppt-brush-float"),[".ppt-viewer-controls", ".ppt-viewer-nav"])
                        screenshot(page,f"ppt-brush-expanded-{width}")
                    page.get_by_role("button",name="画笔工具自由",exact=True).click()
                    expect(page.get_by_test_id("ppt-brush-mini")).to_be_visible()
                    canvas=page.locator(".ppt-viewer-slide canvas")
                    draw(page,canvas)
                    first_ink=ink(canvas)
                    assert first_ink>0, first_ink
                    draw(page,canvas,offset=85)
                    two_ink=ink(canvas)
                    assert two_ink>first_ink, (first_ink,two_ink)
                    page.keyboard.press("ArrowRight")
                    expect(page.locator(".ppt-viewer-counter")).to_have_text("2 / 2")
                    page.wait_for_timeout(150)
                    assert ink(canvas)==0
                    page.get_by_role("button",name="上一页",exact=True).click()
                    expect(page.locator(".ppt-viewer-counter")).to_have_text("1 / 2")
                    restored=wait_for_ink(page,canvas,two_ink)
                    screenshot(page,"ppt-restored-ink")
                    page.get_by_role("button",name="展开画笔设置",exact=True).click()
                    page.get_by_role("button",name="撤销",exact=True).click()
                    after_undo=wait_for_ink(page,canvas,first_ink)
                    page.get_by_role("button",name="清空",exact=True).click()
                    wait_for_ink(page,canvas,0)
                    expect(page.get_by_role("button",name="撤销",exact=True)).to_be_enabled()
                    expect(page.get_by_role("button",name="清空",exact=True)).to_be_disabled()
                    page.get_by_role("button",name="撤销",exact=True).click()
                    after_clear_undo=wait_for_ink(page,canvas,first_ink)
                    screenshot(page,"ppt-per-stroke-clear-undo")
                    page.keyboard.press("Escape")
                    expect(page.locator(".ppt-brush-float")).to_have_count(0)
                    expect(page.locator(".ppt-viewer-backdrop")).to_be_visible()
                    page.keyboard.press("Escape")
                    expect(page.locator(".ppt-viewer-backdrop")).to_have_count(0)
                    expect(page.locator(".ppt-viewer-dock")).to_be_visible()
                    expect(page.locator(".map-brush-panel")).to_have_count(0)
                    expect(page.get_by_test_id("map-brush-overlay")).to_be_visible()
                    assert ink(page.get_by_test_id("map-brush-overlay")) > 0
                    for width,height in [(1366,768),(1920,1080)]:
                        page.set_viewport_size({"width":width,"height":height})
                        verify_clearance(page,page.locator(".ppt-viewer-dock"),[".profile-manager-bar", ".right-rail"])
                        window=move_profile(page,lower_right=True)
                        assert overlap(window,page.locator(".ppt-viewer-dock-main"))
                        screenshot(page,f"ppt-collapsed-dock-{width}")
                        page.locator(".ppt-viewer-dock-main").click()
                        expect(page.locator(".ppt-viewer-backdrop")).to_be_visible()
                        page.locator(".ppt-viewer-controls").get_by_role("button",name="收起",exact=True).click()
                    page.locator(".ppt-viewer-dock-main").click()
                    expect(page.locator(".ppt-viewer-counter")).to_have_text("1 / 2")
                    page.wait_for_timeout(300)
                    assert ink(canvas)>0
                    page.locator(".ppt-viewer-stage").click(position={"x":10,"y":10})
                    page.keyboard.press("p")
                    page.locator(".ppt-viewer-controls").get_by_role("button",name="收起",exact=True).click()
                    expect(page.locator(".map-brush-panel")).to_have_count(0)
                    expect(page.locator(".map-operation-controls")).to_have_count(0)
                    page.locator(".ppt-viewer-dock-main").click()
                    return {"first_stroke_pixels":first_ink,"two_stroke_pixels":two_ink,"page_return_pixels":restored,"after_one_undo_pixels":after_undo,"after_clear_pixels":0,"after_clear_undo_pixels":after_clear_undo,"dock_click_over_profile_sizes":[1366,1920]}
                step(page,"PPT 工具不遮挡、选后收起、绘制/切页/恢复/撤销与两次 Esc",brush_flow)

                def input_guard():
                    # This is the app's real POI input, still rendered while the presentation is open.
                    field=page.get_by_role("textbox",name="POI 检索关键词")
                    field.focus()
                    field.fill("phase c input")
                    for key in ["ArrowRight","ArrowLeft","Home","End","Space","p","m"]:
                        field.press(key)
                    expect(page.locator(".ppt-viewer-counter")).to_have_text("1 / 2")
                    expect(page.locator(".ppt-brush-float")).to_have_count(0)
                    screenshot(page,"ppt-input-guard")
                step(page,"真实输入控件方向键/空格/工具键不误翻页",input_guard)
        finally:
            (EVIDENCE/"results.json").write_text(json.dumps({"results":RESULTS,"browser_errors":errors},ensure_ascii=False,indent=2),encoding="utf-8")
            context.close()
            browser.close()
    return 1 if any(not result["ok"] for result in RESULTS) else 0


if __name__ == "__main__":
    raise SystemExit(main())
