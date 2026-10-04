from pathlib import Path
import re
import json
import shutil
from collections import Counter
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from PIL import Image

QA = Path(__file__).resolve().parents[1] / 'backend/data/outputs'
WT = Path(__file__).resolve().parents[1]
OUT = WT / 'docs'
MD = (OUT / '教师版产品说明.md').read_text(encoding='utf-8')
FONT = '宋体'
UIFONT = '微软雅黑'
BLACK = '000000'
TEAL = '075D83'
SOFT = '496278'
WIDTH = 16.7

doc = Document()
doc.core_properties.title = 'GeoBot 智能教学平台 产品介绍与教师使用指南'
doc.core_properties.subject = '教师使用说明与中期答辩阶段成果'
doc.core_properties.author = 'WebGIS-AI 项目组'
doc.core_properties.keywords = 'GeoBot 地理教学 WebGIS 产品介绍 教师使用指南'
doc.core_properties.comments = ''

def font_settings(style, size, bold=False, face=FONT, color=BLACK):
    style.font.name = face
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor.from_string(color)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), face)
    style.paragraph_format.widow_control = True

font_settings(doc.styles['Normal'], 12)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.16
normal.space_after = Pt(7)
normal.space_before = Pt(0)
for name, size in [('Title', 32), ('Subtitle', 17), ('Heading 1', 20), ('Heading 2', 14)]:
    font_settings(doc.styles[name], size, name != 'Subtitle', UIFONT)
    pf = doc.styles[name].paragraph_format
    pf.space_before = Pt(16 if name == 'Heading 1' else 12)
    pf.space_after = Pt(9)
    pf.keep_with_next = True
    for node in list(doc.styles[name].element.iter()):
        if node.tag == qn('w:pBdr'):
            node.getparent().remove(node)
doc.styles['Heading 1'].paragraph_format.page_break_before = False
for name, size, face, color in [
    ('Guide Front Heading', 18, UIFONT, BLACK),
    ('Guide Caption', 10.5, FONT, SOFT),
    ('Guide Table', 10.5, FONT, BLACK),
    ('Guide Cover Meta', 11, UIFONT, SOFT),
    ('TOC 1', 11, UIFONT, BLACK),
    ('TOC 2', 10.5, FONT, BLACK),
]:
    st = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles['Normal']
    font_settings(st, size, name == 'Guide Front Heading', face, color)
    st.paragraph_format.space_after = Pt(3 if name.startswith('TOC') else 7)
    st.paragraph_format.line_spacing = 1.08 if name.startswith('TOC') else 1.16
    if name.startswith('TOC'):
        st.paragraph_format.space_before = Pt(0)
    if name == 'TOC 2':
        st.paragraph_format.left_indent = Cm(0.55)
    if name == 'Guide Front Heading':
        st.paragraph_format.keep_with_next = True

for name in ['Header', 'Footer']:
    font_settings(doc.styles[name], 9, False, UIFONT)
    doc.styles[name].paragraph_format.space_after = Pt(0)

def section_settings(s):
    # Preserve the report's existing A4 submission format.
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.15)
    s.top_margin, s.bottom_margin = Cm(2), Cm(2.1)
    s.header_distance, s.footer_distance = Cm(0.8), Cm(0.8)
    s.header.is_linked_to_previous = False
    s.footer.is_linked_to_previous = False

def field(p, instr, placeholder=''):
    r = p.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._r.append(b)
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; r._r.append(it)
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate'); r._r.append(sep)
    if placeholder:
        t = OxmlElement('w:t'); t.text = placeholder; r._r.append(t)
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)

def page_number(s, start=1, fmt='decimal'):
    n = OxmlElement('w:pgNumType')
    n.set(qn('w:start'), str(start)); n.set(qn('w:fmt'), fmt)
    s._sectPr.append(n)
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(p, ' PAGE ')

def band(header, top, height, color):
    xml = '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><v:rect style="position:absolute;left:0;top:%spt;width:595.3pt;height:%spt;z-index:-251654144;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" fillcolor="#%s" stroked="f"><v:fill color="#%s"/></v:rect></w:pict>' % (top, height, color, color)
    header.paragraphs[0].add_run()._r.append(parse_xml(xml))

def inline(p, text):
    text = re.sub(r'\[([^]]+)\]\([^)]+\)', r'\1', text)
    pieces = re.split(r'(\*\*.*?\*\*)', text)
    for part in pieces:
        if not part:
            continue
        r = p.add_run(part[2:-2] if part.startswith('**') and part.endswith('**') else part)
        if part.startswith('**') and part.endswith('**'):
            r.bold = True
    if len(pieces) == 3 and not pieces[0] and not pieces[2]:
        p.paragraph_format.keep_with_next = True
    if text.startswith(('**推进环节。**', '**控制揭示。**')):
        p.paragraph_format.keep_with_next = True

def para(text='', style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    inline(p, text)
    return p

def centered(text, style, before=0, after=6, size=None, color=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    if size is not None:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = RGBColor.from_string(color)
    return p

cover = doc.sections[0]
section_settings(cover)
band(cover.header, 0, 42, 'EEF3F8')
band(cover.header, 719, 123, 'F3F7FB')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(10)
p.add_run().add_picture(str(OUT / 'images/00-product-logo.png'), width=Cm(3.6))
centered('GeoBot', 'Title', before=8, after=7, size=34)
centered('智能教学平台', 'Title', before=0, after=12, size=25)
centered('产品介绍与教师使用指南', 'Subtitle', before=0, after=7)
centered('中期答辩阶段成果说明', 'Guide Cover Meta', after=14)
centered('课前备课  /  课中教学  /  课后复盘', 'Guide Cover Meta', after=16, size=12, color=TEAL)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
p.add_run().add_picture(str(OUT / 'images/21-teacher-finland.jpg'), width=Cm(WIDTH))
centered('以交互地图支持地理课堂的观察与探究', 'Guide Cover Meta', after=5, size=10.5)
centered('WebGIS-AI 项目组', 'Guide Cover Meta', before=20, after=4)
centered('2026 年 10 月', 'Guide Cover Meta', after=0)

front = doc.add_section(WD_SECTION_START.NEW_PAGE)
section_settings(front)
page_number(front, 1, 'lowerRoman')
doc.add_paragraph('阅读导引', 'Guide Front Heading')
intro = MD.split('## 阅读导引\n', 1)[1].split('## 目录\n', 1)[0].strip()
for block in intro.split('\n\n'):
    para(block.replace('\n', ''))
p = doc.add_paragraph('目录', 'Guide Front Heading')
p.paragraph_format.space_before = Pt(12)
field(doc.add_paragraph(), ' TOC \\o "1-2" \\h \\z \\u ', '更新目录')

main = doc.add_section(WD_SECTION_START.NEW_PAGE)
section_settings(main)
page_number(main)
main.header.paragraphs[0].text = ''
main.header.paragraphs[0].style = doc.styles['Header']

def set_cell_width(cell, cm):
    cell.width = Cm(cm)
    w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
    w.set(qn('w:type'), 'dxa'); w.set(qn('w:w'), str(round(cm / 2.54 * 1440)))

def make_table(lines):
    rows = [[x.strip() for x in l.strip().strip('|').split('|')] for l in lines]
    rows.pop(1)
    cols = len(rows[0])
    if cols == 2:
        widths = [4.1, 12.6]
    elif '建议用时' in rows[0]:
        widths = [3.2, 2.05, 11.45]
    else:
        widths = [3.25, 6.15, 7.3]
    assert len(widths) == cols
    t = doc.add_table(rows=len(rows), cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, width in zip(t.columns, widths):
        c.width = Cm(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + side)
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9')
        borders.append(b)
    pr.append(borders)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        nosplit = OxmlElement('w:cantSplit'); trpr.append(nosplit)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci, cell in enumerate(row.cells):
            set_cell_width(cell, widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cpr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), 'E8F1F7' if ri == 0 else ('F8FBFD' if ri % 2 == 0 else 'FFFFFF'))
            cpr.append(shd)
            margins = OxmlElement('w:tcMar')
            for side, value in [('top', 85), ('bottom', 85), ('left', 100), ('right', 100)]:
                mar = OxmlElement('w:' + side); mar.set(qn('w:w'), str(value)); mar.set(qn('w:type'), 'dxa'); margins.append(mar)
            cpr.append(margins)
            p = cell.paragraphs[0]
            p.style = doc.styles['Guide Table']
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            if '建议用时' in rows[0] and ri < len(rows) - 1:
                p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ri == 0 or ci == 0 or (cols == 3 and '建议用时' in rows[0] and ci == 1)) else WD_ALIGN_PARAGRAPH.LEFT
            inline(p, rows[ri][ci])
            if ri == 0:
                for run in p.runs:
                    run.bold = True
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.line_spacing = 1
    gap.add_run().font.size = Pt(3)

def number_definition():
    root = doc.part.numbering_part.element
    aid = max([int(x.get(qn('w:abstractNumId'))) for x in root.findall(qn('w:abstractNum'))] + [-1]) + 1
    nid = max([int(x.get(qn('w:numId'))) for x in root.findall(qn('w:num'))] + [0]) + 1
    a = OxmlElement('w:abstractNum'); a.set(qn('w:abstractNumId'), str(aid))
    lvl = OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'), '0')
    for tag, value in [('start', '1'), ('numFmt', 'decimal'), ('lvlText', '%1.'), ('lvlJc', 'left')]:
        el = OxmlElement('w:' + tag); el.set(qn('w:val'), value); lvl.append(el)
    pp = OxmlElement('w:pPr'); ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '360'); ind.set(qn('w:hanging'), '240'); pp.append(ind); lvl.append(pp)
    a.append(lvl)
    first_num = root.find(qn('w:num'))
    root.insert(list(root).index(first_num) if first_num is not None else len(root), a)
    num = OxmlElement('w:num'); num.set(qn('w:numId'), str(nid))
    ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), str(aid)); num.append(ref); root.append(num)
    return nid

fig = 0
body = '## 一 产品概览\n' + MD.split('## 一 产品概览\n', 1)[1]
lines = body.splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1; continue
    if line.startswith('### '):
        p = doc.add_paragraph(line[4:], 'Heading 2')
        i += 1; continue
    if line.startswith('## '):
        p = doc.add_paragraph(line[3:], 'Heading 1')
        if line.startswith(('## 五 ', '## 六 ', '## 七 ')):
            p.paragraph_format.page_break_before = True
        if line == '## 一 产品概览':
            p.paragraph_format.page_break_before = False
        i += 1; continue
    if line.startswith('|'):
        ls = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            ls.append(lines[i]); i += 1
        make_table(ls); continue
    im = re.fullmatch(r'!\[([^]]+)\]\(([^)]+)\)', line)
    if im:
        path = OUT / im.group(2)
        assert path.exists(), path
        fig += 1
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.add_run().add_picture(str(path), width=Cm(WIDTH))
        dp = p._p.xpath('.//wp:docPr')[0]; dp.set('descr', im.group(1))
        cp = doc.add_paragraph('图 ' + str(fig) + '  ' + im.group(1), 'Guide Caption')
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(8)
        i += 1; continue
    if re.match(r'^\d+\. ', line):
        nid = number_definition()
        while i < len(lines) and re.match(r'^\d+\. ', lines[i].strip()):
            p = para(re.sub(r'^\d+\. ', '', lines[i].strip()))
            np = p._p.get_or_add_pPr().get_or_add_numPr()
            np.get_or_add_ilvl().val = 0
            np.get_or_add_numId().val = nid
            i += 1
        continue
    block = [line]
    i += 1
    while i < len(lines) and lines[i].strip() and not re.match(r'^(#|\||!\[|\d+\. )', lines[i].strip()):
        block.append(lines[i].strip()); i += 1
    para(''.join(block))

settings = doc.settings.element
update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); settings.append(update)
no_compress = OxmlElement('w:doNotAutoCompressPictures'); settings.append(no_compress)
doc.save(OUT / '使用说明与介绍.docx')
(OUT / '使用说明与介绍.md').write_text(MD, encoding='utf-8')
levels = Counter(p.style.name for p in doc.paragraphs)
print(json.dumps({'chapters': levels['Heading 1'], 'subsections': levels['Heading 2'], 'interface_figures': fig, 'inline_images': len(doc.inline_shapes), 'tables': len(doc.tables)}, ensure_ascii=True))
